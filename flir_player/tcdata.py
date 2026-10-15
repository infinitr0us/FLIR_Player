# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""Thermocouple logger tables: reading, validity windows and failure flags (Qt-free).

A logger export is a sheet (.xlsx) or text table (.csv/.txt) with one time
column and one column per thermocouple, in °C. Layouts vary, so the data block
is found rather than assumed: the longest run of rows whose time cell holds a
number or a date and that carry at least one other number. Rows above it are
headers (their text names the channels; unit-only cells such as ``[sec]`` are
skipped); anything else outside the block is reported as ignored, so a note
such as 0922's ``shft time`` is not silently lost.

Failure flags use the thermocouple data alone, never the IR, so choosing data
by them cannot bias an emissivity fit. A TC is taken to have failed from the
first moment it shows a sustained impossible reading (well below its own
starting level, or outside the plausible range); isolated spikes only drop
themselves.
"""
from __future__ import annotations

import csv
import math
import re
import warnings
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np

_UNIT_ONLY = re.compile(r"^[\[\(]?\s*(°\s*[CFK]|deg\.?\s*[CFK]|[CFK]|degrees?\s*[CFK]?|s|sec|secs|seconds|"
                        r"min|h|hh:mm:ss)\s*[\]\)]?$", re.IGNORECASE)
# whole words (after splitting "ElapsedTime" and "Time_s"): "Section" is no time, "Template" no temperature
_TIME_HINT = re.compile(r"(?<![a-z])(time|times|timestamp|timestamps|sec|secs|second|seconds|elapsed|elapsedtime|"
                        r"date|datetime|clock)(?![a-z])", re.IGNORECASE)  # compounds without case boundaries too
_SHEET_HINT = re.compile(r"(?<![a-z])(tcs?|thermo\w*|temp|temps|temperatures?)(?![a-z])", re.IGNORECASE)
WORKBOOK_SUFFIXES = (".xlsx", ".xlsm")
TEXT_SUFFIXES = (".csv", ".txt", ".tsv", ".dat")
_TEMPERATURE_UNIT = re.compile(r"°\s*[CFK]\b|\bdeg\.?\s*[CFK]\b|^[\[\(]\s*[CFK]\s*[\]\)]$|^[CFK]$", re.IGNORECASE)
# Fahrenheit or kelvin said outright: refused, since readings are taken as °C (a bare "K" is left alone: it usually
# names a type-K TC; a bare "(F)" may name a position, like "(S)" for south)
_NOT_CELSIUS = re.compile(r"°\s*[FK]\b|\bdeg\.?\s*[FK]\b|\bfahrenheit\b|\bkelvins?\b", re.IGNORECASE)


def _unit_patterns(words: str, spelled: str) -> tuple[re.Pattern, re.Pattern, re.Pattern]:
    """(the cell is only the unit, the unit in brackets, the unit spelled out)."""
    return (re.compile(rf"[\[\(]?\s*({words})\s*[\]\)]?", re.IGNORECASE),
            re.compile(rf"[\[\(]\s*({words})\s*[\]\)]", re.IGNORECASE),
            re.compile(rf"(?<![a-z]){spelled}(?![a-z])", re.IGNORECASE))


_TIME_UNITS = ((1.0, _unit_patterns("s|sec|secs|second|seconds", "seconds")),
               (60.0, _unit_patterns("min|mins|minute|minutes", "minutes")),
               (3600.0, _unit_patterns("h|hr|hrs|hour|hours", "hours")))


def _words(text: str) -> str:
    """ "ElapsedTime" → "Elapsed Time", "TCData" → "TC Data", "Time_s" → "Time s" (for whole-word hints)."""
    text = re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", text)
    return text.replace("_", " ")
_TIME_RATIOS = (1.0, 60.0, 3600.0, 1 / 60.0, 1 / 3600.0)  # one time column against another in s, min or h


@dataclass(frozen=True)
class TcTable:
    """Logger data: ``values[row, channel]`` in °C at ``times[row]`` seconds."""

    times: np.ndarray
    values: np.ndarray
    names: tuple[str, ...]
    source: str = ""
    time_label: str = ""
    notes: tuple[str, ...] = ()
    sheet: str = ""  # the workbook sheet read ("" for a text table)
    # the sheet rows (1-based, one per row of ``times``) and column of the time cells, and the time unit in s
    # (for the logger clock)
    time_rows: tuple[int, ...] = ()
    time_col: int = 0
    time_unit_s: float = 1.0
    clock_zero: datetime | None = None  # the logger's wall clock at logger time 0 (a lab workbook's LogFile sheet)
    clock_note: str = ""  # where clock_zero comes from, or why there is none

    @property
    def interval(self) -> float:
        """Median sampling interval (s); 1.0 when it cannot be told."""
        steps = np.diff(self.times)
        steps = steps[np.isfinite(steps) & (steps > 0)]
        return float(np.median(steps)) if steps.size else 1.0

    def index(self, key: str | int) -> int:
        """Channel index from a 1-based number, an exact name, or a unique name prefix."""
        if isinstance(key, (int, np.integer)):
            if not 1 <= int(key) <= len(self.names):
                raise KeyError(f"No TC channel {key} (1 to {len(self.names)})")
            return int(key) - 1
        text = str(key).strip()
        if text.isdigit():
            return self.index(int(text))
        for k, name in enumerate(self.names):
            if name.casefold() == text.casefold():
                return k
        matches = [k for k, name in enumerate(self.names) if name.casefold().startswith(text.casefold())]
        if len(matches) == 1:
            return matches[0]
        raise KeyError(f"No single TC channel named {text!r} (channels: {', '.join(self.names)})")


# --- reading ---------------------------------------------------------------------------


def read_tc_table(path: str | Path, *, sheet: str | None = None, time_column: int | str | None = None) -> TcTable:
    """Read a logger export (.xlsx/.xlsm/.csv/.txt/.tsv); raises ValueError on no data.

    Without ``sheet``, a workbook's sheet named like TC data ("TC", "Thermocouples",
    "Temperature") is preferred, else the one with the most rows; a note names the sheet used.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in WORKBOOK_SUFFIXES:
        sheets = _xlsx_rows(path, only=sheet)
    elif suffix in TEXT_SUFFIXES:
        sheets = {path.stem: _text_rows(path)}
    else:
        raise ValueError(f"{path.name}: unsupported TC file type (use .xlsx, .csv or .txt)")
    if sheet is not None:
        if sheet not in sheets:
            raise ValueError(f"{path.name} has no sheet {sheet!r} (sheets: {', '.join(sheet_names(path))})")
        sheets = {sheet: sheets[sheet]}
    tables: list[tuple[str, TcTable]] = []
    errors = []
    for name, rows in sheets.items():
        try:
            tables.append((name, parse_rows(rows, time_column=time_column, source=f"{path.name}" +
                                            (f" [{name}]" if len(sheets) > 1 or sheet else ""))))
        except ValueError as exc:
            errors.append(f"{name}: {exc}")
    if not tables:
        raise ValueError(f"{path.name}: no logger data found ({'; '.join(errors) or 'empty file'})")
    hinted = [item for item in tables if _SHEET_HINT.search(_words(item[0]))]
    name, best = max(hinted or tables, key=lambda item: item[1].times.size)
    if len(tables) > 1:
        why = "its name suggests TC data" if hinted else "it has the most rows"
        others = ", ".join(repr(other) for other, _t in tables if other != name)
        best = replace(best, notes=best.notes + (f"sheet {name!r} used ({why}); other sheets with data: {others}",))
    if suffix not in WORKBOOK_SUFFIXES:
        return replace(best, sheet="", clock_note="a text table does not give the logger's clock")
    clock, note = logger_clock(path, name, best)
    return replace(best, sheet=name, clock_zero=clock, clock_note=note)


def sheet_names(path: str | Path) -> list[str]:
    """The sheets of a workbook (fast: no cells are read); a text table counts as one sheet."""
    path = Path(path)
    if path.suffix.lower() not in WORKBOOK_SUFFIXES:
        return [path.stem]
    book = _open_workbook(path)
    try:
        return [ws.title for ws in book.worksheets]
    finally:
        book.close()


def suggest_sheet(names: Sequence[str]) -> str | None:
    """The sheet whose name suggests TC data ("TC", "Thermocouples", "Temperature"), if exactly one does."""
    hinted = [name for name in names if _SHEET_HINT.search(_words(name))]
    return hinted[0] if len(hinted) == 1 else None


# A lab workbook (the engineers' template) keeps the raw log as a sheet "LogFile_<start>" whose "Date/Time"
# cell says when logging started, next to a "Log" sheet of elapsed times, and computes the TC times as
# "=Log!E362-Log!E$360": logger time 0 is the elapsed time of the row subtracted.
_LOGFILE_SHEET = re.compile(r"^\s*logfile", re.IGNORECASE)
_LOGFILE_STAMP = re.compile(r"(\d{4})-(\d{2})-(\d{2})[-_ ](\d{2})-(\d{2})-(\d{2})")
_START_HEADER = re.compile(r"^\s*date\s*/?\s*time\s*$", re.IGNORECASE)
_REF = r"(?:'(?P<{0}q>(?:[^']|'')+)'|(?P<{0}n>[A-Za-z_][\w.]*))!\$?(?P<{0}c>[A-Za-z]{{1,3}})\$?(?P<{0}r>\d+)"
_DIFFERENCE = re.compile(r"^=\s*\(?\s*" + _REF.format("a") + r"\s*-\s*" + _REF.format("b") + r"\s*\)?\s*$")


def logger_clock(path: str | Path, sheet: str, table: TcTable) -> tuple[datetime | None, str]:
    """The logger's wall clock at logger time 0 of ``table`` (read from ``sheet`` of the workbook ``path``)
    and where it comes from; (None, why not) when the workbook does not tell.

    Only a lab workbook tells (see ``_DIFFERENCE``), and only when every time cell the table was read from
    subtracts the same Log row of an elapsed-time column that starts at 0 (the start of logging) and holds
    the difference its formula names, so a half-edited sheet, another column or a stale value gives no
    clock rather than a wrong one. Never raises.
    """
    if not table.time_rows or table.time_col < 1:
        return None, "the TC times do not come from a sheet"
    try:
        return _logger_clock(Path(path), sheet, table)
    except Exception as exc:  # any other layout: no clock, never an error
        return None, f"the logger's clock could not be read: {type(exc).__name__} {exc}"


_NOT_FORMULAS = "the TC times are not formulas on the logger's elapsed time, as in a lab workbook"


def _logger_clock(path: Path, sheet: str, table: TcTable) -> tuple[datetime | None, str]:
    import openpyxl
    from openpyxl.utils import column_index_from_string

    kept = set(table.time_rows)  # rows the reader dropped (no readable time, a repeated time) are not checked
    first, last, col = min(kept), max(kept), table.time_col
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        formulas = _column(book[sheet], col, first, last)
        logs = [ws.title for ws in book.worksheets if _LOGFILE_SHEET.match(ws.title)]
    finally:
        book.close()
    refs: dict[int, tuple[str, str, int, int]] = {}  # sheet row → (source sheet, column, its row, row of time 0)
    for row, formula in zip(range(first, last + 1), formulas):
        if row not in kept:
            continue
        match = _DIFFERENCE.match(formula.strip()) if isinstance(formula, str) else None
        if match is None:
            return None, _NOT_FORMULAS
        (a_sheet, a_col, a_row), (b_sheet, b_col, b_row) = (
            ((match[f"{k}q"].replace("''", "'") if match[f"{k}q"] else match[f"{k}n"]), match[f"{k}c"].upper(),
             int(match[f"{k}r"])) for k in "ab")
        if (a_sheet, a_col) != (b_sheet, b_col):
            return None, "the TC time formulas do not subtract one logger time from another"
        refs[row] = (a_sheet, a_col, a_row, b_row)
    if not refs:
        return None, _NOT_FORMULAS
    origins = sorted({(s, c, z) for s, c, _a, z in refs.values()})
    if len(origins) != 1:
        rows = ", ".join(f"{s}!{c}{z}" for s, c, z in origins[:3]) + (", ..." if len(origins) > 3 else "")
        return None, f"the TC time formulas subtract different logger times ({rows}): a half-edited sheet"
    (source, letter, zero_row), = origins
    if len(logs) != 1:
        return None, ("no LogFile sheet gives the logger's start" if not logs
                      else f"{len(logs)} LogFile sheets: which one started the log is unclear")
    end = max(zero_row, max(a for _s, _c, a, _z in refs.values()))
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        shown = _column(book[sheet], col, first, last)
        elapsed = _column(book[source], column_index_from_string(letter), 1, end)
        start = _log_start(book[logs[0]])
    finally:
        book.close()
    header = elapsed[0]
    numbers = [_number(v) for v in elapsed[1:]]  # sheet rows 2 to end
    if not (isinstance(header, str) and re.search(r"time", header, re.IGNORECASE)):
        return None, f"{source}!{letter}1 does not head an elapsed time"
    if next((v for v in numbers if v is not None), None) != 0:
        return None, f"{source}!{letter} does not start at 0, so it does not count from the start of logging"

    def at(row: int) -> float | None:
        return numbers[row - 2] if 2 <= row <= end else None

    zero = at(zero_row)
    for row, (_s, _c, a_row, _z) in refs.items():
        now, cell = at(a_row), _number(shown[row - first])
        if zero is None or now is None or cell is None:
            return None, f"the logger times at {source}!{letter}{a_row} and {letter}{zero_row} are not numbers"
        if abs((now - zero) - cell) > 1e-6 * max(1.0, abs(cell)):
            return None, (f"the TC time in row {row} does not hold the difference its formula names; the sheet "
                          "may need recalculating")
    source_text = "its Date/Time cell"
    if start is None:
        stamp = _LOGFILE_STAMP.search(logs[0])
        if stamp is None:
            return None, f"sheet {logs[0]!r} does not say when logging started"
        start, source_text = datetime(*(int(v) for v in stamp.groups())), "its name"
    offset = zero * table.time_unit_s
    return start + timedelta(seconds=offset), (
        f"logging started {start:%Y-%m-%d %H:%M:%S} by sheet {logs[0]!r} ({source_text}); logger time 0 is "
        f"{source}!{letter}{zero_row}, {offset:g} s later")


def _column(ws, col: int, first: int, last: int) -> list[Any]:
    """The values of column ``col`` in rows ``first`` to ``last`` (None past the sheet's end)."""
    cells = [row[0] if row else None
             for row in ws.iter_rows(min_row=first, max_row=last, min_col=col, max_col=col, values_only=True)]
    return (cells + [None] * (last - first + 1))[:last - first + 1]


def _log_start(ws) -> datetime | None:
    """The date and time under a "Date/Time" header in the first rows of a raw log sheet."""
    rows = [tuple(r) for r in ws.iter_rows(min_row=1, max_row=6, values_only=True)]
    for i, cells in enumerate(rows[:-1]):
        for j, text in enumerate(cells):
            if isinstance(text, str) and _START_HEADER.match(text):
                below = rows[i + 1][j] if j < len(rows[i + 1]) else None
                if isinstance(below, datetime):
                    return below
    return None


def _open_workbook(path: Path):
    try:
        import openpyxl
    except ModuleNotFoundError as exc:  # pragma: no cover - environment
        raise ValueError("Reading .xlsx TC files needs the openpyxl package") from exc
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def _xlsx_rows(path: Path, only: str | None = None) -> dict[str, list[tuple]]:
    """Cell rows of every sheet, or of the sheet ``only`` (an absent name gives no sheets)."""
    book = _open_workbook(path)
    try:
        return {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in book.worksheets
                if only is None or ws.title == only}
    finally:
        book.close()


def _text_rows(path: Path) -> list[tuple]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    rows = []
    for raw in csv.reader(text.splitlines(), delimiter=delimiter):
        rows.append(tuple(_text_cell(cell) for cell in raw))
    return rows


def _text_cell(cell: str) -> Any:
    cell = cell.strip()
    if not cell:
        return None
    try:
        return float(cell)
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S",
                "%m/%d/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%H:%M:%S.%f", "%H:%M:%S"):
        try:
            stamp = datetime.strptime(cell, fmt)
        except ValueError:
            continue
        return stamp.time() if fmt.startswith("%H") else stamp
    return cell


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        v = float(value)
        return v if math.isfinite(v) else None
    return None


def _time_like(value: Any) -> bool:
    return _number(value) is not None or isinstance(value, (datetime, date, dtime, timedelta))


def parse_rows(rows: Sequence[Sequence[Any]], *, time_column: int | str | None = None, source: str = "") -> TcTable:
    """Find the data block in raw cell rows and build a table (see the module docstring)."""
    rows = [tuple(row) for row in rows]
    width = max((len(row) for row in rows), default=0)
    rows = [row + (None,) * (width - len(row)) for row in rows]
    if not rows or width < 2:
        raise ValueError("fewer than two columns")

    def is_data(row: tuple) -> bool:
        """Two numbers or dates (time and a reading; text such as a status may sit beside them), or one
        with nothing else (a time whose readings are missing)."""
        timed = sum(_time_like(v) for v in row)
        if timed >= 2:
            return True
        return timed == 1 and not any(isinstance(v, str) and v.strip() for v in row)

    data = [is_data(row) for row in rows]
    # a time with text but no other number (a reading "OPEN", or a blank reading beside a status) is
    # a data row when data rows surround it; above the data such a row is a note
    weak = [sum(_time_like(v) for v in row) == 1 and not d for row, d in zip(rows, data)]
    for i in range(len(rows)):
        if weak[i]:
            above = next((data[j] for j in range(i - 1, -1, -1) if not weak[j]), False)
            below = next((data[j] for j in range(i + 1, len(rows)) if not weak[j]), False)
            data[i] = above and below
    # longest run of data rows
    best = (0, -1)
    start = None
    for i, row in enumerate(rows + [(None,) * width]):
        if i < len(rows) and data[i]:
            start = i if start is None else start
            continue
        if start is not None and i - start > best[1] - best[0] + 1:
            best = (start, i - 1)
        start = None
    first, last = best

    def filled_cells(row: tuple) -> int:
        return sum(v is not None and not (isinstance(v, str) and not v.strip()) for v in row)

    # a lone value at an edge (a note such as 0922's "2" above the data) is not a data row
    while first < last and filled_cells(rows[first]) <= 1:
        first += 1
    while last > first and filled_cells(rows[last]) <= 1:
        last -= 1
    if last < first or last - first + 1 < 3:
        raise ValueError("no block of at least three data rows")
    block = rows[first:last + 1]
    headers = rows[:first]
    cut = []  # rows left out where the time column stops rising (notes)
    try:
        tcol, scale = _time_column(block, headers, width, time_column)
    except ValueError as exc:
        if "no increasing time column" not in str(exc):
            raise
        # A lab workbook whose time formulas were not filled down after its shift row changed, or
        # that runs past the end of its log, rises only between such rows: use that stretch.
        tcol, a, z = _rising_run(block, headers, width, time_column)
        if a > 0:
            cut.append(f"rows {first + 1}-{first + a}" if a > 1 else f"row {first + 1}")
        if z < len(block) - 1:
            cut.append(f"rows {first + z + 2}-{last + 1}" if z < len(block) - 2 else f"row {last + 1}")
        scale = _time_column(block[a:z + 1], headers, width, tcol)[1]
        # Times as read over the whole block (dates and clock times keep their first row as zero);
        # the rows left out are blanked so that they name no channel and are not reported again.
        seconds = _seconds([row[tcol] for row in block])
        for i in [*range(a), *range(z + 1, len(block))]:
            rows[first + i] = (None,) * width
        block = [row[:tcol] + ((float(seconds[i]),) if np.isfinite(seconds[i]) else (row[tcol],)) + row[tcol + 1:]
                 for i, row in enumerate(block) if a <= i <= z]
        first, last = first + a, first + z
        headers = rows[:first]
    timed = [_time_like(row[tcol]) for row in block]  # rows at the edges without a time are not data
    lead = timed.index(True) if True in timed else 0
    trail = len(timed) - 1 - timed[::-1].index(True) if True in timed else len(timed) - 1
    headers = rows[:first + lead]
    block = block[lead:trail + 1]
    last = first + trail
    first = first + lead
    channels = [c for c in range(width) if c != tcol
                and sum(_number(row[c]) is not None for row in block) >= 0.5 * len(block)]
    counters = [c for c in channels if _counter([row[c] for row in block])]
    raw_times = _seconds([row[tcol] for row in block])
    # dropping a column needs a time word ("Time From Ignition"); a unit alone is too weak for that
    copies = [c for c in channels if c not in counters and _time_words(headers, c)
              and not _temperature_unit(headers, c) and _time_copy([row[c] for row in block], raw_times)]
    channels = [c for c in channels if c not in counters and c not in copies]
    if not channels:
        raise ValueError("no numeric TC columns next to the time column")
    other = [c for c in channels if any(isinstance(row[c], str) and _NOT_CELSIUS.search(row[c]) for row in headers)]
    if other:
        raise ValueError("TC readings must be in °C, but " + ", ".join(
            repr(name) for name in _names(headers, other)) + " " + ("is" if len(other) == 1 else "are")
            + " headed in °F or kelvin")
    # rows without any reading at the edges of the block are not data (e.g. a lone number above it)
    filled = [any(_number(row[c]) is not None for c in channels) for row in block]
    lead = filled.index(True)
    trail = len(filled) - 1 - filled[::-1].index(True)
    headers = rows[:first + lead]
    block = block[lead:trail + 1]
    last = first + trail
    first = first + lead
    if len(block) < 3:
        raise ValueError("no block of at least three data rows")
    times = _seconds([row[tcol] for row in block]) * scale
    sheet_rows = np.arange(first + 1, first + 1 + len(block))  # 1-based, kept with ``times`` below
    values = np.array([[_number(row[c]) if _number(row[c]) is not None else np.nan for c in channels]
                       for row in block], dtype=np.float64)
    keep = np.isfinite(times)
    notes = []
    if scale != 1.0:
        notes.append(f"the time column is in {'minutes' if scale == 60.0 else 'hours'}; converted to seconds")
    if cut:
        notes.append(f"the time column stops rising, so {' and '.join(cut)} were left out (in a lab workbook: "
                     "time formulas not filled down after the shift row changed, or rows past the end of the log)")
    if not keep.all():
        notes.append(f"{int((~keep).sum())} row(s) without a readable time were skipped")
    times, values, sheet_rows = times[keep], values[keep], sheet_rows[keep]
    order = np.argsort(times, kind="stable")
    if np.any(np.diff(order) < 0):
        notes.append("rows were sorted by time")
    times, values, sheet_rows = times[order], values[order], sheet_rows[order]
    duplicate = np.concatenate([[False], np.diff(times) <= 0])
    if duplicate.any():
        notes.append(f"{int(duplicate.sum())} row(s) repeating an earlier time were dropped")
        times, values, sheet_rows = times[~duplicate], values[~duplicate], sheet_rows[~duplicate]
    # header rows name the channels; other text above the data is reported, not used
    titled = [row for row in headers if any(isinstance(row[c], str) and row[c].strip() for c in channels)]
    names = _names([row for row in titled if not _uniform(row, channels)] or titled, channels)
    time_label = " ".join(reversed(_texts(titled, tcol))) or "Time"
    ignored = []
    for i, row in enumerate(headers):
        texts = [str(v) for v in row if v is not None and str(v).strip()]
        if texts and not any(row is t for t in titled):
            ignored.append(f"row {i + 1} ({', '.join(texts)[:40]})")
    trailing = len(rows) - last - 1 - sum(1 for row in rows[last + 1:] if all(v is None for v in row))
    if ignored:
        notes.append("ignored above the data: " + "; ".join(ignored))
    for c in counters:
        label = " ".join(reversed(_texts(titled, c))) or f"column {c + 1}"
        notes.append(f"column {label!r} counts rows (1, 2, 3, ...), so it is not used as a TC")
    for c in copies:
        label = " ".join(reversed(_texts(titled, c))) or f"column {c + 1}"
        notes.append(f"column {label!r} is the time again in other units, so it is not used as a TC")
    if trailing > 0:
        notes.append(f"{trailing} non-empty row(s) after the data block were ignored")
    return TcTable(times=times, values=values, names=names, source=source, time_label=time_label,
                   notes=tuple(notes), time_rows=tuple(int(r) for r in sheet_rows), time_col=tcol + 1,
                   time_unit_s=scale)


def _counter(cells: list[Any]) -> bool:
    """A row counter (0 or 1, then one more each row) rather than a measurement."""
    values = [_number(v) for v in cells]
    if len(values) < 3 or any(v is None for v in values) or values[0] not in (0.0, 1.0):
        return False
    return all(b - a == 1 for a, b in zip(values, values[1:]))


def _time_copy(cells: list[Any], times: np.ndarray) -> bool:
    """The time column again in other units (minutes beside seconds), not a measurement."""
    values = np.array([np.nan if _number(v) is None else _number(v) for v in cells], dtype=np.float64)
    ok = np.isfinite(values) & np.isfinite(times)
    if ok.sum() < 3 or not np.ptp(times[ok]) > 0:
        return False
    x, y = times[ok], values[ok]
    slope, intercept = np.polyfit(x, y, 1)
    if not any(math.isclose(slope, ratio, rel_tol=1e-6) for ratio in _TIME_RATIOS):
        return False
    return float(np.max(np.abs(y - (slope * x + intercept)))) <= 1e-6 * max(float(np.ptp(y)), 1.0)


def _uniform(row: tuple, channels: list[int]) -> bool:
    """One text over all channels (e.g. FlexLogger's "Free_0"): it tells no channel apart."""
    texts = [row[c].strip() for c in channels
             if isinstance(row[c], str) and row[c].strip() and not _UNIT_ONLY.match(row[c].strip())]
    return len(texts) >= 2 and len(set(texts)) == 1


def _time_words(headers: list[tuple], column: int) -> bool:
    return any(_TIME_HINT.search(_words(t)) for t in _texts(headers, column))


def _temperature_unit(headers: list[tuple], column: int) -> bool:
    return any(isinstance(row[column], str) and _TEMPERATURE_UNIT.search(row[column].strip()) for row in headers)


def _time_evidence(headers: list[tuple], column: int) -> bool:
    """The header says time (a time word, or a time unit such as "[min]") and gives no temperature unit."""
    if _temperature_unit(headers, column):
        return False
    return _time_words(headers, column) or _time_scale(headers, column) is not None


def _time_scale(headers: list[tuple], column: int) -> float | None:
    """Seconds per unit of a time column from its header (1, 60 or 3600); None when no unit is given.

    A unit counts when the cell is only the unit ("[min]", "s"), when it is spelled out ("minutes"),
    or when it is in brackets beside a time word ("Time [s]"); a letter in brackets inside a name
    ("TC (S)", south) does not.
    """
    for row in reversed(headers):  # nearest the data first
        value = row[column]
        if not (isinstance(value, str) and value.strip()):
            continue
        text = value.strip()
        worded = bool(_TIME_HINT.search(_words(text)))
        for factor, (alone, bracketed, spelled) in _TIME_UNITS:
            if alone.fullmatch(text) or spelled.search(text) or (worded and bracketed.search(text)):
                return factor
    return None


def _texts(headers: list[tuple], column: int) -> list[str]:
    out = []
    for row in reversed(headers):  # nearest the data first
        value = row[column]
        if isinstance(value, str) and value.strip() and not _UNIT_ONLY.match(value.strip()):
            text = " ".join(value.split())
            if text not in out:
                out.append(text)
    return out


def _names(headers: list[tuple], channels: list[int]) -> tuple[str, ...]:
    names = []
    for k, c in enumerate(channels):
        texts = _texts(headers, c)
        name = texts[0] if texts else f"TC {k + 1}"
        if len(texts) > 1:
            name += " (" + ", ".join(texts[1:]) + ")"
        names.append(name)
    seen: dict[str, int] = {}
    unique = []
    for name in names:
        seen[name] = seen.get(name, 0) + 1
        unique.append(name if seen[name] == 1 else f"{name} [{seen[name]}]")
    return tuple(unique)


def _time_column(block: list[tuple], headers: list[tuple], width: int, requested) -> tuple[int, float]:
    """(column, seconds per unit) of the time column.

    Preferred: a header that says time (a word such as time, elapsed, date, clock, or a time unit),
    then a column in seconds or dates, then one without a unit, then minutes or hours (converted).
    """
    def scale_of(c: int) -> float:
        if any(isinstance(row[c], (datetime, date, dtime, timedelta)) for row in block):
            return 1.0
        return _time_scale(headers, c) or 1.0

    if requested is not None:
        if isinstance(requested, int):
            if not 0 <= requested < width:
                raise ValueError(f"time column {requested} is outside the table")
            return requested, scale_of(requested)
        for c in range(width):
            if any(str(requested).casefold() == text.casefold() for text in _texts(headers, c)):
                return c, scale_of(c)
        raise ValueError(f"no column headed {requested!r}")
    candidates = []
    headed = False  # a column headed as time holds times (rising or not)
    for c in range(width):
        if _temperature_unit(headers, c):  # a TC heating steadily is no time axis
            continue
        # between its first and last time, (nearly) every row holds one; rows at the edges need not
        # (a summary row such as "Max" under the data)
        headed = headed or _time_evidence(headers, c)  # even with too few times, or none, to be read
        values = [row[c] for row in block]
        ok = [_time_like(v) for v in values]
        if sum(ok) < 3:
            continue
        a, z = ok.index(True), len(ok) - 1 - ok[::-1].index(True)
        cells = [v for v, good in zip(values[a:z + 1], ok[a:z + 1]) if good]
        if len(cells) < 0.9 * (z - a + 1):
            continue
        seconds = _seconds(cells)
        if np.all(np.isfinite(seconds)) and np.all(np.diff(seconds) >= 0) and seconds[-1] > seconds[0]:
            hinted = _time_evidence(headers, c)
            dated = any(isinstance(v, (datetime, date, dtime, timedelta)) for v in cells)
            unit = 1.0 if dated else _time_scale(headers, c)
            rank = 0 if unit == 1.0 else (1 if unit is None else 2)
            candidates.append((not hinted, rank, c))
    if not candidates or (headed and all(unhinted for unhinted, _rank, _c in candidates)):
        # a time column that stops rising goes to _rising_run; a column without a time header (a scan
        # number, a reading) must not take its place
        raise ValueError("no increasing time column")
    c = min(candidates)[2]
    return c, scale_of(c)


def _rising_run(block: list[tuple], headers: list[tuple], width: int, requested) -> tuple[int, int, int]:
    """(column, first, last) of the longest stretch of rows over which a time column rises.

    For a block in which no column rises throughout. The stretch must hold at least half of the
    block's rows. When a column is headed as time, only such columns count (a TC heating steadily
    must not stand in for a broken time column); a column with a temperature unit never does. Then
    the longer stretch wins, then a column in seconds.
    """
    if isinstance(requested, int):
        columns = [requested]
    elif requested is not None:
        columns = [_time_column(block, headers, width, requested)[0]]
    else:
        timed = [c for c in range(width) if _time_evidence(headers, c) and not _temperature_unit(headers, c)]
        columns = timed or [c for c in range(width) if not _temperature_unit(headers, c)]
    best = None
    for c in columns:
        seconds = _seconds([row[c] for row in block])
        if np.isfinite(seconds).sum() < 3:
            continue
        rising = np.concatenate([[False], np.isfinite(seconds[1:]) & np.isfinite(seconds[:-1])
                                 & (np.diff(seconds) > 0)])
        start, a, z = 0, 0, -1
        for i in range(1, len(seconds) + 1):
            if i < len(seconds) and rising[i]:
                continue
            if i - 1 - start > z - a:
                a, z = start, i - 1
            start = i
        if z - a + 1 < max(3, 0.5 * len(block)):
            continue
        unit = _time_scale(headers, c)
        rank = 0 if unit == 1.0 else (1 if unit is None else 2)
        key = (-(z - a), rank, c)
        if best is None or key < best[0]:
            best = (key, c, a, z)
    if best is None:
        raise ValueError("no increasing time column")
    return best[1], best[2], best[3]


def _seconds(cells: Iterable[Any]) -> np.ndarray:
    """Seconds: numbers as given, dates and times relative to the first row."""
    cells = list(cells)
    out = np.full(len(cells), np.nan)
    origin = None
    for i, value in enumerate(cells):
        number = _number(value)
        if number is not None:
            out[i] = number
        elif isinstance(value, timedelta):
            out[i] = value.total_seconds()
        elif isinstance(value, datetime):
            origin = origin or value
            out[i] = (value - origin).total_seconds()
        elif isinstance(value, date):
            origin = origin or datetime.combine(value, dtime())
            out[i] = (datetime.combine(value, dtime()) - origin).total_seconds()
        elif isinstance(value, dtime):
            out[i] = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond * 1e-6
    if any(isinstance(v, dtime) for v in cells):  # times of day: unwrap midnight, start at 0
        finite = np.isfinite(out)
        steps = np.diff(out[finite])
        out[finite] = out[finite] + 86400 * np.concatenate([[0], np.cumsum(steps < -43200)])
        out = out - out[finite][0]
    return out


# --- validity ---------------------------------------------------------------------------


def parse_windows(text: str) -> tuple[tuple[float, float], ...]:
    """``"0-1150, 1300-1500"`` → ((0, 1150), (1300, 1500)); an open end may be left empty."""
    windows = []
    for part in str(text).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        match = re.fullmatch(r"(-?[\d.]*)\s*(?:-|to|\.\.)\s*(-?[\d.]*)", part)
        if not match:
            raise ValueError(f"{part!r} is not a time window such as 800-1150")
        low = float(match.group(1)) if match.group(1) else -math.inf
        high = float(match.group(2)) if match.group(2) else math.inf
        if not low < high:
            raise ValueError(f"time window {part!r} is empty")
        windows.append((low, high))
    return tuple(windows)


def window_mask(times: np.ndarray, windows: Sequence[tuple[float, float]]) -> np.ndarray:
    """True inside any window (inclusive); all True when no window is given."""
    times = np.asarray(times, dtype=np.float64)
    if not windows:
        return np.ones(times.shape, dtype=bool)
    mask = np.zeros(times.shape, dtype=bool)
    for low, high in windows:
        mask |= (times >= low) & (times <= high)
    return mask


@dataclass(frozen=True)
class ChannelFlags:
    """TC-only failure flags of one channel."""

    flagged: np.ndarray  # bool per row: do not use
    failed_from: float | None  # logger time of a sustained failure, None if none
    spikes: int
    baseline_c: float
    reasons: tuple[str, ...] = field(default_factory=tuple)


def failure_flags(table: TcTable, channel: int, *, baseline_s: float = 60.0, below_k: float = 15.0,
                  low_c: float = -40.0, high_c: float = 1400.0, spike_k: float = 50.0,
                  sustain: int = 3, sustain_s: float = 10.0) -> ChannelFlags:
    """Flag readings a working TC on a heated surface cannot produce.

    * below its own starting level by more than ``below_k`` (a heated cell does
      not cool far below ambient; 0922's TC1 dropped to -3300 °C);
    * outside ``low_c``..``high_c``;
    * isolated spikes: one reading more than ``spike_k`` away from both
      neighbours and from the local median.

    ``sustain`` impossible readings within ``sustain_s`` mean the TC failed:
    everything from the first of them on is flagged.
    """
    t = table.times
    x = table.values[:, channel]
    finite = np.isfinite(x)
    early = finite & (t <= t[0] + baseline_s)
    if early.sum() < 3:
        early = finite & (np.arange(x.size) < max(3, x.size // 10))
    # the lower of the first readings' median and the first minute's: a TC that heats from the
    # start must not make its own first readings look impossibly low
    levels = [float(np.median(x[np.flatnonzero(finite)[:5]]))] if finite.any() else []
    if early.any():
        levels.append(float(np.median(x[early])))
    baseline = min(levels) if levels else float("nan")
    impossible = finite & ((x < low_c) | (x > high_c))
    if math.isfinite(baseline):
        impossible |= finite & (x < baseline - below_k)
    reasons = []
    failed_from = None
    hits = np.flatnonzero(impossible)
    for k in range(hits.size):
        window = hits[(t[hits] >= t[hits[k]]) & (t[hits] <= t[hits[k]] + sustain_s)]
        if window.size >= sustain:
            failed_from = float(t[hits[k]])
            break
    flagged = ~finite | impossible
    if failed_from is not None:
        flagged |= t >= failed_from
        bad = x[impossible]
        reasons.append(f"failed from {failed_from:.0f} s (impossible readings {bad.min():.0f} to "
                       f"{bad.max():.0f} °C; it started at {baseline:.0f} °C)")
    spikes = 0
    if x.size >= 5:
        pad = np.pad(np.where(finite, x, np.nan), 2, mode="edge")
        windows = np.lib.stride_tricks.sliding_window_view(pad, 5)
        with warnings.catch_warnings():  # all-NaN windows
            warnings.simplefilter("ignore", RuntimeWarning)
            median = np.nanmedian(windows, axis=1)
        prev = np.concatenate([[np.nan], x[:-1]])
        nxt = np.concatenate([x[1:], [np.nan]])
        spike = (finite & (np.abs(x - median) > spike_k) & (np.abs(x - prev) > spike_k)
                 & (np.abs(x - nxt) > spike_k))
        spikes = int((spike & ~flagged).sum())
        flagged |= spike
        if spikes:
            reasons.append(f"{spikes} isolated spike(s) dropped")
    if impossible.any() and failed_from is None:
        reasons.append(f"{int(impossible.sum())} impossible reading(s) dropped")
    return ChannelFlags(flagged=flagged, failed_from=failed_from, spikes=spikes, baseline_c=baseline,
                        reasons=tuple(reasons))


def resample(table: TcTable, channel: int, grid: np.ndarray, usable: np.ndarray | None = None,
             max_gap: float | None = None) -> np.ndarray:
    """Channel values at ``grid`` seconds (linear), NaN off the data or across gaps.

    ``usable`` (bool per row) drops rows first; two neighbours further apart
    than ``max_gap`` (default 1.5 × the logger interval) are not interpolated.
    """
    t = table.times
    x = table.values[:, channel].astype(np.float64)
    ok = np.isfinite(x) if usable is None else (np.isfinite(x) & usable)
    t, x = t[ok], x[ok]
    out = np.full(np.shape(grid), np.nan)
    if t.size < 2:
        return out
    gap = 1.5 * table.interval if max_gap is None else max_gap
    grid = np.asarray(grid, dtype=np.float64)
    j = np.searchsorted(t, grid, side="right") - 1
    inside = (j >= 0) & (j < t.size)
    exact = inside & (t[np.clip(j, 0, t.size - 1)] == grid)
    between = inside & (j < t.size - 1)
    jj = np.clip(j, 0, t.size - 2)
    span = t[jj + 1] - t[jj]
    between &= span <= gap
    w = np.where(span > 0, (grid - t[jj]) / np.where(span > 0, span, 1), 0.0)
    out = np.where(between, x[jj] + w * (x[jj + 1] - x[jj]), out)
    out = np.where(exact, x[np.clip(j, 0, t.size - 1)], out)
    return out
