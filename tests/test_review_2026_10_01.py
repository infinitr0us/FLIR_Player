# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""Regressions for the 2026-10-01 review of v0.5.0 / v0.5.1 (finding IDs in the names)."""
from __future__ import annotations

import json
import os
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

openpyxl = pytest.importorskip("openpyxl")

from flir_player import workbook  # noqa: E402
from flir_player.excel_export import (  # noqa: E402
    ExportData,
    ExportOptions,
    _ClockSearch,
    load_roi_set,
)
from flir_player.jobs import JobCancelled  # noqa: E402
from flir_player.main_window import MainWindow  # noqa: E402
from flir_player.models import RoiShape  # noqa: E402
from flir_player.sdktime import TimestampRepair  # noqa: E402
from flir_player.workbook import write_workbook  # noqa: E402

from conftest import wait_until  # noqa: E402
from test_excel_workbook import _source  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "local" / "data"
# captured before the autouse guard in conftest replaces it for every test
_SHOW_ERROR = MainWindow._show_error


def _data(sources, rows=12, **options):
    return ExportData(times=np.arange(rows, dtype=float) - 2.0, sources=sources,
                      options=ExportOptions(bins=16, **options), created=datetime(2026, 10, 1),
                      versions={"tool": "test", "sdk": ""})


# --- Excel export -------------------------------------------------------------------------


def test_c01_export_bins_more_areas_when_their_pixels_overflow_the_sheet(tmp_path) -> None:
    from flir_player.excel_export import SourceSpec, run_export

    # 50 boxes of 400 pixels need 20,000 Pixels columns at the default limit;
    # the export must lower the limit (bins) instead of failing
    rois = tuple(RoiShape(k + 1, "rect", ((2 + (k % 10) * 21, 2 + (k // 10) * 21),
                                          (22 + (k % 10) * 21, 22 + (k // 10) * 21)), f"Box{k + 1}")
                 for k in range(50))
    dest = tmp_path / "many.xlsx"
    message = run_export(dest, [SourceSpec(path=SAMPLES / "2.seq", rois=rois)],
                         ExportOptions(start_s=0.0, end_s=2.0, step_s=1.0, sheets=frozenset()))
    assert dest.is_file() and "50 ROIs" in message
    pixels = openpyxl.load_workbook(dest, read_only=True)["Pixels"]
    assert pixels.max_column == 1 + 50 * 256  # every box binned (2 × 128 bins)


def test_a05_charts_plot_the_tc_series_from_its_hidden_column(tmp_path) -> None:
    path = tmp_path / "tc.xlsx"
    write_workbook(path, _data([_source("Camera A", 12, np.random.default_rng(2))]))
    with zipfile.ZipFile(path) as archive:
        charts = [archive.read(name).decode("utf-8") for name in archive.namelist()
                  if name.startswith("xl/charts/chart")]
    with_tc = [xml for xml in charts if "TC for " in xml]
    assert with_tc, "no chart carries a TC series"
    for xml in with_tc:
        assert '<c:plotVisOnly val="1"/>' not in xml  # Excel skips hidden columns otherwise


def _out_of_domain(clamped: bool):
    source = _source("Camera A", 12, np.random.default_rng(3))
    spot = source.rois[0]
    spot.counts["value"][0] = 0.0  # validation row 0: below the curve's domain
    spot.validation_clamped[0] = clamped
    return source


@pytest.mark.parametrize("clamped", [True, False])
def test_a01_validation_never_writes_a_nan_result(tmp_path, clamped) -> None:
    path = tmp_path / "validation.xlsx"
    write_workbook(path, _data([_out_of_domain(clamped)]))
    book = openpyxl.load_workbook(path, data_only=True)  # "nan" made this raise
    sheet = book["Validation"]
    row = next(r for r in range(1, sheet.max_row + 1)
               if sheet.cell(r, 2).value == "TC1" and sheet.cell(r, 5).value == 1)
    assert sheet.cell(row, 7).value == "#NUM!" and sheet.cell(row, 8).value == "#NUM!"
    summary = next(r for r in range(1, 12) if sheet.cell(r, 4).value is not None
                   and str(sheet.cell(r, 5).value or "").startswith(("PASS", "CHECK")))
    if clamped:  # skipped rows do not count
        assert str(sheet.cell(summary, 5).value).startswith("PASS")
    else:  # outside the domain although the SDK did not clamp: a real mismatch
        assert sheet.cell(summary, 4).value == "#NUM!" and sheet.cell(summary, 5).value == "CHECK"


def test_a06_an_absent_tc_sheet_imposes_no_column_limit(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(workbook, "EXCEL_COLUMNS", 30)  # TC Compare would need 15 + 8 × 3 = 39
    source = _source("Camera A", 6, np.random.default_rng(4))
    source.rois = [source.rois[0]] * 3  # three spots
    write_workbook(tmp_path / "no-tc.xlsx", _data([source], rows=6, sheets=frozenset()))
    with pytest.raises(ValueError, match="TC Compare"):
        write_workbook(tmp_path / "tc.xlsx", _data([source], rows=6, sheets=frozenset({"tc"})))


def test_a07_cancelling_late_leaves_no_temporary_files(tmp_path, monkeypatch) -> None:
    scratch = _scratch(tmp_path, monkeypatch)
    calls = []

    def abort():  # cancel at the fourth check: after Data, while writing Counts
        calls.append(1)
        return len(calls) >= 4

    dest = tmp_path / "cancel.xlsx"
    with pytest.raises(JobCancelled):
        write_workbook(dest, _data([_source("Camera A", 600, np.random.default_rng(5))], rows=600),
                       abort=abort)
    assert len(calls) == 4
    assert not dest.exists()
    assert list(scratch.iterdir()) == []


def _scratch(tmp_path, monkeypatch) -> Path:
    """A private system temp folder, to see what writing leaves behind."""
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    return scratch


def test_a07_a_failure_while_creating_sheets_leaves_no_temporary_files(tmp_path, monkeypatch) -> None:
    import xlsxwriter

    scratch = _scratch(tmp_path, monkeypatch)
    original, created = xlsxwriter.Workbook.add_worksheet, []

    def add_worksheet(self, name=None, *args, **kwargs):
        created.append(name)
        if len(created) == 3:  # two sheets already hold open row files
            raise OSError("simulated failure")
        return original(self, name, *args, **kwargs)

    monkeypatch.setattr(xlsxwriter.Workbook, "add_worksheet", add_worksheet)
    with pytest.raises(OSError, match="simulated"):
        write_workbook(tmp_path / "fail.xlsx", _data([_source("Camera A", 12, np.random.default_rng(6))]))
    assert list(scratch.iterdir()) == []


def test_a07_a_failure_inside_close_leaves_no_temporary_files(tmp_path, monkeypatch) -> None:
    from xlsxwriter.exceptions import FileCreateError
    from xlsxwriter.worksheet import Worksheet

    scratch = _scratch(tmp_path, monkeypatch)
    original = Worksheet._write_optimized_sheet_data  # constant-memory sheets

    def write_sheet_data(self):  # runs while the sheet's XML part is open
        if self.name == "Data":
            raise OSError("disk full")
        return original(self)

    monkeypatch.setattr(Worksheet, "_write_optimized_sheet_data", write_sheet_data)
    with pytest.raises(FileCreateError, match="disk full"):  # XlsxWriter wraps the OSError
        write_workbook(tmp_path / "full.xlsx", _data([_source("Camera A", 12, np.random.default_rng(7))]))
    assert list(scratch.iterdir()) == []


@pytest.mark.parametrize("stage", ["_data", "_counts", "_pixels", "_tc"])
def test_x12_every_long_sheet_honours_cancellation(tmp_path, monkeypatch, stage) -> None:
    scratch = _scratch(tmp_path, monkeypatch)
    polled = []

    def abort():  # called from _check, called from the sheet writer
        polled.append(sys._getframe(2).f_code.co_name)
        return polled[-1] == stage

    with pytest.raises(JobCancelled):
        write_workbook(tmp_path / "cancel.xlsx", _data([_source("Camera A", 12, np.random.default_rng(8))]),
                       abort=abort)
    assert polled[-1] == stage
    assert list(scratch.iterdir()) == []


def _first_columns(sheet, row: int) -> dict:
    columns = {}
    for col in range(1, sheet.max_column + 1):
        columns.setdefault(sheet.cell(row, col).value, col)
    return columns


def test_x8_charts_plot_the_tc_in_the_display_unit_column(tmp_path) -> None:
    import re

    from openpyxl.utils import get_column_letter

    path = tmp_path / "tc-unit.xlsx"
    write_workbook(path, _data([_source("Camera A", 12, np.random.default_rng(2))], unit="°F"))
    display = _first_columns(openpyxl.load_workbook(path)["TC Compare"], 9)["TC (display unit)"]
    with zipfile.ZipFile(path) as archive:
        xml = "".join(archive.read(name).decode("utf-8") for name in archive.namelist()
                      if name.startswith("xl/charts/chart"))
    series = next(block for block in xml.split("<c:ser>") if "TC for Camera A · TC1" in block)
    reference = re.search(r"<c:yVal><c:numRef><c:f>([^<]+)</c:f>", series).group(1)  # scatter chart
    assert reference.startswith(f"'TC Compare'!${get_column_letter(display)}$")


def test_x4_regression_helpers_need_an_ir_and_a_tc_sample(tmp_path) -> None:
    path = tmp_path / "helpers.xlsx"
    write_workbook(path, _data([_source("Camera A", 12, np.random.default_rng(2))]))
    tc = openpyxl.load_workbook(path)["TC Compare"]
    columns = _first_columns(tc, 9)
    for label in ("x·y", "x²"):
        assert tc.cell(10, columns[label]).value.startswith("=IF(AND(ISNUMBER('Counts'!")


def test_x11_summary_results_are_cached_from_the_data(tmp_path) -> None:
    path = tmp_path / "summary.xlsx"
    write_workbook(path, _data([_source("Camera A", 12, np.random.default_rng(3))]))
    book = openpyxl.load_workbook(path, data_only=True)
    summary, data = book["Summary"], book["Data"]
    times = [data.cell(r, 1).value for r in range(6, 18)]
    values = [data.cell(r, 3).value for r in range(6, 18)]  # first ROI (spot), display unit
    present = [(t, v) for t, v in zip(times, values) if isinstance(v, float)]
    first, last = summary["E4"].value, summary["E5"].value  # baseline rows (1-based)
    baseline = float(np.mean([v for v in values[first - 1:last] if isinstance(v, float)]))
    peak_time, peak = max(present, key=lambda pair: pair[1])
    row = 14  # first ROI's results
    assert summary.cell(row, 4).value == pytest.approx(baseline)
    assert summary.cell(row, 5).value == pytest.approx(peak)
    assert summary.cell(row, 6).value == pytest.approx(peak_time)
    assert summary.cell(row, 7).value == pytest.approx(peak - baseline)
    for j, cell in enumerate(("B6", "B7", "B8")):
        threshold = summary[cell].value
        reached = next((t for t, v in present if v >= threshold), None)
        cached = summary.cell(row, 8 + j).value
        assert cached == "not reached" if reached is None else cached == pytest.approx(reached)


def test_x6_the_clock_time_base_exports_end_to_end(tmp_path) -> None:
    from flir_player.excel_export import SourceSpec, run_export

    spot = RoiShape(1, "cursor", ((10.5, 10.5),), "Spot")
    times = {}
    for base in ("frames", "clock"):
        dest = tmp_path / f"{base}.xlsx"
        run_export(dest, [SourceSpec(path=SAMPLES / "1.ats", rois=(spot,))],
                   ExportOptions(step_s=0.5, time_base=base, sheets=frozenset()))
        sheet = openpyxl.load_workbook(dest, read_only=True)["Data"]
        times[base] = [row[0] for row in sheet.iter_rows(min_row=6, values_only=True) if row[0] is not None]
    assert times["clock"][0] == 0.0 and times["clock"][1] == 0.5  # ignition = frame 1
    assert abs(len(times["clock"]) - len(times["frames"])) <= 1


# --- ATS clock rollover ---------------------------------------------------------------------


def _written(tmp_path, when: datetime) -> Path:
    path = tmp_path / "clip.ats"
    path.write_bytes(b"x")
    os.utime(path, (when.timestamp(), when.timestamp()))
    return path


@pytest.mark.parametrize("order", [1, -1])
def test_a03_a_recording_over_new_year_keeps_counting_forward(tmp_path, order) -> None:
    # written 1 January 2026; the day counter rolls from day 365 (30 Dec 1976) to day 1
    path = _written(tmp_path, datetime(2026, 1, 1, 0, 30))
    raw = [datetime(1976, 12, 30, 23, 59, 59), datetime(1976, 1, 1, 0, 0, 0), datetime(1977, 1, 1, 0, 0, 1)]
    expected = [datetime(2025, 12, 31, 23, 59, 59), datetime(2026, 1, 1, 0, 0, 0), datetime(2026, 1, 1, 0, 0, 1)]
    repair = TimestampRepair(path)
    pairs = list(zip(raw, expected))[::order]
    assert [repair(stamp) for stamp, _ in pairs] == [want for _, want in pairs]


def test_a03_clock_search_measures_repaired_stamps(tmp_path) -> None:
    stamps = [datetime(1976, 12, 30, 23, 59, 58), datetime(1976, 12, 30, 23, 59, 59),
              datetime(1976, 1, 1, 0, 0, 0), datetime(1976, 1, 1, 0, 0, 1)]

    class Clock:
        def get_frame(self, index):
            self.index = index

        @property
        def frame_info(self):
            return SimpleNamespace(time=stamps[self.index])

    repair = TimestampRepair(_written(tmp_path, datetime(2026, 1, 1, 0, 30)))
    search = _ClockSearch(Clock(), 4, repair(stamps[0]), repair)
    assert [search.time(i) for i in range(4)] == [0.0, 1.0, 2.0, 3.0]
    assert search.nearest(2.2, guess=0, tolerance=0.5) == 2


# --- ROI sets -----------------------------------------------------------------------------------


def _roi_file(tmp_path, payload) -> Path:
    path = tmp_path / "set.rois.json"
    path.write_text(payload if isinstance(payload, str) else json.dumps(payload), encoding="utf-8")
    return path


GOOD = {"format": "flir-player-roi-set", "rois": [{"kind": "cursor", "points": [[3, 4]], "name": "TC1"}]}


@pytest.mark.parametrize("payload", [
    "[]",
    "not json",
    {"format": "flir-player-roi-set", "rois": {}},
    {"format": "flir-player-roi-set", "rois": ["box"]},
    {"format": "flir-player-roi-set", "rois": [{"kind": "rect", "points": [1, 2]}]},
    {"format": "flir-player-roi-set", "rois": [{"kind": "cursor", "points": [["NaN", 4]]}]},
    {"format": "flir-player-roi-set", "rois": [{"kind": "cursor", "points": [[1e999, 4]]}]},
    {"format": "flir-player-roi-set", "rois": [{"kind": "cursor", "points": [[10 ** 400, 4]]}]},
    dict(GOOD, ignition_frame=[3]),
    dict(GOOD, size="wide"),
])
def test_b04_malformed_roi_sets_raise_value_error(tmp_path, payload) -> None:
    with pytest.raises(ValueError):
        load_roi_set(_roi_file(tmp_path, payload))


def test_b05_roi_names_are_read_as_one_line(tmp_path) -> None:
    payload = {"format": "flir-player-roi-set",
               "rois": [{"kind": "cursor", "points": [[3, 4]], "name": "north\r\n  south\n"},
                        {"kind": "cursor", "points": [[5, 6]], "name": "\n"}]}
    names = [shape.name for shape in load_roi_set(_roi_file(tmp_path, payload)).rois]
    assert names == ["north south", "ROI 2"]


def test_b04_the_player_reports_a_malformed_roi_set(qapp, tmp_path, monkeypatch) -> None:
    from PySide6.QtWidgets import QFileDialog

    window = MainWindow()
    window.show()
    try:
        window.open_path(SAMPLES / "2.seq")
        assert wait_until(qapp, lambda: window.current_packet is not None and not window._busy)
        bad = _roi_file(tmp_path, {"format": "flir-player-roi-set",
                                   "rois": [{"kind": "cursor", "points": [["NaN", 4]]}]})
        monkeypatch.setattr(QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(bad), "")))
        errors = []
        monkeypatch.setattr(MainWindow, "_show_error", lambda self, title, message: errors.append(title))
        window._load_roi_set()
        assert errors == ["Load ROI set"]
    finally:
        window.close()
        assert wait_until(qapp, lambda: not window.decoder.isRunning())


# --- jobs and dialogs ---------------------------------------------------------------------------


def test_b02_a_bitmask_job_never_closes_another_exports_dialog() -> None:
    notices = []
    window = SimpleNamespace(_export_progress="Excel export's dialog", _notify=notices.append,
                             _show_error=lambda *a: notices.append(a))
    MainWindow._on_bitmasks_finished(window, True, "Wrote 1 bitmask file(s)")
    assert window._export_progress == "Excel export's dialog"
    assert notices == ["Wrote 1 bitmask file(s)"]


def test_b03_no_error_dialog_once_the_player_is_closing(monkeypatch) -> None:
    from flir_player.widgets import MessageDialog

    shown = []
    monkeypatch.setattr(MessageDialog, "critical", staticmethod(lambda *a, **k: shown.append(a[1])))
    _SHOW_ERROR(SimpleNamespace(_closing=True), "Export failed", "disk full")
    assert shown == []
    _SHOW_ERROR(SimpleNamespace(_closing=False), "Export failed", "disk full")
    assert shown == ["Export failed"]


# --- SDK preload --------------------------------------------------------------------------------


@pytest.mark.skipif(os.name != "nt", reason="the preload only runs on Windows")
def test_b07_preload_loads_the_sdk_dlls_by_full_path_in_order(tmp_path, monkeypatch) -> None:
    import flir_player

    package = tmp_path / "fnv"
    library = package / "_lib"
    library.mkdir(parents=True)
    for name in ("fnvfile.dll", "fnv.dll", "zlib1.dll", "CharLS.dll"):  # minizip/fnvreduce missing
        (library / name).write_bytes(b"")
    monkeypatch.delitem(sys.modules, "fnv", raising=False)
    monkeypatch.setattr(flir_player.importlib.util, "find_spec",
                        lambda name: SimpleNamespace(origin=str(package / "__init__.py"),
                                                     submodule_search_locations=[str(package)]))
    loaded = []

    def load(path, reserved, flags):  # kernel32.LoadLibraryExW
        loaded.append((Path(path), reserved, flags))
        return 1

    monkeypatch.setattr(flir_player.ctypes, "WinDLL", lambda *a, **k: SimpleNamespace(LoadLibraryExW=load))
    flir_player._preload_file_sdk()
    assert loaded == [(library / name, None, 0x8)
                      for name in ("CharLS.dll", "zlib1.dll", "fnv.dll", "fnvfile.dll")]
