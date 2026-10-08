# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""Emissivity from TCs in the player (v0.6.0): the TC file helpers, fitted emissivities in the
workbook, ``run_analysis`` and the setup and results dialogs, and the job through the decoder.
"""
from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QDialog, QLineEdit

from conftest import wait_until
from flir_player import tcmatch
from flir_player.excel_export import TcPrefill
from flir_player.models import PresetRange, RoiShape, VideoMetadata
from flir_player.tcdata import TcTable, read_tc_table, sheet_names, suggest_sheet
from flir_player.tcmatch import MatchOptions, RunOutput, Samples
from test_tcmatch import PARAMS, make_scene

SAMPLES = Path(__file__).resolve().parents[1] / "local" / "data"


def _write_csv(path: Path, names, rows) -> Path:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Time (s)", *names])
        writer.writerows(rows)
    return path


# --- TC file helpers ------------------------------------------------------------------------


def test_sheet_listing_suggestion_and_reading_one_sheet(tmp_path) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    log = book.active
    log.title = "Log"
    log.append(["Time", "Duct", "O2"])
    for k in range(50):
        log.append([k, 20.0 + k, 21.0])
    tc = book.create_sheet("TC")
    tc.append(["Time (s)", "S1", "S2"])
    for k in range(20):
        tc.append([k, 30.0 + k, 40.0 + 2 * k])
    book.create_sheet("Notes").append(["checked by", "staff"])
    path = tmp_path / "lab.xlsx"
    book.save(path)

    assert sheet_names(path) == ["Log", "TC", "Notes"]
    assert suggest_sheet(sheet_names(path)) == "TC"
    assert suggest_sheet(["TC", "Temperatures"]) is None  # two candidates: the user picks
    table = read_tc_table(path, sheet="TC")
    assert table.sheet == "TC" and table.names == ("S1", "S2") and table.times.size == 20
    assert read_tc_table(path).sheet == "TC"  # named like TC data, though smaller than Log
    assert read_tc_table(path, sheet="Log").sheet == "Log"
    with pytest.raises(ValueError, match="sheets: Log, TC, Notes"):
        read_tc_table(path, sheet="Missing")
    text = _write_csv(tmp_path / "tc.csv", ["A"], [[k, 20.0 + k] for k in range(5)])
    assert sheet_names(text) == ["tc"] and read_tc_table(text).sheet == ""


# --- fitted emissivities in the workbook -----------------------------------------------------


def test_prefill_checks_roi_emissivities() -> None:
    base = dict(times=(0.0, 1.0), columns=((1.0, 2.0),), names=("A",))
    TcPrefill(**base, roi_eps=(("spot", 0.8), ("box", 1.0))).check(10, 10)
    for bad in (0.0, 1.2, math.nan):
        with pytest.raises(ValueError, match="emissivity"):
            TcPrefill(**base, roi_eps=(("spot", bad),)).check(10, 10)


def test_roi_emissivity_fills_the_override_and_the_cached_values(tmp_path) -> None:
    from datetime import datetime as _dt

    import openpyxl
    from test_excel_workbook import T650, _source

    from flir_player.excel_export import ExportData, ExportOptions
    from flir_player.radiometry import object_temperature
    from flir_player.workbook import write_workbook

    rows = 12
    source = _source("Camera A", rows, np.random.default_rng(3))
    prefill = TcPrefill(times=tuple(float(t) for t in range(rows)), columns=((20.0,) * rows,), names=("TC A",),
                        roi_tc=(("TC1", "TC A"),), roi_eps=(("TC1", 0.8),))
    data = ExportData(times=np.arange(rows, dtype=float), sources=[source],
                      options=ExportOptions(bins=16, tc_prefill=prefill), created=_dt(2026, 10, 7),
                      versions={"tool": "test", "sdk": ""})
    path = tmp_path / "book.xlsx"
    write_workbook(path, data)
    settings = openpyxl.load_workbook(path, data_only=True)["Settings"]
    by_name = {settings.cell(r, 3).value: r for r in range(1, settings.max_row + 1)
               if settings.cell(r, 3).value in ("TC1", "TC2", "Box")}
    assert settings.cell(by_name["TC1"], 6).value == 0.8  # the override cell
    assert settings.cell(by_name["TC2"], 6).value is None
    assert settings.cell(by_name["TC1"], 9).value == pytest.approx(0.8)  # "Emissivity used" cache
    assert settings.cell(by_name["Box"], 9).value == pytest.approx(source.initial.emissivity)
    # the Data sheet's cached temperatures follow the override (viewers that do not recalculate)
    data_sheet = openpyxl.load_workbook(path, data_only=True)["Data"]
    spot = next(roi for roi in source.rois if roi.shape.name == "TC1")
    tau_params = source.initial.with_(emissivity=0.8)
    from flir_player.radiometry import transmission_used
    tau_params = tau_params.with_(transmission=transmission_used(T650, tau_params))
    expected = object_temperature(T650, tau_params, spot.counts["value"][0]) - 273.15
    cached = [data_sheet.cell(r, c).value for r in range(1, 12) for c in range(1, 12)]
    assert any(isinstance(v, float) and abs(v - expected) < 1e-6 for v in cached)


def test_prefill_from_a_result_carries_each_tcs_emissivity() -> None:
    samples, table = make_scene()
    result, _location, _series = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    prefill = tcmatch.tc_prefill(result, table)
    eps = dict(prefill.roi_eps)
    a, b, _c = result.channels
    assert eps == {"TC A spot": round(a.eps[1], 3), "TC A 1×1": round(a.eps[1], 3),
                   "TC B spot": round(b.eps[1], 3), "TC B 1×1": round(b.eps[1], 3)}
    assert all(v == pytest.approx(0.85, abs=0.01) for v in eps.values())
    prefill.check(100_000, 10)


# --- engine additions ------------------------------------------------------------------------


def test_a_given_frame_is_timed_like_the_statistics() -> None:
    samples, table = make_scene()
    by_seconds, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=137.0))
    by_frame, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_frame=137 * 30))
    assert by_frame.lag_fixed and by_frame.lag_s == by_seconds.lag_s == 137.0
    assert by_frame.ignition_frame == by_seconds.ignition_frame
    # on a superframing recording frames are timed by the clock: the inverse of frame_at
    clock = np.cumsum(np.r_[0.0, np.full(9, 0.5), 2.0, np.full(9, 0.5)])  # a dropped stretch at frame 10
    clocked = Samples(path=Path("s.ats"), times=np.arange(3.0), mean=np.zeros((3, 1, 1)), spread=np.zeros((3, 1, 1)),
                      counts=np.ones(3, np.int32), first_frame=np.arange(3), origin=(0, 0), size=(1, 1), fps=2.0,
                      step=1.0, report=None, params=PARAMS, frame_seconds=clock)
    for frame in (0, 5, 10, 19):
        assert clocked.seconds_at(frame) == clock[frame]
        assert clocked.frame_at(clocked.seconds_at(frame)) == frame
    assert clocked.seconds_at(21) == pytest.approx(clock[-1] + 1.0)


def test_run_analysis_returns_the_table_the_files_and_the_parameters(tmp_path, monkeypatch) -> None:
    samples, table = make_scene()
    samples.info["frames"] = int(samples.counts.sum())
    seen = {}

    def fake_sample(path, options, *, step_s, im, progress, abort, cache_dir, stage):
        seen.update(im=im, cache_dir=cache_dir)
        stage("Reading 1 frames")
        return samples

    monkeypatch.setattr(tcmatch, "read_tc_table", lambda path, sheet=None, time_column=None: table)
    monkeypatch.setattr(tcmatch, "sample_recording", fake_sample)
    stages = []
    parameters = {"emissivity": 0.95, "reflected_temp": 303.15, "atmosphere_temp": 293.15, "distance": 8.0,
                  "relative_humidity": 0.4, "est_atmospheric_transmission": 1.0, "ext_optics_temp": 293.15,
                  "ext_optics_transmission": 1.0, "can_change": True}
    handle = object()
    output = tcmatch.run_analysis("synthetic.seq", "tc.csv", MatchOptions(spot=1), out_dir=tmp_path / "out",
                                  cache_dir=tmp_path / "cache", im=handle, parameters=parameters,
                                  stage=stages.append)
    assert isinstance(output, RunOutput) and output.table is table
    assert seen == {"im": handle, "cache_dir": tmp_path / "cache"}
    assert output.result.reflected_c == pytest.approx(30.0)  # the player's reflected temperature
    assert output.parameters == parameters and output.parameters is not parameters
    assert set(output.files) == {"json", "summary", "rois", "figure"}
    assert all(path.exists() for path in output.files.values())
    assert stages[0] == "Reading the TC file" and "Saving the results" in stages
    assert any(text.startswith("Matching") for text in stages)


# --- dialogs ---------------------------------------------------------------------------------


def _metadata(tmp_path, *, presets=()) -> VideoMetadata:
    return VideoMetadata(path=tmp_path / "test.csq", width=640, height=480, num_frames=3000, start_time=None,
                         end_time=None, duration_seconds=100.0, nominal_fps=30.0, camera_model="FLIR T650sc",
                         presets=tuple(presets))


def _loaded(qapp, dialog) -> bool:
    return wait_until(qapp, lambda: dialog.table is not None or "Could not" in dialog.table_label.text(),
                      timeout=10)


def test_setup_dialog_reads_the_tc_file_and_builds_the_options(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog, results_folder

    tc = _write_csv(tmp_path / "tc.csv", ["S1", "S2", "TC Tree (1.5 m)"],
                    [[k, 20.0 + k, 25.0 + k, 30.0] for k in range(60)])
    box = RoiShape(4, "rect", ((100.2, 50.0), (160.0, 90.5)), "Box 1")
    rois = (RoiShape(1, "cursor", ((5.5, 5.5),), "Spot 1"), box)
    dialog = TcCalibrationDialog(_metadata(tmp_path), rois, current_frame=41, selected_roi_id=4)
    run = dialog.button_box.button(dialog.button_box.StandardButton.Ok)
    assert not run.isEnabled() and "TC file first" in dialog.estimate_label.text()
    assert dialog.out_edit.text() == str(results_folder(tmp_path / "test.csq"))
    assert dialog.region() == (100, 50, 160, 91)  # the selected box, every pixel it touches
    assert dialog.preset_combo is None  # one preset
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog) and dialog.table.names == ("S1", "S2", "TC Tree (1.5 m)")
    assert dialog.chosen_channels() == ["S1", "S2"]  # the air TC starts unticked
    assert run.isEnabled() and "frame statistics" in dialog.estimate_label.text()
    dialog.channel_table.item(1, 2).setText("0-40")
    dialog.time_combo.setCurrentIndex(dialog.time_combo.findData("frame"))
    dialog.current_button.click()
    options = dialog.options()
    assert options.roi == (100, 50, 160, 91) and options.channels == ("S1", "S2")
    assert options.valid == {"S2": ((0.0, 40.0),)} and options.lag_frame == 41 and options.preset is None
    params = dialog.parameters()
    assert params["tc_file"] == str(tc) and params["sheet"] is None and params["parameters_from"] == "player"
    assert params["cache_dir"] == str(Path(params["out_dir"]) / "cache")
    dialog._check_all(False)
    assert not run.isEnabled() and "at least one" in dialog.estimate_label.text()
    dialog.deleteLater()


def test_setup_dialog_guards_memory_windows_and_old_results(qapp, tmp_path, monkeypatch) -> None:
    from flir_player import tc_dialog
    from flir_player.tc_dialog import TcCalibrationDialog

    tc = _write_csv(tmp_path / "tc.csv", ["A"], [[k, 20.0 + k] for k in range(60)])
    metadata = _metadata(tmp_path, presets=(PresetRange(1, 100, 523.15, 873.15), PresetRange(2, 100, 773.15, 1473.15)))
    dialog = TcCalibrationDialog(metadata, (), current_frame=0)
    assert dialog.preset_combo is not None and dialog.preset_combo.currentData() == 1  # the lowest range
    assert "250 to 600 °C" in dialog.preset_combo.itemText(0)
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog)
    run = dialog.button_box.button(dialog.button_box.StandardButton.Ok)
    dialog.cache_check.setChecked(False)  # whole image, 101 bins: 248 MB fits in memory
    assert run.isEnabled()
    monkeypatch.setattr(tc_dialog, "MEMORY_LIMIT", 1e8)
    dialog._update()
    assert not run.isEnabled() and "too large" in dialog.estimate_label.text()
    dialog.cache_check.setChecked(True)
    assert run.isEnabled()
    warnings = []
    monkeypatch.setattr(tc_dialog.MessageDialog, "warning", lambda parent, title, text: warnings.append(text))
    dialog.channel_table.item(0, 2).setText("soon")
    dialog.accept()
    assert dialog.result() != QDialog.DialogCode.Accepted and "A:" in warnings[0]
    dialog.channel_table.item(0, 2).setText("")
    out = Path(dialog.out_edit.text())
    out.mkdir()
    (out / "summary.txt").write_text("old", encoding="utf-8")
    asked = []
    monkeypatch.setattr(tc_dialog, "confirm_replace", lambda parent, paths, confirmed: asked.append(paths))
    dialog.accept()  # declined (None): stays open
    assert asked and out / "summary.txt" in asked[0] and dialog.result() != QDialog.DialogCode.Accepted
    dialog.deleteLater()


def test_enter_only_finishes_typing_and_focus_changes_do_not_reread(qapp, tmp_path) -> None:
    from PySide6.QtTest import QTest

    from flir_player.tc_dialog import TcCalibrationDialog

    tc = _write_csv(tmp_path / "tc.csv", ["A"], [[k, 20.0 + k] for k in range(60)])
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0)
    dialog.show()
    QTest.mouseClick(dialog.file_edit, Qt.MouseButton.LeftButton)  # as a user would: the dialog has focus
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog)
    request = dialog._request
    dialog._file_edited()  # the path field loses focus to the Run button: no re-read, Run stays enabled
    assert dialog._request == request and dialog.table is not None
    QTest.keyClick(dialog.file_edit, Qt.Key.Key_Return)  # Enter in a field does not start the run
    table = dialog.channel_table
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=table.visualItemRect(table.item(0, 2)).center())
    editor = table.viewport().findChild(QLineEdit)  # one click opened the cell's editor
    assert editor is not None
    QTest.keyClicks(editor, "0-30")
    QTest.keyClick(editor, Qt.Key.Key_Return)
    qapp.processEvents()
    assert dialog.isVisible() and dialog.result() != QDialog.DialogCode.Accepted
    assert dialog.options().valid == {"A": ((0.0, 30.0),)}
    dialog.close()
    dialog.deleteLater()


def test_setup_dialog_starts_from_the_previous_run(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog

    tc = _write_csv(tmp_path / "tc.csv", ["A", "B"], [[k, 20.0 + k, 30.0 + k] for k in range(60)])
    previous = {"tc_file": str(tc), "sheet": None, "out_dir": str(tmp_path / "res"), "cache_dir": None,
                "parameters_from": "file", "channel_rows": [("A", False, ""), ("B", True, "5-50")],
                "options": MatchOptions(spot=5, tc_min_c=80.0, jump_k=40.0, jump_hold_s=0.0, lag_frame=99,
                                        below_range=True)}
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0, previous=previous, has_results=True)
    assert _loaded(qapp, dialog)
    assert dialog.chosen_channels() == ["B"] and dialog.options().valid == {"B": ((5.0, 50.0),)}
    options = dialog.options()
    assert (options.spot, options.tc_min_c, options.jump_k, options.jump_hold_s, options.lag_frame,
            options.below_range) == (5, 80.0, 40.0, 0.0, 99, True)
    assert dialog.params_combo.currentData() == "file" and not dialog.cache_check.isChecked()
    assert dialog.out_edit.text() == str(tmp_path / "res")
    last = [b for b in dialog.button_box.buttons() if b.text() == "Show Last Results"]
    assert last
    last[0].click()
    assert dialog.result() == TcCalibrationDialog.LAST_RESULTS
    dialog.deleteLater()


def _output(**overrides) -> RunOutput:
    samples, table = make_scene()
    result, _location, _series = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    for key, value in overrides.items():
        setattr(result, key, value)
    return RunOutput(result=result, table=table)


def test_results_dialog_lists_the_tcs_and_offers_their_values(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcResultsDialog

    output = _output()
    dialog = TcResultsDialog(output, tmp_path)
    grid = dialog.tc_table
    assert [grid.item(r, 0).text() for r in range(grid.rowCount())] == ["TC A", "TC B", "TC C"]
    assert grid.item(0, 1).text() == "9, 12" and grid.item(0, 3).text() == "0.85"
    assert grid.item(2, 1).text() == "-" and "failed" in grid.item(2, 6).text()
    items = [dialog.eps_combo.itemText(k) for k in range(dialog.eps_combo.count())]
    assert items[0].startswith("All TCs agree: 0.85") and any(t.startswith("TC B: 0.85") for t in items)
    values, added = [], []
    dialog.emissivity_requested.connect(values.append)
    dialog.add_rois_requested.connect(lambda: added.append(True))
    dialog.eps_button.click()
    dialog.rois_button.click()
    assert values == [pytest.approx(output.result.common_eps)] and added == [True]
    assert not dialog.rois_button.isEnabled()  # once
    assert dialog.workbook_button.isEnabled() and dialog.folder_button.isEnabled()
    dialog.workbook_button.click()
    assert dialog.result() == TcResultsDialog.WORKBOOK
    dialog.deleteLater()


def test_results_dialog_without_pixels_or_with_presets(qapp) -> None:
    from flir_player.tc_dialog import TcResultsDialog

    output = _output()
    for ch in output.result.channels:
        ch.match.found = False
        ch.eps = None
    dialog = TcResultsDialog(output, None)
    assert not dialog.rois_button.isEnabled() and not dialog.workbook_button.isEnabled()
    assert not dialog.eps_button.isEnabled() and not dialog.folder_button.isEnabled()
    dialog.deleteLater()
    dialog = TcResultsDialog(_output(preset=1), None)
    assert dialog.rois_button.isEnabled() and not dialog.workbook_button.isEnabled()
    assert "superframing" in dialog.workbook_button.toolTip()
    dialog.deleteLater()


# --- the whole flow on a real recording --------------------------------------------------------


def _synthetic_tc(path: Path, recording: Path, pixel, lag_s: float) -> Path:
    """A TC that reads the recording's own SDK temperature at ``pixel`` (1 s means), logger 0 = ``lag_s``."""
    import flir_player  # noqa: F401  (FileSDK DLL preload)
    import fnv
    import fnv.file

    from flir_player.calibration import read_array, set_unit_safely

    im = fnv.file.ImagerFile(str(recording))
    try:
        set_unit_safely(im, fnv.Unit.TEMPERATURE_FACTORY, fnv.TempType.CELSIUS)
        fps = 30.0
        n = int(im.num_frames)
        values = np.array([read_array(im, i)[pixel] for i in range(n)], dtype=np.float64)
    finally:
        im.close()
    bins = np.floor(np.arange(n) / fps + 0.5).astype(int)
    per_second = np.array([values[bins == k].mean() for k in range(bins.max() + 1)])
    start = int(lag_s)
    rows = [[k, float(per_second[start + k])] for k in range(per_second.size - start - 1)]
    return _write_csv(path, ["Probe TC", "Air (tree)"], [row + [25.0] for row in rows])


def test_the_whole_flow_through_the_player(qapp, tmp_path, monkeypatch) -> None:
    from flir_player import main_window as mw
    from flir_player.main_window import MainWindow
    from flir_player.tc_dialog import TcCalibrationDialog, TcResultsDialog

    recording = SAMPLES / "2.seq"
    tc = _synthetic_tc(tmp_path / "probe.csv", recording, (256, 52), lag_s=10)
    window = MainWindow()
    try:
        window.open_path(recording)
        assert wait_until(qapp, lambda: window.current_packet is not None and not window._busy, timeout=20)
        window._add_roi("rect", ((30.0, 236.0), (80.0, 276.0)))
        box_id = window._selected_roi_id
        built = {}

        def fake_exec(self):
            self.file_edit.setText(str(tc))
            self._load(sheet=None)
            assert _loaded(qapp, self)
            self.out_edit.setText(str(tmp_path / "results"))
            built["params"] = self.parameters()
            return QDialog.DialogCode.Accepted

        shown = []

        def fake_results(self):
            shown.append(self.output)
            return TcResultsDialog.WORKBOOK

        monkeypatch.setattr(TcCalibrationDialog, "exec", fake_exec)
        monkeypatch.setattr(TcResultsDialog, "exec", fake_results)
        monkeypatch.setattr(mw.QFileDialog, "getSaveFileName",
                            staticmethod(lambda *a, **k: (str(tmp_path / "book.xlsx"), "")))
        stages = []
        window.decoder.export_stage.connect(stages.append)
        window._open_tc_dialog()
        assert built["params"]["options"].roi == (30, 236, 80, 276)
        assert built["params"]["options"].channels == ("Probe TC",)  # the air TC starts unticked
        assert wait_until(qapp, lambda: shown, timeout=120)
        result = shown[0].result
        match = result.channels[0].match
        assert abs(result.lag_s - 10.0) <= 1.0 and match.found
        assert abs(match.pixel[0] - 256) <= 2 and abs(match.pixel[1] - 52) <= 3
        assert any(text.startswith("Reading") for text in stages)
        assert (tmp_path / "results" / "summary.txt").exists()
        assert (tmp_path / "results" / "cache").is_dir()
        # the workbook job (the results dialog asked for it) lends the open handle again
        assert wait_until(qapp, lambda: (tmp_path / "book.xlsx").exists() and not window._busy
                          and window._export_progress is None, timeout=120)
        # the player gets the ROIs and the ignition frame; a second call replaces them
        before = len(window._rois)
        window._add_tc_rois(shown[0])
        window._add_tc_rois(shown[0])
        names = [roi.name for roi in window._rois]
        assert len(window._rois) == before + 1 and names.count("Probe TC 3×3") == 1
        assert box_id in [roi.id for roi in window._rois]
        assert window._ignition_frame() == tcmatch.workbook_origin(result)[0]
        # the player reads frames as before once the borrowed handle is back
        unit = window.current_packet.unit.key
        window.seek_to(100)
        assert wait_until(qapp, lambda: window.current_packet.index == 100 and not window._busy, timeout=20)
        assert window.current_packet.unit.key == unit
    finally:
        window.close()


# --- Codex review 2026-10-07 (local/notes/2026-10-07-tc-dialog/codex/, run review-tc-dialog) -------


def _renamed(result, names):
    """The scene's two found TCs repeated under ``names``, at distinct pixels."""
    import dataclasses

    a, b, _c = result.channels
    channels = []
    for k, name in enumerate(names):
        source = (a, b)[k % 2]
        eps = (0.5 + 0.04 * k,) * 3
        channels.append(dataclasses.replace(source, name=name, eps=eps, eps_events=1))
    return dataclasses.replace(result, channels=channels)


def test_c1_rois_carry_their_own_tc_even_when_names_collide() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    names = ("TC (front)", "TC (back)", "TC 1", "TC 1 wall")
    result = _renamed(result, names)
    table = TcTable(times=table.times, values=np.column_stack([table.values[:, 0]] * 4), names=names)
    pairs = tcmatch.tc_rois(result)
    assert [channel for _shape, channel in pairs] == [n for n in names for _ in (0, 1)]
    labels = [shape.name for shape, _channel in pairs]
    assert len(set(labels)) == len(labels)
    assert labels[:2] == ["TC (front) spot", "TC (front) 1×1"]  # shared short name "TC": full names
    assert labels[4:] == ["TC 1 spot", "TC 1 1×1", "TC 1 wall spot", "TC 1 wall 1×1"]
    prefill = tcmatch.tc_prefill(result, table)
    eps = dict(prefill.roi_eps)
    tc = dict(prefill.roi_tc)
    for shape, channel in pairs:
        assert tc[shape.name] == channel
        assert eps[shape.name] == round(next(ch.eps[1] for ch in result.channels if ch.name == channel), 3)


def test_c5_every_tc_gets_its_emissivity_beyond_the_ten_compare_columns() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    names = tuple(f"T{k + 1}" for k in range(11))
    result = _renamed(result, names)
    table = TcTable(times=table.times, values=np.column_stack([table.values[:, 0]] * 11), names=names)
    prefill = tcmatch.tc_prefill(result, table)
    assert len(prefill.names) == 10 and "T11 spot" not in dict(prefill.roi_tc)
    assert dict(prefill.roi_eps)["T11 spot"] == round(result.channels[10].eps[1], 3)
    assert len(prefill.roi_eps) == 22


def test_c2_a_given_frame_between_bins_stays_that_frame() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_frame=4121))
    assert result.lag_fixed and result.ignition_frame == 4121
    assert result.lag_s == pytest.approx(4121 / 30.0)
    frame, offset = tcmatch.workbook_origin(result)
    assert frame == 4121 and offset == pytest.approx(0.0, abs=1e-9)


def test_c6_a_job_dropped_from_the_queue_still_reports(qapp) -> None:
    from flir_player.decoder import DecoderThread

    decoder = DecoderThread()  # not started: the jobs wait in its queue
    seen = []
    decoder.tc_finished.connect(lambda ok, message, output: seen.append(("tc", ok, message, output)))
    decoder.export_finished.connect(lambda ok, message: seen.append(("export", ok, message)))
    decoder.request_tc_match({"recording": "a.seq"})
    decoder.request_tc_workbook({"recording": "a.seq"})
    decoder.request_open("b.seq")
    assert seen == [("tc", False, "Cancelled", None), ("export", False, "Export cancelled")]
    assert [command for command, _payload in decoder._commands.queue] == ["open"]


def _fake_run(monkeypatch, samples, table):
    def fake_sample(path, options, *, step_s, im, progress, abort, cache_dir, stage):
        stage("Reading 1 frames")
        return samples

    monkeypatch.setattr(tcmatch, "read_tc_table", lambda path, sheet=None, time_column=None: table)
    monkeypatch.setattr(tcmatch, "sample_recording", fake_sample)


def test_c7_cancelling_late_is_a_cancel_and_publishes_nothing(tmp_path, monkeypatch) -> None:
    from flir_player.jobs import JobCancelled

    samples, table = make_scene()
    samples.info["frames"] = int(samples.counts.sum())
    _fake_run(monkeypatch, samples, table)
    with pytest.raises(JobCancelled):  # cancelled during the calibration check: not "no calibration"
        tcmatch.run_analysis("synthetic.seq", "tc.csv", MatchOptions(spot=1), abort=lambda: True)
    state = {"saving": False}

    def at_saving(text):
        state["saving"] = state["saving"] or text == "Saving the results"

    out = tmp_path / "out"
    with pytest.raises(JobCancelled):
        tcmatch.run_analysis("synthetic.seq", "tc.csv", MatchOptions(spot=1), out_dir=out, stage=at_saving,
                             abort=lambda: state["saving"])
    assert not (out / "summary.txt").exists() and not (out / "result.json").exists()


def test_c8_results_replace_only_what_the_user_confirmed(tmp_path, monkeypatch) -> None:
    samples, table = make_scene()
    samples.info["frames"] = int(samples.counts.sum())
    _fake_run(monkeypatch, samples, table)
    out = tmp_path / "out"
    out.mkdir()
    (out / "summary.txt").write_text("someone else's", encoding="utf-8")
    with pytest.raises((OSError, ValueError)):
        tcmatch.run_analysis("synthetic.seq", "tc.csv", MatchOptions(spot=1), out_dir=out, replace=[])
    assert (out / "summary.txt").read_text(encoding="utf-8") == "someone else's"
    assert not (out / "result.json").exists()  # all or nothing
    output = tcmatch.run_analysis("synthetic.seq", "tc.csv", MatchOptions(spot=1), out_dir=out,
                                  replace=[out / "summary.txt"])
    assert (out / "summary.txt").read_text(encoding="utf-8").startswith("Recording:")
    assert all(path.exists() for path in output.files.values())


def test_c9_a_read_for_an_earlier_path_does_not_land(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog

    tc = _write_csv(tmp_path / "tc.csv", ["A"], [[k, 20.0 + k] for k in range(60)])
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0)
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    dialog.file_edit.setText(str(tmp_path / "missing.csv"))
    dialog._file_edited()
    wait_until(qapp, lambda: False, timeout=1.0)  # let the first read finish
    run = dialog.button_box.button(dialog.button_box.StandardButton.Ok)
    assert dialog.table is None and not run.isEnabled() and "No such file" in dialog.table_label.text()
    dialog.deleteLater()


def test_c10_the_size_estimate_uses_the_engines_bin_width(qapp, tmp_path) -> None:
    import dataclasses

    from flir_player.tc_dialog import TcCalibrationDialog

    tc = _write_csv(tmp_path / "tc.csv", ["A"], [[k / 100.0, 20.0 + k] for k in range(600)])  # 0.01 s logger
    box = RoiShape(1, "rect", ((0.0, 0.0), (200.0, 200.0)), "Box 1")
    metadata = VideoMetadata(path=tmp_path / "r.csq", width=640, height=480, num_frames=3000, start_time=None,
                             end_time=None, duration_seconds=100.0, nominal_fps=30.0)
    dialog = TcCalibrationDialog(metadata, (box,), current_frame=0, selected_roi_id=1)
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog) and dialog.table.interval == pytest.approx(0.01)
    # bins follow the frames (1/30 s), not the 0.01 s logger: 3,001 bins of 200 × 200 pixels
    assert dialog._statistics_bytes() == pytest.approx(8.0 * (100.0 * 30 + 1) * 200 * 200)
    presets = (PresetRange(1, 1500, 523.15, 873.15), PresetRange(2, 1500, 773.15, 1473.15))
    dialog.metadata = dataclasses.replace(metadata, presets=presets)
    assert dialog._statistics_bytes() == pytest.approx(8.0 * (100.0 * 15 + 1) * 200 * 200)  # one preset: 15 Hz
    dialog.deleteLater()
