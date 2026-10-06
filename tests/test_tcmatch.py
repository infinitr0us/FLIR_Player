# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""TC calibration engine: logger tables, the offset/pixel search, events and emissivity fits.

Synthetic scenes use the A700's factory calibration (0922 battery test) and
known offsets, pixels and emissivities, so the engine must recover them. The
0922 regression at the end needs the full recording and runs only with
``FLIR_TCMATCH_0922=1``.
"""
from __future__ import annotations

import json
import math
import os
from datetime import datetime, time as dtime, timedelta
from pathlib import Path

import numpy as np
import pytest

from flir_player import tcmatch
from flir_player.calibration import CalibrationReport
from flir_player.geometry import roi_coordinates
from flir_player.radiometry import Calibration, MeasurementParameters, PlanckConstants, RangeLimits, coefficients
from flir_player.tcdata import (
    TcTable,
    failure_flags,
    parse_rows,
    parse_windows,
    read_tc_table,
    resample,
    window_mask,
)
from flir_player.tcmatch import MatchOptions, Samples

SAMPLES = Path(__file__).resolve().parents[1] / "local" / "data"
CLIP = "a700_clip.seq"
A700 = Calibration(planck=PlanckConstants(R=176682.01169644916, B=1488.0, F=1.5, O=-5772.0),
                   limits=RangeLimits(calibrated=(273.15, 923.15), raw=(5936.0, 58611.0)), source="A700 FFF")
REPORT = CalibrationReport(calibration=A700, status="verified", message="synthetic A700")
PARAMS = MeasurementParameters(emissivity=0.95, reflected_k=293.15, atmosphere_k=293.15, transmission=1.0)


def counts_for(temp_c, eps, params=PARAMS):
    """Raw counts a surface at ``temp_c`` with emissivity ``eps`` gives."""
    k1, k2 = coefficients(A700, params.with_(emissivity=eps))
    return (A700.planck.signal(np.asarray(temp_c, dtype=np.float64) + 273.15) + k2) / k1


# --- logger tables ---------------------------------------------------------------------------


def test_parse_rows_finds_the_0922_layout() -> None:
    rows = [("ELAPSED TIME", "BatteryCell#3", "BatteryCell#6", "BatteryCell#9"),
            ("[sec]", "T1", "T2", "T3"),
            (2, None, None, None),
            ("shft time", None, None, None)]
    rows += [(k, 12.0 + 0.5 * k, 11.5, 12.1) for k in range(10)]
    rows += [(None, None, None, None), ("end", None, None, None)]
    table = parse_rows(rows, source="x.xlsx")
    assert table.names == ("T1 (BatteryCell#3)", "T2 (BatteryCell#6)", "T3 (BatteryCell#9)")
    assert table.time_label == "ELAPSED TIME"
    np.testing.assert_array_equal(table.times, np.arange(10.0))
    assert table.values.shape == (10, 3) and table.values[4, 0] == 14.0
    assert table.interval == 1.0
    assert any("shft time" in note for note in table.notes)  # reported, not silently dropped
    assert table.index("T2") == 1 and table.index(3) == 2 and table.index("t3 (batterycell#9)") == 2
    with pytest.raises(KeyError):
        table.index("T")  # three channels start with T


def test_parse_rows_dates_and_times_of_day() -> None:
    start = datetime(2026, 9, 22, 23, 59, 58)
    stamped = parse_rows([("Date/Time", "TC A")] + [(start + timedelta(seconds=2 * k), 20.0 + k) for k in range(5)])
    np.testing.assert_allclose(stamped.times, [0, 2, 4, 6, 8])
    assert stamped.interval == 2.0
    clock = parse_rows([("Time", "A")] + [(dtime(23, 59, 59), 20.0), (dtime(0, 0, 0), 21.5), (dtime(0, 0, 1), 23.0)])
    np.testing.assert_allclose(clock.times, [0, 1, 2])  # unwrapped over midnight


def test_parse_rows_prefers_a_time_column_over_a_scan_number() -> None:
    rows = [("Scan", "Elapsed (s)", "TC1")] + [(k + 1, 0.5 * k, 30.0) for k in range(6)]
    table = parse_rows(rows)
    np.testing.assert_allclose(table.times, 0.5 * np.arange(6))
    assert table.names == ("TC1",)  # a scan counter is not a thermocouple
    assert any("Scan" in note for note in table.notes)


def test_read_tc_table_csv_and_xlsx(tmp_path) -> None:
    csv_path = tmp_path / "log.csv"
    csv_path.write_text("time;TC 1;TC 2\n[s];[°C];[°C]\n0;20;21\n1;22;23\n2;24;25\n3;26;27\n", encoding="utf-8")
    table = read_tc_table(csv_path)
    assert table.names == ("TC 1", "TC 2")
    np.testing.assert_allclose(table.values[:, 1], [21, 23, 25, 27])

    import xlsxwriter
    xlsx_path = tmp_path / "log.xlsx"
    book = xlsxwriter.Workbook(str(xlsx_path))
    sheet = book.add_worksheet("Notes")
    sheet.write(0, 0, "nothing here")
    data = book.add_worksheet("Logger")
    data.write_row(0, 0, ["ELAPSED TIME", "Cell#3"])
    for k in range(5):
        data.write_row(1 + k, 0, [k, 100.0 + 2 * k])
    book.close()
    table = read_tc_table(xlsx_path)
    assert table.names == ("Cell#3",) and table.values[-1, 0] == 108.0
    assert "Logger" in table.source
    with pytest.raises(ValueError):
        read_tc_table(xlsx_path, sheet="Notes")


def test_failure_flags_catch_a_broken_tc_but_not_a_quench() -> None:
    t = np.arange(400.0)
    good = np.where(t < 100, 12.0, np.minimum(12 + 3 * (t - 100), 500.0))
    good[350:] = 18.0  # pushed into water: a real drop, still above the starting level
    broken = np.full_like(t, 12.0)
    broken[67:] = np.where(np.arange(333) % 7 < 4, -300.0, 40.0)  # shorted from 67 s
    spiky = good.copy()
    spiky[200] += 400.0
    table = TcTable(times=t, values=np.column_stack([good, broken, spiky]), names=("good", "broken", "spiky"))
    ok = failure_flags(table, 0)
    assert ok.failed_from is None and not ok.flagged.any() and ok.reasons == ()
    bad = failure_flags(table, 1)
    assert bad.failed_from == 67.0 and bad.flagged[67:].all() and not bad.flagged[:67].any()
    assert "failed from 67 s" in bad.reasons[0]
    spikes = failure_flags(table, 2)
    assert spikes.failed_from is None and spikes.spikes == 1 and spikes.flagged.sum() == 1 and spikes.flagged[200]


def test_windows_and_resampling() -> None:
    assert parse_windows("800-1150, 1300-") == ((800.0, 1150.0), (1300.0, math.inf))
    assert parse_windows("-5..10") == ((-5.0, 10.0),)
    with pytest.raises(ValueError):
        parse_windows("1150-800")
    mask = window_mask(np.array([0.0, 900, 1200, 1400]), parse_windows("800-1150,1300-"))
    assert mask.tolist() == [False, True, False, True]
    table = TcTable(times=np.array([0.0, 1, 2, 10, 11]), values=np.array([[0.0], [1], [2], [10], [11]]), names=("a",))
    out = resample(table, 0, np.array([0.5, 1.0, 5.0, 10.5, 20.0]))
    np.testing.assert_allclose(out, [0.5, 1.0, np.nan, 10.5, np.nan])  # no bridge over the 8 s gap


# --- the offset scan ---------------------------------------------------------------------------


def test_lag_scan_matches_a_brute_force_pearson() -> None:
    rng = np.random.default_rng(3)
    n_ir, n_tc, P, C = 90, 50, 7, 2
    X = rng.normal(size=(n_ir, P)).cumsum(axis=0)
    present = rng.random(n_ir) > 0.15
    Y = rng.normal(size=(n_tc, C)).cumsum(axis=0)
    usable = rng.random((n_tc, C)) > 0.2
    lags = np.arange(-45, 80)
    best, where = tcmatch.lag_scan(np.where(present[:, None], X, 0.0), present, Y, usable, lags, min_samples=20,
                                   chunk=3)
    for i, L in enumerate(lags):
        for c in range(C):
            values = []
            for p in range(P):
                j = np.arange(n_tc)
                k = L + j
                ok = usable[:, c] & (k >= 0) & (k < n_ir)
                ok[ok] &= present[k[ok]]
                if ok.sum() < 20:
                    values.append(-np.inf)
                    continue
                values.append(np.corrcoef(X[k[ok], p], Y[ok, c])[0, 1])
            values = np.array(values)
            if np.isfinite(values).any():
                assert best[i, c] == pytest.approx(values.max(), abs=1e-9)
                assert where[i, c] == int(np.argmax(values))
            else:
                assert np.isnan(best[i, c])


# --- synthetic scenes ------------------------------------------------------------------------


def _rise(tau, start, amplitude, tau_c, bump_at=None, bump=0.0, cool_at=None):
    x = np.where(tau < start, 0.0, amplitude * (1 - np.exp(-np.maximum(tau - start, 0) / tau_c)))
    if bump_at is not None:
        x = x + bump * np.exp(-((tau - bump_at) / 12.0) ** 2) * (tau > start)
    if cool_at is not None:
        peak = amplitude * (1 - np.exp(-(cool_at - start) / tau_c))
        x = np.where(tau > cool_at, peak * np.exp(-(tau - cool_at) / 70.0), x)
    return x


SPOT_A = (9, 12)  # (row, col) in the image
SPOT_B = (14, 27)


def make_scene(*, lag=137, eps=0.85, n_ir=760, n_log=480, origin=(3, 2), shape=(22, 34), flames=(),
               low_after=None, frames_per_bin=30):
    """(Samples, TcTable): two TC spots with conduction halos, a broken third TC.

    TC A heats from logger 60 s with a bump at 180 s; TC B from 150 s with a
    bump at 260 s. Neighbours of a spot see its curve delayed by 4 s per pixel
    of distance and weaker, so the spot itself correlates best. ``flames``
    lists logger windows of strong flicker in front of spot B.
    """
    rng = np.random.default_rng(7)
    t_ir = np.arange(n_ir, dtype=np.float64)
    bg = 15.0 + 0.002 * t_ir + 0.3 * np.sin(t_ir / 37.0)
    curves = {"A": lambda tau: _rise(tau, 60, 240, 45, 180, 35, 330),
              "B": lambda tau: _rise(tau, 150, 180, 35, 260, 30, 390)}
    h, w = shape
    x0, y0 = origin
    temps = bg[:, None, None] + rng.normal(0, 0.05, (n_ir, h, w))
    for (row, col), key in ((SPOT_A, "A"), (SPOT_B, "B")):
        for dr in range(-4, 5):
            for dc in range(-4, 5):
                d = math.hypot(dr, dc)
                weight = max(0.0, 1 - d / 5)
                if weight <= 0:
                    continue
                temps[:, row - y0 + dr, col - x0 + dc] += weight * curves[key](t_ir - lag - 4 * d)
    counts = counts_for(temps, eps).astype(np.float32)
    spread = np.full_like(counts, 2.0)  # sensor noise in counts
    rb, cb = SPOT_B[0] - y0, SPOT_B[1] - x0
    for start, end in flames:
        sel = slice(int(start + lag), int(end + lag) + 1)
        spread[sel, rb - 3:rb + 4, cb - 3:cb + 4] = 800.0
    samples = Samples(path=Path("synthetic.seq"), times=t_ir, mean=counts, spread=spread,
                      counts=np.full(n_ir, frames_per_bin, dtype=np.int32),
                      first_frame=np.arange(n_ir, dtype=np.int64) * 30, origin=origin, size=(48, 40), fps=30.0,
                      step=1.0, report=REPORT, params=PARAMS, info={"camera": "Synthetic", "region": [0, 0, 1, 1]})
    tau = np.arange(n_log, dtype=np.float64)
    t_rec = (tau + lag).astype(int)
    a = bg[t_rec] + curves["A"](tau)
    b = bg[t_rec] + curves["B"](tau)
    if low_after is not None:
        b = np.where(tau > low_after, 0.55 * b, b)  # TC lifts off: reads low, keeps the shape
    broken = bg[t_rec].copy()
    broken[40:] = -400.0
    table = TcTable(times=tau, values=np.column_stack([a, b, broken]), names=("TC A", "TC B", "TC C"),
                    source="synthetic")
    return samples, table


def test_synthetic_scene_recovers_offset_pixels_and_emissivity() -> None:
    samples, table = make_scene()
    options = MatchOptions(spot=1, tc_min_c=50.0)
    result, location, series = tcmatch.analyse(samples, table, options)
    assert result.lag_s == 137.0 and not result.lag_fixed
    assert result.ignition_frame == 137 * 30
    a, b, c = result.channels
    assert a.match.pixel == SPOT_A and b.match.pixel == SPOT_B
    assert a.match.r > 0.999 and SPOT_A in a.match.candidates
    assert not c.match.found and c.flags and "failed from 40 s" in c.flags[0]
    for ch in (a, b):
        assert ch.events and all(e.physical for e in ch.events)
        for e in ch.events:
            assert e.fit.eps == pytest.approx(0.85, abs=0.002)
            assert e.fit.eps_ls == pytest.approx(0.85, abs=0.002)
            assert e.fit.rmse_k < 1.0 and e.excluded_bound == 0
    assert result.pooled.eps == pytest.approx(0.85, abs=0.002)
    assert result.consistent and result.spread < 0.01
    assert result.cross and all(x["rmse_k"] < 1.5 for x in result.cross)
    assert any("TC C" in w for w in result.warnings)  # not searched for
    json.dumps(result.to_dict())  # serialisable
    text = tcmatch.summary_text(result)
    assert "logger time 0 = recording time 137 s" in text and "agree" in text


def test_partial_overlap_and_a_fixed_offset() -> None:
    samples, table = make_scene(lag=-40, n_ir=500, n_log=480)  # the logger started 40 s before the camera
    result, _location, _series = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    assert result.lag_s == -40.0
    assert result.channels[0].match.pixel == SPOT_A
    fixed, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=-40.0,
                                                                 pixels={"TC B": (SPOT_B[0], SPOT_B[1] + 1)}))
    assert fixed.lag_fixed and fixed.lag_s == -40.0
    assert fixed.channels[1].match.fixed and fixed.channels[1].match.pixel == (SPOT_B[0], SPOT_B[1] + 1)


def test_events_skip_flames_and_respect_validity_windows() -> None:
    samples, table = make_scene(flames=[(300, 330)], low_after=340)
    # TC B reads low after 340 s; its shape still fixes the offset and the pixel
    options = MatchOptions(spot=1, valid={"TC B": ((0, 340),)})
    result, _location, series = tcmatch.analyse(samples, table, options)
    b = result.channels[1]
    assert result.lag_s == 137.0 and b.match.pixel == SPOT_B
    for e in b.events:
        assert e.end_s <= 340 and not (e.start_s <= 315 <= e.end_s)  # no flames, no low readings
        assert e.fit.eps == pytest.approx(0.85, abs=0.003)
    rules = tcmatch.event_rules(series["TC B"]["channel"], series["TC B"]["spot"], options, True)
    flicker = rules["no flames (flicker)"]
    assert not flicker[300:331].any() and flicker[200:290].all()
    # without the window (offset and pixels given), the low readings are "IR above any match"
    loose, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=137.0,
                                                                 pixels={"TC A": SPOT_A, "TC B": SPOT_B}))
    late = [e for e in loose.channels[1].events if e.end_s > 345]
    assert late and all(not e.physical for e in late)
    assert any("not used for the emissivity" in w for w in loose.warnings)
    assert all(e.fit.eps == pytest.approx(0.85, abs=0.003) for e in loose.channels[1].events if e.physical)
    # and the search, given the low readings, is pulled off: the reason the windows apply to it by default
    pulled, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, search_outside_windows=True,
                                                                  valid={"TC B": ((0, 340),)}))
    assert pulled.lag_s != 137.0


def test_low_frame_rate_disables_the_flame_filter_with_a_warning() -> None:
    samples, table = make_scene(frames_per_bin=1)
    samples.spread[:] = np.nan
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    assert any("flames" in w for w in result.warnings)
    assert result.channels[0].events and result.channels[0].events[0].fit.eps == pytest.approx(0.85, abs=0.002)


def test_spot_block_and_roi_shapes_cover_the_same_pixels() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=3))
    shapes = tcmatch.roi_shapes(result)
    assert [s.kind for s in shapes] == ["cursor", "rect", "cursor", "rect"]
    spot, box = shapes[0], shapes[1]
    ys, xs = roi_coordinates(spot, 40, 48)
    assert (int(ys[0]), int(xs[0])) == SPOT_A
    ys, xs = roi_coordinates(box, 40, 48)
    r0, r1, c0, c1 = tcmatch.block_bounds(SPOT_A, 3, samples)
    oy, ox = samples.origin[1], samples.origin[0]
    assert sorted(zip(ys.tolist(), xs.tolist())) == [(r + oy, c + ox) for r in range(r0, r1) for c in range(c0, c1)]
    # a 3 × 3 block averages the halo: a lower apparent temperature, so the fitted ε drops below the truth
    e = result.channels[0].events[0]
    assert e.fit.eps < 0.85 and e.pixel_eps[0] <= e.fit.eps <= e.pixel_eps[2]


def test_fit_eps_bound_and_unconstrained() -> None:
    tc = np.linspace(100, 400, 50)
    counts = counts_for(tc, 0.7)
    fit = tcmatch.fit_eps(A700, PARAMS, counts, tc)
    assert fit.eps == pytest.approx(0.7, abs=1e-3) and fit.eps_ls == pytest.approx(0.7, abs=1e-6)
    assert not fit.at_bound and fit.matching_median == pytest.approx(0.7, abs=1e-6)
    hot = tcmatch.fit_eps(A700, PARAMS, counts_for(tc + 60, 1.0), tc)  # IR above any match
    assert hot.at_bound and hot.eps == 1.0 and hot.eps_ls > 1.0
    empty = tcmatch.fit_eps(A700, PARAMS, np.array([]), np.array([]))
    assert empty.n == 0 and math.isnan(empty.eps)


# --- sampling ---------------------------------------------------------------------------------


class _FakeImager:
    def __init__(self, frames: np.ndarray, presets=None):
        self.frames = frames
        self.num_frames, self.height, self.width = frames.shape
        self._i = 0
        self._presets = presets

    def get_frame(self, index):
        self._i = int(index)

    @property
    def final(self):
        return self.frames[self._i].ravel()

    @property
    def frame_info(self):
        return [{"name": "Preset", "value": str(self._presets[self._i])}] if self._presets is not None else []


def test_accumulate_mean_and_detrended_spread(tmp_path) -> None:
    rng = np.random.default_rng(5)
    fps, n = 10.0, 47
    t = np.arange(n) / fps
    frames = (1000 + 30 * t[:, None, None] + rng.normal(0, 2, (n, 4, 5))).astype(np.float32)
    im = _FakeImager(frames)
    seen = []
    times, mean, spread, counts, first = tcmatch._accumulate(
        im, np.arange(n), fps, 1.0, (1, 0, 4, 3), lambda d, t_: seen.append((d, t_)), lambda: False, None, None)
    bins = np.floor(t + 0.5).astype(int)
    assert times.tolist() == [0.0, 1.0, 2.0, 3.0, 4.0, 5.0] and counts.tolist() == np.bincount(bins).tolist()
    assert first.tolist() == [int(np.flatnonzero(bins == k)[0]) for k in range(6)]
    for k in range(6):
        block = frames[bins == k][:, 0:3, 1:4].astype(np.float64)
        np.testing.assert_allclose(mean[k], block.mean(axis=0), rtol=1e-6)
        tau = t[bins == k]
        if tau.size < 3:  # no spread from two frames
            assert np.isnan(spread[k]).all()
            continue
        for r in range(3):
            for c in range(3):
                y = block[:, r, c]
                resid = y - np.polyval(np.polyfit(tau, y, 1), tau)
                assert spread[k, r, c] == pytest.approx(resid.std(), rel=1e-4, abs=1e-4)
    assert seen[-1] == (n, n)
    # the cache keeps the same statistics
    cached = tcmatch._accumulate(im, np.arange(n), fps, 1.0, (1, 0, 4, 3), None, lambda: False, tmp_path, "k")
    loaded = tcmatch._load_cache(tmp_path, "k")
    for got, want in zip(loaded, cached):
        np.testing.assert_array_equal(np.asarray(got), np.asarray(want))
    assert tcmatch._load_cache(tmp_path, "other") is None


def test_accumulate_refuses_huge_statistics_without_a_cache() -> None:
    im = _FakeImager(np.zeros((2, 2, 2), np.float32))
    with pytest.raises(ValueError, match="smaller region"):
        tcmatch._accumulate(im, np.array([0, 10 ** 8]), 1.0, 1.0, (0, 0, 2, 2), None, lambda: False, None, None)


class _Preset:
    def __init__(self, available, frames, low=0.0, high=0.0):
        self.available, self.num_frames, self.calibrated = available, frames, high > 0
        self.min_temp, self.max_temp = low, high


class _Info:
    def __init__(self, presets):
        self.preset_info = presets


def test_preset_frames_of_a_superframing_recording() -> None:
    frames = np.zeros((200, 2, 2), np.float32)
    im = _FakeImager(frames, presets=[1 + (i % 2) for i in range(200)])
    im.source_info = _Info([_Preset(False, 0), _Preset(True, 100, 523.15, 873.15), _Preset(True, 100, 773.15, 1473.15)])
    layout, ranges = tcmatch.preset_frames(im)
    assert layout == {1: (0, 2), 2: (1, 2)}
    assert ranges == {1: (523.15, 873.15), 2: (773.15, 1473.15)}
    im._presets[100:] = [1] * 100  # a broken pattern
    with pytest.raises(ValueError, match="alternate"):
        tcmatch.preset_frames(im)
    single = _FakeImager(frames)
    single.source_info = _Info([_Preset(True, 200, 273.15, 923.15)])
    assert tcmatch.preset_frames(single) == ({}, {})


def test_sample_recording_on_the_a700_clip() -> None:
    import flir_player  # noqa: F401  (FileSDK DLL preload)
    import fnv.file

    from flir_player.calibration import read_array

    path = SAMPLES / CLIP
    samples = tcmatch.sample_recording(path, MatchOptions(roi=(300, 400, 360, 470)), step_s=1.0)
    assert samples.fps == pytest.approx(30.0)
    assert samples.report.verified and samples.flicker_available
    assert samples.mean.shape == (samples.times.size, 70, 60)
    im = fnv.file.ImagerFile(str(path))
    try:
        im.unit = fnv.Unit.COUNTS
        frames = [k for k in range(int(im.num_frames)) if math.floor(k / samples.fps + 0.5) == 2]
        direct = np.mean([read_array(im, k)[400:470, 300:360].astype(np.float64) for k in frames], axis=0)
    finally:
        im.close()
    assert samples.counts[2] == len(frames) == 30
    np.testing.assert_allclose(samples.mean[2], direct, rtol=1e-6)


# --- 0922 regression (opt-in: reads a 65 GB recording) ------------------------------------------


NOTES_0922 = Path(__file__).resolve().parents[1] / "local" / "notes" / "2026-10-04-0922-fire-test"


@pytest.mark.skipif(os.environ.get("FLIR_TCMATCH_0922") != "1"
                    or not (SAMPLES / "0922" / "0922_A700.seq").exists(),
                    reason="set FLIR_TCMATCH_0922=1 with local/data/0922/0922_A700.seq present")
def test_0922_a700_reproduces_the_one_off_analysis(tmp_path) -> None:
    options = MatchOptions(roi=(250, 380, 560, 480), scope=(0, 2087), valid={"T2": ((0, 1150),)})
    result = tcmatch.run(SAMPLES / "0922" / "0922_A700.seq", NOTES_0922 / "inputs" / "0922_TC_Temperature.xlsx",
                         options, out_dir=tmp_path, cache_dir=SAMPLES / "0922" / "cache" / "tcmatch")
    assert abs(result.lag_s - 730) <= 2
    t1, t2, t3 = result.channels
    assert not t1.match.found and "failed from 67 s" in t1.flags[0]
    # the one-off analysis's pixels lie among the engine's near-equal candidates (a few pixels apart)
    for match, (row, col) in ((t2.match, (462, 428)), (t3.match, (469, 356))):
        rows = [p[0] for p in match.candidates]
        cols = [p[1] for p in match.candidates]
        assert min(rows) <= row <= max(rows) and min(cols) <= col <= max(cols)
        assert abs(match.pixel[0] - row) <= 4 and abs(match.pixel[1] - col) <= 4
    # 0.80 in the one-off analysis; ±3 px and a few seconds move it within 0.7-0.9 on this edge-on view
    assert 0.65 <= result.pooled.eps <= 0.9
    assert not t3.trusted and not any(e.used for e in t3.events)  # TC3 mostly wants ε > 1
    assert (tmp_path / "summary.txt").exists() and (tmp_path / "overview.png").exists()


# --- workbook prefill ---------------------------------------------------------------------------


def test_workbook_prefill_fills_the_tc_sheet(tmp_path) -> None:
    from datetime import datetime as _dt

    import openpyxl
    from test_excel_workbook import _source

    from flir_player.excel_export import ExportData, ExportOptions, TcPrefill
    from flir_player.workbook import TC_COLUMNS, TC_FIRST, write_workbook

    rows = 12
    times = tuple(float(t) for t in range(-3, 15))  # more logger rows than IR rows
    b = [float(v) for v in range(18)]
    b[4] = math.nan
    prefill = TcPrefill(times=times, columns=((20.0,) * 18, tuple(b)), names=("TC A", "TC B"),
                        roi_tc=(("TC1", "TC B"), ("Box", "TC A")), window=(1.0, 5.0), offset_s=0.5)
    data = ExportData(times=np.arange(rows, dtype=float) - 2.0, sources=[_source("Camera A", rows,
                                                                                  np.random.default_rng(7))],
                      options=ExportOptions(bins=16, tc_prefill=prefill), created=_dt(2026, 10, 6),
                      versions={"tool": "test", "sdk": ""})
    path = tmp_path / "book.xlsx"
    write_workbook(path, data)
    book = openpyxl.load_workbook(path)
    tc = book["TC Compare"]
    head = TC_FIRST + 1  # Excel row of the headers
    assert [tc.cell(head, c).value for c in range(1, 5)] == ["Time (s)", "TC A", "TC B", "TC 3"]
    assert tc.cell(6, 2).value == 0.5
    for i, t in enumerate(times):
        assert tc.cell(head + 1 + i, 1).value == t
        assert tc.cell(head + 1 + i, 2).value == 20.0
        assert tc.cell(head + 1 + i, 3).value == (None if i == 4 else b[i])
    # the comparison formulas of the same rows survive (constant-memory rows are written once)
    for i in range(rows):
        assert str(tc.cell(head + 1 + i, TC_COLUMNS + 3).value).startswith("=")
    assert tc.cell(head + 1 + rows, TC_COLUMNS + 3).value is None
    settings = book["Settings"]
    mapped = {settings.cell(r, 3).value: settings.cell(r, 8).value for r in range(1, settings.max_row + 1)
              if settings.cell(r, 3).value in ("TC1", "TC2", "Box")}
    assert mapped == {"TC1": "TC B", "TC2": None, "Box": "TC A"}
    summary = book["Summary"]
    assert (summary.cell(9, 2).value, summary.cell(10, 2).value) == (1.0, 5.0)


def test_tc_prefill_checks_its_shape() -> None:
    from flir_player.excel_export import TcPrefill

    ok = TcPrefill(times=(0.0, 1.0), columns=((1.0, 2.0),), names=("A",), roi_tc=(("spot", "A"),))
    ok.check(10, 10)
    with pytest.raises(ValueError, match="1 to 10"):
        TcPrefill(times=(0.0,), columns=((1.0,),) * 11, names=tuple("ABCDEFGHIJK")).check(10, 10)
    with pytest.raises(ValueError, match="one value per logger time"):
        TcPrefill(times=(0.0, 1.0), columns=((1.0,),), names=("A",)).check(10, 10)
    with pytest.raises(ValueError, match="No TC column named B"):
        TcPrefill(times=(0.0,), columns=((1.0,),), names=("A",), roi_tc=(("spot", "B"),)).check(10, 10)
    with pytest.raises(ValueError, match="exceed"):
        TcPrefill(times=(0.0, 1.0), columns=((1.0, 2.0),), names=("A",)).check(1, 10)


def test_tc_prefill_from_a_result_maps_spots_and_boxes() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=3))
    prefill = tcmatch.tc_prefill(result, table)
    assert prefill.names == ("TC A", "TC B", "TC C")
    assert dict(prefill.roi_tc) == {"TC A spot": "TC A", "TC A 3×3": "TC A", "TC B spot": "TC B",
                                    "TC B 3×3": "TC B"}
    best = max((e for ch in result.channels for e in ch.events if e.used), key=lambda e: e.fit.n)
    assert prefill.window == (best.start_s, best.end_s)
    np.testing.assert_array_equal(np.array(prefill.columns[1]), table.values[:, 1])
    prefill.check(100_000, 10)


# --- Codex review 2026-10-06 (local/notes/2026-10-06-tc-calibration/codex/review_probes.py) -------


def _one_pixel_scene(tc, counts=None, params=PARAMS, eps=0.8):
    tc = np.asarray(tc, dtype=np.float64)
    n = tc.size
    mean = np.asarray(counts_for(tc, eps, params) if counts is None else counts, np.float32).reshape(n, 1, 1)
    samples = Samples(path=Path("synthetic.seq"), times=np.arange(n, dtype=float), mean=mean,
                      spread=np.zeros_like(mean), counts=np.full(n, 30, dtype=np.int32),
                      first_frame=np.arange(n, dtype=np.int64) * 30, origin=(0, 0), size=(1, 1), fps=30.0, step=1.0,
                      report=REPORT, params=params, info={"camera": "synthetic", "frames": n * 30})
    return samples, TcTable(times=np.arange(n, dtype=float), values=tc[:, None], names=("TC A",))


def _bound_scene():
    tc = 100.0 + np.arange(400) * 0.1
    counts = counts_for(tc, 0.8)
    counts[70:] = counts_for(tc[70:] + 100.0, 1.0)  # from 70 s the IR is far above any match
    samples, table = _one_pixel_scene(tc, counts)
    samples.spread[60:70] = 5000.0  # flames split the two events
    return samples, table, MatchOptions(spot=1, lag_s=0.0, pixels={"TC A": (0, 0)})


def _replace(obj, **changes):
    from dataclasses import replace
    return replace(obj, **changes)


def test_r01_an_event_no_emissivity_explains_still_counts_against_trust() -> None:
    samples, table, options = _bound_scene()
    result, _l, _s = tcmatch.analyse(samples, table, options)
    ch = result.channels[0]
    assert [(e.start_s, e.end_s) for e in ch.events] == [(0.0, 59.0), (70.0, 399.0)]
    unexplained = ch.events[1]
    assert unexplained.fit.n == 0 and unexplained.excluded_bound == 330 and not unexplained.physical
    assert not ch.trusted and result.pooled is None
    assert "no fit" in tcmatch.summary_text(result)


def test_r12_cancelling_during_the_fits_stops_the_analysis(monkeypatch) -> None:
    samples, table, options = _bound_scene()
    cancel = [False]
    original = tcmatch.fit_eps

    def cancelling(*args, **kwargs):
        value = original(*args, **kwargs)
        cancel[0] = True
        return value

    monkeypatch.setattr(tcmatch, "fit_eps", cancelling)
    with pytest.raises(tcmatch.JobCancelled):
        tcmatch.analyse(samples, table, options, abort=lambda: cancel[0])


def test_r02_a_tc_is_found_only_at_the_shared_offset_with_enough_overlap() -> None:
    rng = np.random.default_rng(421)
    tc = rng.uniform(90.0, 120.0, 200)
    samples, table = _one_pixel_scene(tc, counts_for(np.r_[tc[:2], tc[:-2]], 0.8))  # IR 2 s late
    options = MatchOptions(spot=1, lag_s=0.0)
    grid, channels = tcmatch.prepare_channels(table, options, 1.0)
    match = tcmatch.locate(samples, grid, channels, options).channels[0]
    assert match.own_lag_s == 2.0 and match.own_r > 0.99  # its own offset fits...
    assert abs(match.r) < 0.2 and not match.found  # ...but the given one does not
    late = counts_for(np.r_[np.full(170, 100.0), tc[:30]], 0.8)  # only 30 of 200 samples overlap
    samples, table = _one_pixel_scene(tc, late)
    options = MatchOptions(spot=1, lag_s=170.0)
    grid, channels = tcmatch.prepare_channels(table, options, 1.0)
    assert not tcmatch.locate(samples, grid, channels, options).channels[0].found


def test_r03_a_chosen_preset_brings_its_own_range() -> None:
    layout = {1: (0, 2), 2: (1, 2)}
    ranges = {1: (523.15, 873.15), 2: (773.15, 1473.15)}
    preset, limits = tcmatch.choose_preset(layout, ranges, 2)
    assert preset == 2 and limits.calibrated == (773.15, 1473.15)
    assert tcmatch.choose_preset(layout, ranges)[0] == 1  # default: the lowest range
    assert tcmatch.choose_preset({}, {}) == (None, None)
    with pytest.raises(ValueError, match="Preset 3"):
        tcmatch.choose_preset(layout, ranges, 3)


def test_r04_a_tc_heating_from_the_start_is_not_failed() -> None:
    table = TcTable(times=np.arange(200.0), values=(20.0 + 2.0 * np.arange(200.0))[:, None], names=("A",))
    flags = failure_flags(table, 0)
    assert flags.failed_from is None and not flags.flagged.any() and flags.baseline_c <= 30.0


class _NoWholeCopy:
    """Statistics that fail when converted whole (as a memmap would be read into RAM)."""

    def __init__(self, array):
        self.array, self.shape, self.dtype, self.largest = array, array.shape, array.dtype, 0

    def __getitem__(self, key):
        out = self.array[key]
        self.largest = max(self.largest, np.size(out))
        return out

    def __array__(self, dtype=None, copy=None):
        raise AssertionError("the whole statistics array was converted")


def test_r05_locating_reads_the_statistics_in_chunks() -> None:
    samples, table = make_scene()
    reference = np.asarray(samples.mean, dtype=np.float64)
    guard = _NoWholeCopy(samples.mean)
    samples.mean = guard
    result, location, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, coarse_pixels=200))
    # the binned offset is refined at full resolution (2 × 2 bins alone gave 140 s)
    assert location.binning > 1 and result.lag_s == 137.0
    b = location.binning
    nb, h, w = reference.shape
    pad = 2 * (MatchOptions().search_radius + b) + b
    bounded = max(128 * h * w, table.times.size * 16 * w, nb * pad * pad)  # bins chunk, r-map rows, refine window
    assert guard.largest <= bounded < reference.size
    expected = reference[:, :h // b * b, :w // b * b].reshape(nb, h // b, b, w // b, b).mean(axis=(2, 4))
    np.testing.assert_allclose(tcmatch._binned_counts(reference.astype(np.float32), b, chunk=50),
                               expected.reshape(nb, -1), rtol=1e-6)


def test_r06_a_logger_starting_before_the_camera_still_gets_a_workbook(monkeypatch) -> None:
    import flir_player.excel_export as excel_export

    samples, table = make_scene(lag=-40, n_ir=500, n_log=480)
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    assert result.ignition_frame == -1200
    assert tcmatch.workbook_origin(result) == (0, -40.0)
    seen = {}

    def fake_export(dest, specs, options, **kwargs):
        seen["spec"], seen["options"] = specs[0], options
        return "ok"

    monkeypatch.setattr(excel_export, "run_export", fake_export)
    tcmatch.write_workbook("book.xlsx", "synthetic.seq", result, table)
    prefill = seen["options"].tc_prefill
    assert seen["spec"].ignition_frame == 0 and prefill.offset_s == -40.0
    assert (seen["options"].start_s, seen["options"].end_s) == (-40.0, 439.0)  # logger 0..479 s
    best = max((e for ch in result.channels for e in ch.events if e.used), key=lambda e: e.fit.n)
    assert prefill.window == (best.start_s - 40.0, best.end_s - 40.0)
    frame, offset = tcmatch.workbook_origin(_replace(result, ignition_frame=10 ** 9))
    assert frame == result.num_frames - 1 and offset == pytest.approx(result.lag_s - frame / result.fps)


def test_r07_the_pixel_spread_covers_the_best_correlated_candidates() -> None:
    tc = np.r_[np.full(70, 100.0), np.linspace(100.0, 300.0, 100)]
    samples, table = _one_pixel_scene(tc)
    mean = np.empty((tc.size, 13, 13), np.float32)
    wobble = 25.0 * np.sin(np.arange(tc.size) * 0.29)
    for y in range(13):
        for x in range(13):
            mean[:, y, x] = counts_for(tc, 0.5 + y * 0.03) + wobble
    mean[:, 6, 6] = counts_for(tc, 0.68)
    samples.mean, samples.spread, samples.size = mean, np.zeros_like(mean), (13, 13)
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=0.0, search_radius=6))
    ch = result.channels[0]
    event = ch.events[0]
    assert ch.match.pixel == (6, 6) and ch.match.candidates[0] == (6, 6)
    assert event.pixel_n == len(ch.match.candidates) == 169
    low, _mid, high = event.pixel_eps
    assert low <= event.fit.eps <= high and low <= 0.52 and high >= 0.84


def test_r08_engine_blocks_and_exported_boxes_agree_at_edges() -> None:
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=3))
    x0, y0 = samples.origin
    _nb, h, w = samples.mean.shape
    assert result.region == [x0, y0, x0 + w, y0 + h]
    for pixel in ((y0, x0), (y0 + h - 1, x0 + w - 1), SPOT_A):
        moved = _replace(result, channels=[_replace(result.channels[0],
                                                    match=_replace(result.channels[0].match, pixel=pixel))])
        box = tcmatch.roi_shapes(moved)[1]
        ys, xs = roi_coordinates(box, samples.size[1], samples.size[0])
        r0, r1, c0, c1 = tcmatch.block_bounds(pixel, 3, samples)
        engine = sorted((r + y0, c + x0) for r in range(r0, r1) for c in range(c0, c1))
        assert sorted(zip(ys.tolist(), xs.tolist())) == engine


def test_r09_hot_surroundings_do_not_make_a_valid_surface_impossible() -> None:
    tc = np.linspace(100.0, 200.0, 200)
    params = PARAMS.with_(reflected_k=573.15)  # surroundings at 300 °C, hotter than the TC
    samples, table = _one_pixel_scene(tc, counts_for(tc, 0.7, params), params=params)
    assert np.all(samples.apparent(samples.mean[:, 0, 0]) > tc + 20)  # IR above the TC, yet ε 0.7 fits
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=0.0, pixels={"TC A": (0, 0)},
                                                                  reflected_c=300.0))
    event = result.channels[0].events[0]
    assert event.excluded_bound == 0 and event.used and event.fit.eps == pytest.approx(0.7, abs=0.002)
    hopeless = tcmatch.no_emissivity_fits(np.array([350.0, 150.0, 90.0, -10.0]), np.full(4, 150.0), 20.0, 20.0)
    assert hopeless.tolist() == [True, False, False, True]  # above both, at TC, towards T_refl, below both


def test_r10_a_missing_reading_keeps_the_rest_of_the_table() -> None:
    start = datetime(2026, 10, 6, 12)
    for dated in (False, True):
        rows = [("Time", "TC A")] + [(start + timedelta(seconds=k) if dated else k, None if k == 4 else 100.0 + k)
                                     for k in range(10)]
        table = parse_rows(rows)
        np.testing.assert_array_equal(table.times, np.arange(10.0))
        assert np.isnan(table.values[4, 0]) and table.values[3, 0] == 103.0
    table = parse_rows([("Time", "A"), (2, None)] + [(k, 10.0 + 0.5 * k) for k in range(5)])
    np.testing.assert_array_equal(table.times, np.arange(5.0))  # a lone note above the data is no row


def test_r11_a_cancelled_rebuild_leaves_no_complete_cache(tmp_path) -> None:
    import gc

    n = 800
    im = _FakeImager(np.arange(n, dtype=np.float32).reshape(n, 1, 1) + 1000.0)
    cached = tcmatch._accumulate(im, np.arange(n), 10.0, 1.0, (0, 0, 1, 1), None, lambda: False, tmp_path, "k")
    del cached
    gc.collect()
    (tmp_path / "k_spread.npy").write_bytes(b"damaged")
    assert tcmatch._load_cache(tmp_path, "k") is None
    calls = [0]

    def abort():
        calls[0] += 1
        return calls[0] > 1

    with pytest.raises(tcmatch.JobCancelled):
        tcmatch._accumulate(im, np.arange(n), 10.0, 1.0, (0, 0, 1, 1), None, abort, tmp_path, "k")
    gc.collect()
    assert tcmatch._load_cache(tmp_path, "k") is None


# --- Codex follow-up (codex/followup_probes.py) -------------------------------------------------


def test_f1_explicit_limits_win_over_the_file_record() -> None:
    import flir_player  # noqa: F401  (FileSDK DLL preload)
    import fnv.file

    from flir_player.calibration import derive_calibration

    im = fnv.file.ImagerFile(str(SAMPLES / CLIP))
    try:
        report = derive_calibration(im, SAMPLES / CLIP, frames=3, limits=RangeLimits(calibrated=(773.15, 1473.15)))
    finally:
        im.close()
    assert report.verified and "FFF" in report.message
    assert report.calibration.limits.calibrated == (773.15, 1473.15)


def test_f2_cancelling_during_the_pooled_fit_stops_the_analysis(monkeypatch) -> None:
    samples, table = _one_pixel_scene(100.0 + np.arange(200) * 0.5)
    options = MatchOptions(spot=1, lag_s=0.0, pixels={"TC A": (0, 0)})
    calls = [0]
    original = tcmatch.fit_eps

    def counting(*args, **kwargs):
        calls[0] += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(tcmatch, "fit_eps", counting)
    tcmatch.analyse(samples, table, options)
    pooled_call = calls[0]  # the last fit is the pooled one
    calls[0] = 0
    with pytest.raises(tcmatch.JobCancelled):
        tcmatch.analyse(samples, table, options, abort=lambda: calls[0] >= pooled_call)


def test_f3_the_refined_offset_stays_in_the_requested_range() -> None:
    samples, table = make_scene()
    for window in ((140.0, 140.0), (140.0, 150.0), (120.0, 135.0)):
        result, location, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, coarse_pixels=200,
                                                                            lag_range=window))
        assert location.binning > 1 and window[0] <= result.lag_s <= window[1]
    free, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, coarse_pixels=200))
    assert free.lag_s == 137.0


def test_f4_status_text_and_text_readings_do_not_cut_the_table() -> None:
    status = parse_rows([("Time", "TC A", "TC B", "Status")]
                        + [(k, 20.0 + k, 30.0 + 2 * k, "OK") for k in range(10)])
    assert status.names == ("TC A", "TC B")
    np.testing.assert_array_equal(status.times, np.arange(10.0))
    marker = parse_rows([("Time", "TC A", "TC B")]
                        + [(k, 20.0 + k, "OPEN" if k == 4 else 30.0 + 2 * k) for k in range(10)])
    assert marker.names == ("TC A", "TC B")
    np.testing.assert_array_equal(marker.times, np.arange(10.0))
    assert np.isnan(marker.values[4, 1]) and marker.values[4, 0] == 24.0
    # one TC: a text reading, or a blank reading beside a status, is a gap, not a header
    single = parse_rows([("Time", "TC A")] + [(k, "OPEN" if k == 4 else 20.0 + k) for k in range(10)])
    blank = parse_rows([("Time", "TC A", "Status")] + [(k, None if k == 4 else 20.0 + k, "OK") for k in range(10)])
    for table in (single, blank):
        assert table.names == ("TC A",)
        np.testing.assert_array_equal(table.times, np.arange(10.0))
        assert np.isnan(table.values[4, 0]) and table.values[5, 0] == 25.0
    # above the data, a row with one number and text stays a note
    noted = parse_rows([("Interval", 1), ("Time", "TC A")] + [(k, 20.0 + 0.5 * k) for k in range(6)])
    assert noted.names == ("TC A",) and noted.times.size == 6
    # summary rows under the data (two numbers beside text) end it and are reported
    summed = parse_rows([("Time", "A", "B"), ("[s]", "[C]", "[C]")] + [(k, 20.0 + k * 0.5, 30.0) for k in range(5)]
                        + [("Max", 22.0, 30.0), ("Mean", 21.0, 30.0)])
    assert summed.names == ("A", "B") and summed.times.tolist() == [0.0, 1.0, 2.0, 3.0, 4.0]
    assert any("after the data" in note for note in summed.notes)
    # the 0922 layout (a note and its value above the data) still parses
    table = read_tc_table(NOTES_0922 / "inputs" / "0922_TC_Temperature.xlsx") if (
        NOTES_0922 / "inputs" / "0922_TC_Temperature.xlsx").exists() else None
    if table is not None:
        assert table.names[1] == "T2 (BatteryCell#6)" and table.times[0] == 0.0 and table.times.size == 2248


def test_f5_the_bounded_fit_is_the_same_in_blocks(monkeypatch) -> None:
    rng = np.random.default_rng(11)
    tc = rng.uniform(80.0, 400.0, 500)
    counts = counts_for(tc, 0.77) + rng.normal(0, 30, tc.size)
    whole = tcmatch.fit_eps(A700, PARAMS, counts, tc)
    monkeypatch.setattr(tcmatch, "FIT_BLOCK", 37)
    blocked = tcmatch.fit_eps(A700, PARAMS, counts, tc)
    assert blocked.eps == whole.eps and blocked.rmse_k == pytest.approx(whole.rmse_k, rel=1e-12)
    assert whole.eps == pytest.approx(0.77, abs=0.01)
