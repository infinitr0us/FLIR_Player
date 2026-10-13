# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""0.7.3: the time match near the camera and logger clocks (a lab workbook's LogFile sheet and the Log row of
logger time 0, against the recording's first frame), and the lab workbooks' "Cell Temp" and "TRIM temp"
unticked like the TC tree (the engineers' reply of 2026-10-13: their 1006 A700 fit matched 33 minutes off)."""

from __future__ import annotations

import re
import zipfile
from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from PySide6.QtWidgets import QLabel

from flir_player import tcmatch
from flir_player.models import VideoMetadata
from flir_player.tcdata import TcTable, logger_clock, read_tc_table
from flir_player.tcmatch import CLOCK_WINDOW_S, MatchOptions, MatchResult, clock_estimate
from test_tc_dialog import _loaded, _write_csv
from test_tcmatch import make_scene

DATA = Path(__file__).resolve().parents[1] / "local" / "data"
START = datetime(2026, 10, 6, 10, 46, 36, 701000)


def _lab_workbook(path: Path, *, shift: int = 360, rows: int = 900, half_edited: bool = False,
                  formulas: bool = True, logfile: bool = True, date_cell: bool = True, stale_from: int | None = None,
                  refs=None, column: str = "E", elapsed_from: float = 0.0) -> Path:
    """A workbook laid out like the engineers' (LogFile, Log and TC sheets; TC times as formulas with their
    values cached, as Excel saves them). ``half_edited``: rows 5-6 subtract the Log row 217, the rest 360;
    ``refs(i)`` gives (Log row, Log row of time 0) of data row i instead; ``stale_from``: cached times 5 s off
    their formulas from that data row on; ``column``: the Log column subtracted (E elapsed seconds counting
    from ``elapsed_from``, F a temperature)."""
    if refs is None:
        def refs(i):
            if half_edited and i < 2:
                return (217, 217) if i == 0 else (shift + 1, 217)
            return shift + i, shift
    pairs = [refs(i) for i in range(rows)]
    book = openpyxl.Workbook()
    book.active.title = "Constants"
    if logfile:
        ws = book.create_sheet("LogFile_2026-10-06-10-46-36 (ro")
        ws.append(["Root Name", "Title", "Author", "Date/Time" if date_cell else "Created", "Groups"])
        ws.append(["LogFile_2026-10-06-10-46-36", None, None, START if date_cell else "today", 2])
    log = book.create_sheet("Log")
    log.append(["A_Time*", "A", "B_Time*", "B", "Tree_Time*", "Tree"])
    logged = {"E": lambda k: k + elapsed_from, "F": lambda k: 20.0 + 0.5 * k}  # Log row k + 2 holds second k
    for k in range(max(max(pair) for pair in pairs)):
        log.append([k, 0.0, k, 0.0, logged["E"](k), logged["F"](k)])
    tc = book.create_sheet("TC")
    tc.append(["Time", "ELAPSED TIME", None, None, None, "Free_0r", "Free_0r", "Free_0r"])
    tc.append(["From Ignition", "[sec]", None, None, None, "T1", "Cell Temp", "TRIM temp"])
    tc.append([None, shift])
    tc.append(["[mim]", "shft time"])
    cached = {}
    for i, (source, ref) in enumerate(pairs):
        r = 5 + i
        value = logged[column](source - 2) - logged[column](ref - 2)
        if stale_from is not None and i >= stale_from:
            value += 5
        b = f"=Log!{column}{source}-Log!{column}${ref}" if formulas else value
        tc.append([f"=(B{r}-0)/60" if formulas else value / 60, b, None, None, None,
                   20.0 + 10.0 * abs(np.sin(i / 30.0)), 600.0 - 0.1 * i, 650.0 - 0.1 * i])  # no rising column
        cached[f"B{r}"], cached[f"A{r}"] = value, value / 60
    book.save(path)
    if formulas:  # cache the formulas' values, as Excel does (openpyxl writes none)
        index = book.sheetnames.index("TC") + 1
        name = f"xl/worksheets/sheet{index}.xml"
        with zipfile.ZipFile(path) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        xml = parts[name].decode("utf-8")
        xml = re.sub(r'<c r="([A-Z]+\d+)"><f>([^<]*)</f><v\s*/>',
                     lambda m: f'<c r="{m[1]}"><f>{m[2]}</f><v>{cached[m[1]]}</v>', xml)
        parts[name] = xml.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for n, blob in parts.items():
                z.writestr(n, blob)
    return path


# --- Cell Temp and TRIM temp ------------------------------------------------------------------


def test_cell_and_trim_temps_start_unticked_like_the_tc_tree(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import _AIR_HINT, TcCalibrationDialog

    for name in ("Cell Temp", "TRIM temp", "cell  temp", "Bottom TC tree (1 m)", "TC Tree (1.5 m)"):
        assert _AIR_HINT.search(name), name
    for name in ("T1", "S2", "Free_0r T3", "Cell 3", "Cell Temperature", "Trimmed edge", "cells temp"):
        assert not _AIR_HINT.search(name), name
    tc = _write_csv(tmp_path / "tc.csv", ["T1", "T2", "Cell Temp", "TRIM temp", "TC Tree (2.0 m)"],
                    [[k, 20.0 + k, 21.0 + k, 600.0, 650.0, 30.0] for k in range(60)])
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0)
    dialog.file_edit.setText(str(tc))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog) and dialog.chosen_channels() == ["T1", "T2"]
    dialog.deleteLater()


# --- the logger's clock ---------------------------------------------------------------------------


def test_a_lab_workbook_gives_the_logger_clock_at_time_0(tmp_path) -> None:
    table = read_tc_table(_lab_workbook(tmp_path / "lab.xlsx"))
    assert table.sheet == "TC" and table.times[0] == 0.0 and table.time_rows == tuple(range(5, 905)) and table.time_col == 2
    assert table.clock_zero == START + timedelta(seconds=358)  # Log row 360 = elapsed 358 s
    assert table.clock_note == ("logging started 2026-10-06 10:46:36 by sheet 'LogFile_2026-10-06-10-46-36 (ro' "
                                "(its Date/Time cell); logger time 0 is Log!E360, 358 s later")
    shifted = read_tc_table(_lab_workbook(tmp_path / "lab2.xlsx", shift=2))  # no shift: time 0 is the start
    assert shifted.clock_zero == START


def test_a_half_edited_sheet_takes_the_clock_of_the_rows_read(tmp_path) -> None:
    table = read_tc_table(_lab_workbook(tmp_path / "lab.xlsx", half_edited=True))
    assert table.time_rows[0] == 7 and table.time_col == 2 and any("left out" in n for n in table.notes)
    assert table.clock_zero == START + timedelta(seconds=358)  # rows 7+ subtract Log row 360, not 217


def test_no_clock_unless_the_workbook_says_it(tmp_path) -> None:
    values = read_tc_table(_lab_workbook(tmp_path / "values.xlsx", formulas=False))
    assert values.clock_zero is None and "not formulas" in values.clock_note
    no_log = read_tc_table(_lab_workbook(tmp_path / "nolog.xlsx", logfile=False))
    assert no_log.clock_zero is None and "no LogFile sheet" in no_log.clock_note
    stale = read_tc_table(_lab_workbook(tmp_path / "stale.xlsx", stale_from=0))
    assert stale.clock_zero is None and "row 5 does not hold the difference" in stale.clock_note
    by_name = read_tc_table(_lab_workbook(tmp_path / "name.xlsx", date_cell=False))
    assert by_name.clock_zero == START.replace(microsecond=0) + timedelta(seconds=358)
    assert "its name" in by_name.clock_note
    csv = read_tc_table(_write_csv(tmp_path / "tc.csv", ["T1"], [[k, 20.0 + k] for k in range(10)]))
    assert csv.clock_zero is None and "text table" in csv.clock_note
    table = read_tc_table(_lab_workbook(tmp_path / "lab.xlsx"))
    assert logger_clock(tmp_path / "missing.xlsx", "TC", table)[0] is None  # never raises


@pytest.mark.parametrize("name, sheet, zero", [
    ("1006/20261006a.xlsx", "TC", datetime(2026, 10, 6, 10, 52, 34, 701000)),
    ("1006/20261006b.xlsx", "TC", datetime(2026, 10, 6, 14, 14, 54, 927000)),  # half edited: Log row 360
    ("0922/20260922.xlsx", "TC", datetime(2026, 9, 22, 11, 13, 29, 645000)),
    ("0903/20260903.xlsx", "TC", datetime(2026, 9, 3, 10, 57, 38, 343000)),
])
def test_the_engineers_workbooks_give_their_clocks(name, sheet, zero) -> None:
    path = DATA / name
    if not path.is_file():
        pytest.skip(f"{name} not in local/data")
    assert read_tc_table(path, sheet=sheet).clock_zero == zero


# --- the clocks' estimate ---------------------------------------------------------------------------


def test_clock_estimate_sets_whole_hours_aside() -> None:
    log = np.arange(0.0, 2143.0)
    # 1006 A700: logger local, camera UTC (4 h apart); the clocks give 400 s, not 4,000 s
    a700 = clock_estimate(datetime(2026, 10, 6, 10, 52, 34, 701000), datetime(2026, 10, 6, 14, 45, 54, 602000),
                          log, 6581.0)
    assert a700.lag_s == pytest.approx(400.099) and a700.hours == 4 and "4 h" in a700.note
    # 0922 T650: a zone without daylight saving (5 h); the hour before puts too little of the log inside
    t650 = clock_estimate(datetime(2026, 9, 22, 11, 13, 29, 645000), datetime(2026, 9, 22, 16, 2, 34, 9000),
                          np.arange(0.0, 3300.0), 3485.0)
    assert t650.lag_s == pytest.approx(655.636) and t650.hours == 5
    start = datetime(2026, 1, 1, 12)
    # the logger started 10 min after the camera ... or 10 min before it (most of the log still inside)
    assert clock_estimate(start + timedelta(seconds=600), start, log, 6581.0).lag_s == 600.0
    assert clock_estimate(start - timedelta(seconds=600), start, log, 6581.0).lag_s == -600.0
    # a recording shorter than the log: half the recording must lie inside the log
    assert clock_estimate(start - timedelta(seconds=1000), start, log, 300.0).lag_s == -1000.0
    assert clock_estimate(None, start, log, 600.0) is None and clock_estimate(start, None, log, 600.0) is None
    assert clock_estimate(start + timedelta(days=5), start, log, 6581.0) is None  # days apart: no clock


# --- the time match near the clocks -------------------------------------------------------------------


def _clocked(lag_by_clocks: float, **scene):
    samples, table = make_scene(**scene)
    start = datetime(2026, 1, 1, 12)
    samples.info["clock_start"] = start
    table = TcTable(times=table.times, values=table.values, names=table.names, source=table.source,
                    clock_zero=start + timedelta(seconds=lag_by_clocks), clock_note="logging started (test)")
    return samples, table


def test_the_search_stays_near_the_clocks_and_reports_them() -> None:
    samples, table = _clocked(137 + 30)  # clocks 30 s off the true 137 s
    result, _location, _series = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S))
    assert result.lag_s == 137.0 and result.clock_lag_s == 167.0 and result.clock_frame == 167 * 30
    assert not any("clock" in w for w in result.warnings)
    text = tcmatch.summary_text(result)
    assert "(searched within 120 s of the clocks)" in text and "Clocks: logger time 0 at about 167 s" in text
    assert "30 s later than the match" in text and "  Logging started (test); the recording's first frame" in text
    back = MatchResult.from_dict(result.to_dict())
    assert (back.clock_lag_s, back.clock_frame, back.clock_note) == (167.0, 167 * 30, result.clock_note)
    assert back.options["clock_window_s"] == CLOCK_WINDOW_S and back.options["lag_range"] is None
    old = result.to_dict()
    for key in ("clock_lag_s", "clock_frame", "clock_note"):
        old.pop(key)
    assert MatchResult.from_dict(old).clock_lag_s is None  # results from before 0.7.3


def test_a_match_far_from_the_clocks_is_reported() -> None:
    samples, table = _clocked(137 + 200)
    free, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    assert free.lag_s == 137.0
    assert any("clocks put logger time 0 at 337 s" in w and "200 s away" in w for w in free.warnings)
    given, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_s=137.0))
    assert any(w.startswith("The given time match (137 s) is 200 s from") for w in given.warnings)
    edge, _l, _s = tcmatch.analyse(*_clocked(137 + 119), MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S))
    assert edge.lag_s == 137.0 and any("lies at the edge of the window" in w for w in edge.warnings)


def test_without_a_logger_clock_the_whole_recording_is_searched() -> None:
    samples, table = make_scene()
    samples.info["clock_start"] = datetime(2026, 1, 1, 12)
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S))
    assert result.lag_s == 137.0 and result.clock_lag_s is None
    assert any(w.startswith("No time match near the clocks: the TC file does not give the logger's clock")
               for w in result.warnings)
    assert "Clocks: not compared, as the TC file does not give the logger's clock." in tcmatch.summary_text(result)


# --- the dialog ------------------------------------------------------------------------------------------


def _metadata(tmp_path, start=None, duration=6581.0) -> VideoMetadata:
    return VideoMetadata(path=tmp_path / "test.seq", width=640, height=480, num_frames=int(duration * 30) + 1,
                         start_time=start, end_time=None, duration_seconds=duration, nominal_fps=30.0,
                         camera_model="FLIR A700")


def test_the_dialog_searches_near_the_clocks_by_default(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog

    lab = _lab_workbook(tmp_path / "lab.xlsx")
    zero = START + timedelta(seconds=358)
    start = zero - timedelta(hours=4, seconds=400)  # a camera on UTC, started 400 s before logger time 0
    dialog = TcCalibrationDialog(_metadata(tmp_path, start), (), current_frame=0)
    assert dialog.time_combo.currentData() == "clock" and not dialog.clock_label.isVisibleTo(dialog)
    dialog.file_edit.setText(str(lab))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog) and dialog.chosen_channels() == ["T1"]
    assert dialog.clock_label.text() == "Clocks: logger time 0 at about frame 12,001, 6 min 40 s into the recording."
    assert "Log!E360" in dialog.clock_label.toolTip()
    assert dialog.options().clock_window_s == CLOCK_WINDOW_S and dialog.options().lag_frame is None
    dialog.time_combo.setCurrentIndex(dialog.time_combo.findData("auto"))
    assert dialog.options().clock_window_s is None and "more than 120 s from it is reported" in dialog.clock_label.text()
    dialog.time_combo.setCurrentIndex(dialog.time_combo.findData("frame"))
    dialog.frame_spin.setValue(71461)
    assert "The frame given is 33 min from it" in dialog.clock_label.text()
    dialog.deleteLater()
    csv = _write_csv(tmp_path / "tc.csv", ["T1"], [[k, 20.0 + k] for k in range(60)])
    dialog = TcCalibrationDialog(_metadata(tmp_path, start), (), current_frame=0)
    dialog.file_edit.setText(str(csv))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog)
    assert dialog.clock_label.text() == ("Clocks: the TC file does not give the logger's clock, so the whole "
                                         "recording is searched.")
    dialog.deleteLater()


def test_the_dialog_restores_the_time_mode(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog, saved_request

    for options, mode in ((MatchOptions(), "auto"), (MatchOptions(clock_window_s=CLOCK_WINDOW_S), "clock"),
                          (MatchOptions(lag_frame=99), "frame")):
        dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0, previous={"options": options})
        assert dialog.time_combo.currentData() == mode
        dialog.deleteLater()
    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1))
    old = result.to_dict()
    old["options"].pop("clock_window_s")  # a fit from before 0.7.3: offered near the clocks again
    def saved(res):
        return tcmatch.SavedRun(result=res, folder=tmp_path, tc_file=None, tc_name="lab.xlsx", sheet=None,
                                time_column=None, parameters=None, options=tcmatch.options_from_dict(res.options))

    assert saved_request(saved(MatchResult.from_dict(old)), table)["options"].clock_window_s == CLOCK_WINDOW_S
    assert saved_request(saved(result), table)["options"].clock_window_s is None  # "anywhere" was chosen


def test_the_results_headline_names_the_clocks(qapp) -> None:
    from flir_player.tc_dialog import TcResultsDialog

    samples, table = _clocked(137 + 30)
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S))
    dialog = TcResultsDialog(tcmatch.RunOutput(result=result, table=table), None)
    texts = [label.text() for label in dialog.findChildren(QLabel)]
    assert any("The clocks put it at frame 5,011, 30 s later than the match." in t for t in texts)
    dialog.deleteLater()


# --- review findings (Codex, 2026-10-13) ------------------------------------------------------------------


def test_r1_every_time_cell_must_subtract_the_same_log_row_and_hold_its_difference(tmp_path) -> None:
    # row 5 subtracts Log row 360, the rest Log row 960: still rising, so the reader keeps every row
    mixed = read_tc_table(_lab_workbook(tmp_path / "mixed.xlsx",
                                        refs=lambda i: (360, 360) if i == 0 else (960 + i, 960)))
    assert len(mixed.time_rows) == 900 and not any("left out" in n for n in mixed.notes)
    assert mixed.clock_zero is None and "subtract different logger times (Log!E360, Log!E960)" in mixed.clock_note
    stale = read_tc_table(_lab_workbook(tmp_path / "stale.xlsx", stale_from=40))  # later rows not recalculated
    assert stale.clock_zero is None and "row 45 does not hold the difference" in stale.clock_note


def test_r2_only_an_elapsed_time_counting_from_the_start_of_logging_gives_the_clock(tmp_path) -> None:
    shifted = read_tc_table(_lab_workbook(tmp_path / "shifted.xlsx", elapsed_from=-900.0))
    assert shifted.clock_zero is None and "does not start at 0" in shifted.clock_note
    temperature = read_tc_table(_lab_workbook(tmp_path / "temp.xlsx", column="F"))
    assert temperature.clock_zero is None and "Log!F1 does not head an elapsed time" in temperature.clock_note


def test_r3_a_window_without_offsets_searches_everything_and_an_empty_range_is_an_error() -> None:
    samples, table = _clocked(137 + 30.25)  # no whole second within 0.1 s of 167.25 s
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=0.1))
    assert result.lag_s == 137.0 and not result.near_clocks
    assert any("within 0.1 s of the clocks' 167 s too little of the TC data falls on the recording's frames"
               in w for w in result.warnings)
    with pytest.raises(ValueError, match="No time offset from 5000 to 6000 s"):
        tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_range=(5000.0, 6000.0)))


def test_r4_superframing_shows_the_clocks_in_seconds_only(qapp, tmp_path) -> None:
    from flir_player.models import PresetRange
    from flir_player.tc_dialog import TcCalibrationDialog

    lab = _lab_workbook(tmp_path / "lab.xlsx")
    zero = START + timedelta(seconds=358)
    meta = replace(_metadata(tmp_path, zero - timedelta(seconds=400)),
                   presets=(PresetRange(1, 7000, 523.0, 873.0), PresetRange(2, 7000, 773.0, 1473.0)))
    dialog = TcCalibrationDialog(meta, (), current_frame=0)
    dialog.file_edit.setText(str(lab))
    dialog._load(sheet=None)
    assert _loaded(qapp, dialog)
    assert dialog.clock_label.text() == "Clocks: logger time 0 at about 6 min 40 s into the recording."
    dialog.time_combo.setCurrentIndex(dialog.time_combo.findData("frame"))
    dialog.frame_spin.setValue(9000)
    assert "frame given" not in dialog.clock_label.text()  # frames are clock-timed: no frame arithmetic
    dialog.deleteLater()


def test_r5_the_summary_names_the_clock_window_only_when_it_was_used() -> None:
    samples, table = _clocked(137 + 30)
    ranged, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S,
                                                                   lag_range=(100.0, 200.0)))
    assert ranged.lag_s == 137.0 and not ranged.near_clocks
    assert "searched within" not in tcmatch.summary_text(ranged)
    assert MatchResult.from_dict(ranged.to_dict()).near_clocks is False


def test_r6_a_given_frame_is_kept_as_given_and_older_given_offsets_reopen_as_frames(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog, saved_request

    samples, table = make_scene()
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, lag_frame=137 * 30))
    assert result.lag_fixed and result.options["lag_frame"] == 137 * 30 and result.options["lag_s"] is None

    def saved(res):
        return tcmatch.SavedRun(result=res, folder=tmp_path, tc_file=None, tc_name="lab.xlsx", sheet=None,
                                time_column=None, parameters=None, options=tcmatch.options_from_dict(res.options))

    old = result.to_dict()
    old["options"].update(lag_frame=None, lag_s=137.0)  # how fits before 0.7.3 recorded a given frame
    old["options"].pop("clock_window_s")
    previous = saved_request(saved(MatchResult.from_dict(old)), table)
    assert previous["options"].lag_frame == 137 * 30 and previous["options"].clock_window_s is None
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0, previous={"options": previous["options"]})
    assert dialog.time_combo.currentData() == "frame" and dialog.frame_spin.value() == 137 * 30 + 1
    dialog.deleteLater()


def test_r7_rows_the_reader_dropped_do_not_cost_the_clock(tmp_path) -> None:
    path = _lab_workbook(tmp_path / "lab.xlsx", rows=60)
    book = openpyxl.load_workbook(path)
    tc = book["TC"]
    tc["B20"] = "#VALUE!"  # a broken time: the reader skips the row
    tc["B30"] = "=Log!E999-Log!E$5"  # a repeated time with another origin: the reader drops it
    book.save(path)
    with zipfile.ZipFile(path) as z:  # openpyxl dropped the cached values: put them back (B30 repeats B29)
        parts = {n: z.read(n) for n in z.namelist()}
    name = f"xl/worksheets/sheet{book.sheetnames.index('TC') + 1}.xml"
    xml = parts[name].decode("utf-8")
    xml = re.sub(r'<c r="([AB])(\d+)"([^>]*)><f>([^<]*)</f><v\s*/>',
                 lambda m: (f'<c r="{m[1]}{m[2]}"{m[3]}><f>{m[4]}</f><v>'
                            f'{(int(m[2]) - 5 - (m[2] == "30")) / (60 if m[1] == "A" else 1)}</v>'), xml)
    parts[name] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, blob in parts.items():
            z.writestr(n, blob)
    table = read_tc_table(path, sheet="TC")
    assert 20 not in table.time_rows and 30 not in table.time_rows and len(table.time_rows) == 58
    assert any("without a readable time" in n for n in table.notes) and any("repeating" in n for n in table.notes)
    assert table.clock_zero == START + timedelta(seconds=358), table.clock_note


def test_r8_the_window_needs_real_overlap_not_just_an_end_offset() -> None:
    samples, table = _clocked(880.0, n_ir=760)  # the recording goes on (2,000 s) after the statistics end
    samples.info["frames"] = 2000 * 30
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, clock_window_s=CLOCK_WINDOW_S))
    assert not result.near_clocks and result.lag_s == 137.0
    assert any("too little of the TC data falls on the recording's frames" in w for w in result.warnings)


def test_r9_a_given_offset_before_the_recording_reopens_as_given(qapp, tmp_path) -> None:
    from flir_player.tc_dialog import TcCalibrationDialog

    previous = {"options": MatchOptions(lag_frame=-3000)}  # logger time 0 100 s before the first frame
    dialog = TcCalibrationDialog(_metadata(tmp_path), (), current_frame=0, previous=previous)
    assert dialog.time_combo.currentData() == "frame" and dialog.frame_spin.value() == -2999
    assert dialog.options().lag_frame == -3000
    dialog.deleteLater()


@pytest.mark.parametrize("clocks, rows, overlap", [
    (850.0, None, 0.6),  # at most 30 bins overlap where each TC needs 60 % of its data (288)
    (650.0, np.r_[0:20, 300:480], 0.0),  # a gapped log: at most 20 bins overlap, the scan needs 30
])
def test_r10_the_window_needs_the_overlap_the_scan_needs(clocks, rows, overlap) -> None:
    samples, table = _clocked(clocks, n_ir=760)
    samples.info["frames"] = 2000 * 30  # the recording goes on after the statistics end
    if rows is not None:
        table = TcTable(times=table.times[rows], values=table.values[rows], names=table.names, source=table.source,
                        clock_zero=table.clock_zero, clock_note=table.clock_note)
    result, _l, _s = tcmatch.analyse(samples, table, MatchOptions(spot=1, min_overlap=overlap,
                                                                  clock_window_s=CLOCK_WINDOW_S))
    assert not result.near_clocks and result.lag_s == 137.0
    assert sum(ch.match.found for ch in result.channels) == 2
    assert any("too little of the TC data falls on the recording's frames" in w for w in result.warnings)
