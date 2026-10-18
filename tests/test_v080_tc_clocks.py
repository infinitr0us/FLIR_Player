# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""TC Compare from logger data and clocks, without fitting TC pixels or emissivity."""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path
from threading import Event
from types import SimpleNamespace

import flir_player  # load the SDK DLL directory before numpy
import numpy as np
import openpyxl
import pytest
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QDialog, QDialogButtonBox

from conftest import wait_until
from flir_player import excel_dialog, tc_dialog, tcmatch
from flir_player.excel_dialog import ExcelExportDialog
from flir_player.models import PresetRange, RoiShape, VideoMetadata
from flir_player.settings import app_settings
from flir_player.tcdata import TcTable
from flir_player.tcmatch import MatchOptions, RunOutput

START = datetime(2026, 10, 6, 12)


def table(names=("T1", "T2", "T3"), *, lag=12.3):
    times = np.arange(40.0)
    values = np.column_stack([60 + times + k for k in range(len(names))])
    return TcTable(times, values, names, sheet="TC", clock_zero=START + timedelta(seconds=lag))


def test_clock_prefill_offset_order_cap_and_no_fitted_values():
    data = table(tuple(f"T{k}" for k in range(1, 13)))
    filled = tcmatch.clock_prefill(data, lag_s=12.3, fps=10, frame=143,
                                  pairs={"T12": "Cell 9", "T2": "Cell 6"})
    assert filled.offset_s == pytest.approx(-2.0)
    assert filled.names == ("T12", "T2", "T1", "T3", "T4", "T5", "T6", "T7", "T8", "T9")
    assert filled.columns[0] == tuple(data.values[:, 11])
    assert filled.columns[1] == tuple(data.values[:, 1])
    assert filled.times == tuple(data.times)
    assert filled.roi_tc == (("Cell 9", "T12"), ("Cell 6", "T2"))
    assert filled.roi_eps == () and filled.match_exclude == ()
    assert tcmatch.clock_prefill(data, lag_s=12.3, fps=10, frame=103,
                               pairs={"T1": "Cell 3"}).offset_s == pytest.approx(2.0)
    filled.check(1000, 10)


def test_clock_prefill_valid_hot_whole_bins_and_longest_run():
    times = np.arange(0.0, 20.0, 2.0)
    values = np.column_stack((40 + times * 2, 60 + times))
    values[6, 0] = np.nan
    data = TcTable(times, values, ("T1", "T2"))
    options = MatchOptions(valid={"T1": ((4, 16),), "T2": ((2, 6),)}, jump_hold_s=0)
    filled = tcmatch.clock_prefill(data, lag_s=5, fps=10, frame=20,
                                  pairs={"T1": "Cell 3", "T2": "Cell 6"}, options=options)
    # T1: hot at 6, 8, 10 then missing at 12, then 14, 16. Offset +3, half bin 1.
    assert dict(filled.roi_use) == {"Cell 3": ((8.0, 14.0), (16.0, 20.0)),
                                   "Cell 6": ((4.0, 10.0),)}
    assert filled.window == (8.0, 14.0)
    cold = tcmatch.clock_prefill(data, lag_s=0, fps=10, frame=0, pairs={"T1": "Cell 3"},
                                options=replace(options, tc_min_c=500))
    assert cold.roi_use == (("Cell 3", ()),) and cold.window is None


def test_clock_prefill_flame_jump_removes_rows():
    data = table(("T1",))
    data.values[10:, 0] += 80
    options = MatchOptions(jump_window_s=2, jump_hold_s=3)
    filled = tcmatch.clock_prefill(data, lag_s=0, fps=10, frame=0,
                                  pairs={"T1": "Cell 3"}, options=options)
    assert filled.roi_use == (("Cell 3", ((-0.5, 7.5), (14.5, 39.5))),)
    assert filled.window == (14.5, 39.5)


def test_clock_prefill_refuses_a_paired_tc_beyond_sheet_width():
    data = table(tuple(f"T{k}" for k in range(11)))
    with pytest.raises(ValueError, match="10"):
        tcmatch.clock_prefill(data, lag_s=0, fps=10, frame=0,
                             pairs={n: f"Cell {k}" for k, n in enumerate(data.names)})


def test_clock_prefill_keeps_exact_case_sensitive_columns():
    data = table(("t1", "T1"))
    data.values[:, 0] = 20
    filled = tcmatch.clock_prefill(data, lag_s=0, fps=10, frame=0, pairs={"T1": "Cell 3"})
    # the workbook finds columns with Excel's MATCH, which ignores capitals: "t1" stays off the sheet
    assert filled.names == ("T1",)
    assert filled.columns[0] == tuple(data.values[:, 1])
    assert filled.roi_use == (("Cell 3", ((-0.5, 39.5),)),)
    with pytest.raises(ValueError, match="capitals"):
        tcmatch.clock_prefill(data, lag_s=0, fps=10, frame=0, pairs={"T1": "Cell 3", "t1": "Cell 6"})
    _pairs, errors = excel_dialog.parse_tc_pairs("T1 = Cell 3, t1 = Cell 6", ("t1", "T1"), ("Cell 3", "Cell 6"))
    assert any("capitals" in e for e in errors)


@pytest.mark.parametrize("text", ["T1 = Cell 3, T2 = Cell 6", "  T1  =  Cell 3 ; T2=Cell 6  "])
def test_pair_text(text):
    pairs, errors = excel_dialog.parse_tc_pairs(text, ("T1", "T2"), ("Cell 3", "Cell 6"))
    assert pairs == {"T1": "Cell 3", "T2": "Cell 6"} and not errors


@pytest.mark.parametrize("text,rois,error", [
    ("T4 = Cell 3", ("Cell 3",), "Unknown TC"),
    ("t1 = Cell 3", ("Cell 3",), "Unknown TC"),
    ("T1 = cell 3", ("Cell 3",), "Unknown ROI"),
    ("T1 = Cell 3, T1 = Cell 6", ("Cell 3", "Cell 6"), "TC T1 is paired twice"),
    ("T1 = Cell 3; T2 = Cell 3", ("Cell 3",), "ROI Cell 3 is paired twice"),
    ("T1 = Cell 3", ("Cell 3", "Cell 3"), "More than one ROI"),
    ("", ("Cell 3",), "Enter at least one"),
    ("T1 Cell 3", ("Cell 3",), "TC name = ROI name"),
])
def test_pair_text_errors(text, rois, error):
    _pairs, errors = excel_dialog.parse_tc_pairs(text, ("T1", "T2"), rois)
    assert error in " ".join(errors)


@pytest.fixture
def isolated_settings(tmp_path, monkeypatch):
    monkeypatch.setenv("FLIR_SETTINGS_FILE", str(tmp_path / "settings.ini"))
    return app_settings()


@pytest.fixture
def lab_file(tmp_path, monkeypatch):
    # Read a real xlsx on the worker; the LogFile clock parser has its own tests in test_v073.
    path = tmp_path / "tc.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.active.append(["Test notes"])
    sheet = book.create_sheet("TC")
    sheet.append(["Time (s)", "T1", "T2", "T3"])
    for t in range(40):
        sheet.append([t, 60 + t, 70 + t, 80 + t])
    book.save(path)
    book.close()
    original = tc_dialog.read_tc_table

    def read(*args, **kwargs):
        return replace(original(*args, **kwargs), clock_zero=START + timedelta(seconds=12.3))

    monkeypatch.setattr(tc_dialog, "read_tc_table", read)
    return path


def metadata(tmp_path, **changes):
    return replace(VideoMetadata(tmp_path / "recording.seq", 30, 20, 1000, START, None, 100, 10), **changes)


def zones(prefix="Cell"):
    return tuple(RoiShape(k, "rect", ((k * 2.0, 0.0), (k * 2.0 + 2, 5.0)), f"{prefix} {k}")
                 for k in (3, 6, 9))


@pytest.fixture
def make_dialog(qapp, tmp_path, isolated_settings):
    dialogs = []

    def make(**kwargs):
        dialog = ExcelExportDialog(kwargs.pop("meta", metadata(tmp_path)), kwargs.pop("rois", zones()),
                                   current_frame=25, **kwargs)
        dialogs.append(dialog)
        return dialog

    yield make
    for dialog in dialogs:
        dialog.close()
        dialog.deleteLater()
    assert wait_until(qapp, lambda: not tc_dialog._ALIVE)


def load(qapp, dialog, path):
    dialog.tc_file_edit.setText(str(path))
    dialog._load_tc(sheet=None)
    assert wait_until(qapp, lambda: not dialog._tc_reading)
    assert dialog._tc_table is not None, dialog.tc_label.text()


def test_dialog_file_default_clock_time_pairs_and_manual_ignition(qapp, make_dialog, lab_file):
    dialog = make_dialog()
    assert dialog.tc_source_combo.currentData() == "file" and dialog.tc_source_combo.count() == 1
    load(qapp, dialog, lab_file)
    assert dialog.tc_sheet_combo.currentText() == "TC"
    assert dialog.ignition_spin.value() == 124
    assert dialog.start_check.isChecked() and dialog.end_check.isChecked()
    assert (dialog.start_spin.value(), dialog.end_spin.value()) == (0.0, 39.0)
    before = dialog.options().tc_prefill
    assert before.offset_s == pytest.approx(0)
    assert dict(before.roi_tc) == {"Cell 3": "T1", "Cell 6": "T2", "Cell 9": "T3"}
    assert "Clocks: logger time 0 at frame 124 (0 min 12 s into the recording)" in dialog.tc_label.text()
    assert "Time 0 is logger time 0" in dialog.tc_label.text()
    dialog.ignition_spin.setValue(144)
    after = dialog.options().tc_prefill
    assert after.offset_s == pytest.approx(-2)
    assert after.roi_tc == before.roi_tc and after.columns == before.columns
    assert "Time 0 (frame 144) is logger time 2.0 s" in dialog.tc_label.text()
    dialog._update()
    assert dialog.ignition_spin.value() == 144


def test_dialog_fit_default_and_file_prefilled(make_dialog, lab_file):
    data = table()
    channels = [SimpleNamespace(name="T1", match=SimpleNamespace(found=True, pixel=(1, 7)))]
    result = SimpleNamespace(channels=channels, preset=None, fps=10, lag_s=5, ignition_frame=50,
                             num_frames=1000, origin_s=None, step=1)
    # The real fit prefill is covered by test_v071; this stub only avoids constructing a fit for source choice.
    channels[0].eps = None
    channels[0].events = []
    dialog = make_dialog(tc_run=RunOutput(result=result, table=data, tc_file=lab_file))
    assert dialog.tc_source_combo.currentData() == "fit" and dialog.tc_source_combo.count() == 2
    assert dialog.tc_file_edit.text() == str(lab_file)
    assert dialog.ignition_spin.value() == 51 and dialog.options().tc_prefill is not None
    channels[0].match.found = False
    other = make_dialog(tc_run=RunOutput(result=result, table=data, tc_file=lab_file))
    assert other.tc_source_combo.currentData() == "file"


@pytest.mark.parametrize("kind,why", [("logger", "logger's clock"), ("camera", "recording has no clock"),
                                       ("outside", "outside this recording")])
def test_dialog_without_clock_uses_manual_ignition(qapp, make_dialog, lab_file, monkeypatch, tmp_path, kind, why):
    data = replace(table(), clock_zero=None) if kind == "logger" else table(lag=1000 if kind == "outside" else 12)
    monkeypatch.setattr(tc_dialog, "read_tc_table", lambda *a, **k: data)
    dialog = make_dialog(meta=metadata(tmp_path, start_time=None if kind == "camera" else START))
    load(qapp, dialog, lab_file)
    assert dialog.ignition_spin.value() == 26
    assert dialog.options().tc_prefill.offset_s == 0
    assert why in dialog.tc_label.text() and "taken as logger time 0" in dialog.tc_label.text()
    dialog.ignition_spin.setValue(101)
    assert dialog.options().tc_prefill.offset_s == 0


def test_dialog_clamps_clock_before_recording(qapp, make_dialog, lab_file, monkeypatch):
    monkeypatch.setattr(tc_dialog, "read_tc_table", lambda *a, **k: table(lag=-5))
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    assert dialog.ignition_spin.value() == 1
    assert (dialog.start_spin.value(), dialog.end_spin.value()) == (0, 34)
    assert dialog.options().tc_prefill.offset_s == -5
    assert "Time 0 (frame 1) is logger time 5.0 s" in dialog.tc_label.text()


def test_a700_note_and_actual_zone_prefix(qapp, make_dialog, lab_file, tmp_path):
    dialog = make_dialog(meta=metadata(tmp_path, camera_model="FLIR A700"), rois=zones("Module"))
    load(qapp, dialog, lab_file)
    assert dict(dialog.options().tc_prefill.roi_tc) == {"Module 3": "T1", "Module 6": "T2", "Module 9": "T3"}
    assert "A700" in dialog.tc_label.text() and "moved by hand" in dialog.tc_label.text()
    assert dialog.ignition_spin.value() == 124  # no hidden 40 s correction


def test_ambiguous_zone_numbers_need_explicit_pairs(qapp, make_dialog, lab_file):
    dialog = make_dialog(rois=zones("Front") + zones("Back"))
    load(qapp, dialog, lab_file)
    assert dialog.tc_pairs_edit.text() == "" and dialog.options().tc_prefill is None
    dialog.tc_pairs_edit.setText("T1 = Back 3")
    assert dialog.options().tc_prefill.roi_tc == (("Back 3", "T1"),)


def test_unchecking_tc_during_read_preserves_manual_time(qapp, make_dialog, lab_file, monkeypatch):
    gate, started = Event(), Event()

    def read(*args, **kwargs):
        started.set()
        assert gate.wait(5)
        return table()

    monkeypatch.setattr(tc_dialog, "read_tc_table", read)
    dialog = make_dialog()
    dialog.tc_file_edit.setText(str(lab_file))
    dialog._load_tc(sheet=None)
    try:
        assert wait_until(qapp, started.is_set)
        dialog.tc_check.setChecked(False)
        dialog.ignition_spin.setValue(101)
    finally:
        gate.set()
    assert wait_until(qapp, lambda: not dialog._tc_reading)
    assert dialog.ignition_spin.value() == 101 and not dialog.start_check.isChecked()
    assert dialog.options().tc_prefill is None
    dialog.tc_check.setChecked(True)
    assert dialog.ignition_spin.value() == 101
    assert dialog.options().tc_prefill.offset_s == pytest.approx(2.3)


@pytest.mark.parametrize("refusal", ["presets", "clock", "sheet", "duplicate", "added", "pairs", "cap"])
def test_dialog_refusals(qapp, make_dialog, lab_file, tmp_path, refusal, monkeypatch):
    if refusal == "cap":
        monkeypatch.setattr(tc_dialog, "read_tc_table", lambda *a, **k: table(tuple(f"T{k}" for k in range(11))))
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    if refusal == "presets":
        dialog.metadata = replace(dialog.metadata, presets=(PresetRange(1, 500), PresetRange(2, 500)), nominal_fps=0)
    elif refusal == "clock":
        dialog.time_base_combo.setCurrentIndex(1)
    elif refusal == "sheet":
        dialog.sheet_checks["tc"].setChecked(False)
    elif refusal == "duplicate":
        dialog.rois += (replace(dialog.rois[0], id=99),)
    elif refusal == "added":
        dialog._others.append({"path": tmp_path / "other.seq", "rois": (dialog.rois[0], dialog.rois[0])})
    elif refusal == "pairs":
        dialog.tc_pairs_edit.setText("T1 = Missing")
    else:
        dialog.rois = tuple(replace(zones()[0], id=k, name=f"Cell {k}") for k in range(11))
        dialog.tc_pairs_edit.setText(", ".join(f"T{k} = Cell {k}" for k in range(11)))
    dialog._update()
    assert dialog.options().tc_prefill is None
    assert dialog.tc_label.text()
    if refusal == "cap":
        assert "10" in dialog.tc_label.text()
    if refusal == "presets":  # 0.9.0 fills a superframing workbook too: what stops it here is the frame rate
        assert "no frame rate" in dialog.tc_label.text()


def test_settings_only_after_accept(qapp, make_dialog, lab_file, isolated_settings):
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    custom = "T2 = Cell 3; T1 = Cell 6"
    dialog.tc_pairs_edit.setText(custom)
    dialog.reject()
    assert not isolated_settings.contains(excel_dialog._TC_FILE_KEY)
    assert not isolated_settings.contains(excel_dialog._TC_PAIRS_KEY)
    dialog.accept()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert excel_dialog.remembered_tc_file(dialog.metadata.path) == str(lab_file)
    assert isolated_settings.value(excel_dialog._TC_PAIRS_KEY) == custom
    again = make_dialog()
    assert wait_until(qapp, lambda: not again._tc_reading)
    assert again.tc_file_edit.text() == str(lab_file) and again.tc_pairs_edit.text() == custom
    isolated_settings.setValue(excel_dialog._TC_PAIRS_KEY, "Missing = Cell 3")
    fallback = make_dialog()
    assert wait_until(qapp, lambda: not fallback._tc_reading)
    assert fallback.tc_pairs_edit.text() == "T1 = Cell 3, T2 = Cell 6, T3 = Cell 9"
    excel_dialog.remember_tc_file(dialog.metadata.path, lab_file.with_name("missing.xlsx"))
    assert make_dialog().tc_file_edit.text() == ""
    # another recording does not get this recording's TC file (another test's clocks could time it)
    excel_dialog.remember_tc_file(dialog.metadata.path, lab_file)
    assert excel_dialog.remembered_tc_file(dialog.metadata.path.with_name("other.csq")) == ""


def test_export_cannot_overwrite_its_tc_file(qapp, make_dialog, lab_file, monkeypatch, isolated_settings):
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    dialog.path_edit.setText(str(lab_file))
    before = lab_file.read_bytes()
    messages = []
    monkeypatch.setattr(excel_dialog.MessageDialog, "warning", lambda *a: messages.append(a[-1]))
    monkeypatch.setattr(excel_dialog, "confirm_replace", lambda *a: [lab_file])
    dialog.accept()
    assert dialog.result() != QDialog.DialogCode.Accepted
    assert messages and "TC file" in messages[0]
    assert not isolated_settings.contains(excel_dialog._TC_FILE_KEY)
    assert lab_file.read_bytes() == before


def test_declining_replace_does_not_remember(qapp, make_dialog, lab_file, monkeypatch, isolated_settings):
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    monkeypatch.setattr(excel_dialog, "confirm_replace", lambda *a: None)
    dialog.accept()
    assert dialog.result() != QDialog.DialogCode.Accepted
    assert not isolated_settings.contains(excel_dialog._TC_FILE_KEY)


def test_read_error_stays_usable_and_sheet_can_be_changed(qapp, make_dialog, lab_file):
    dialog = make_dialog()
    load(qapp, dialog, lab_file)
    dialog.tc_sheet_combo.setCurrentText("Notes")
    assert wait_until(qapp, lambda: not dialog._tc_reading)
    assert dialog._tc_table is None and dialog.options().tc_prefill is None
    assert "Could not read TC data" in dialog.tc_label.text()
    assert dialog.button_box.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
    dialog.tc_sheet_combo.setCurrentText("TC")
    assert wait_until(qapp, lambda: not dialog._tc_reading)
    assert dialog.options().tc_prefill is not None


def test_read_in_progress_and_stale_result(qapp, make_dialog, lab_file, monkeypatch):
    gate, started = Event(), Event()

    def read(*args, **kwargs):
        started.set()
        assert gate.wait(5)
        return table()

    monkeypatch.setattr(tc_dialog, "read_tc_table", read)
    dialog = make_dialog()
    dialog.tc_file_edit.setText(str(lab_file))
    dialog._load_tc(sheet=None)
    try:
        assert wait_until(qapp, started.is_set)
        assert "Reading" in dialog.tc_label.text()
        assert not dialog.button_box.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
        dialog.tc_check.setChecked(False)
        assert dialog.button_box.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
        dialog.tc_check.setChecked(True)
        dialog.tc_file_edit.setText(str(lab_file.with_name("missing.xlsx")))
        dialog._tc_file_edited()
    finally:
        gate.set()
    assert wait_until(qapp, lambda: not tc_dialog._ALIVE)
    assert dialog._tc_table is None and dialog.options().tc_prefill is None
    assert "No such file" in dialog.tc_label.text()
    assert dialog.ignition_spin.value() == 26


def test_reader_outlives_closed_dialog(qapp, lab_file, tmp_path, isolated_settings, monkeypatch):
    gate, started = Event(), Event()

    def read(*args, **kwargs):
        started.set()
        assert gate.wait(5)
        return table()

    monkeypatch.setattr(tc_dialog, "read_tc_table", read)
    dialog = ExcelExportDialog(metadata(tmp_path), zones(), current_frame=0)
    dialog.tc_file_edit.setText(str(lab_file))
    dialog._load_tc(sheet=None)
    try:
        assert wait_until(qapp, started.is_set)
        dialog.reject()
        dialog.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        assert tc_dialog._ALIVE
    finally:
        gate.set()
    assert wait_until(qapp, lambda: not tc_dialog._ALIVE)


def test_clock_prefill_in_exported_sample_workbook(tmp_path):
    pytest.importorskip("fnv")
    from flir_player.excel_export import ExportOptions, SourceSpec, run_export

    recording = Path(__file__).resolve().parents[1] / "local" / "data" / "2.seq"
    if not recording.exists():
        pytest.skip("sample recording '2.seq' is not present in local/data/")
    data = table(("T1", "T2", "T3"))
    prefill = tcmatch.clock_prefill(data, lag_s=0.2, fps=30, frame=0, pairs={"T3": "Cell 9"})
    path = tmp_path / "book.xlsx"
    run_export(path, [SourceSpec(recording, (zones()[2],))],
               ExportOptions(start_s=0, end_s=1, step_s=0.1, tc_prefill=prefill))
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        assert book["TC Compare"]["B6"].value == pytest.approx(0.2)
        settings = book["Settings"]
        row = next(r for r in range(1, settings.max_row + 1) if settings.cell(r, 3).value == "Cell 9")
        assert settings.cell(row, 8).value == "T3"
    finally:
        book.close()


def test_clock_prefill_written_values_and_chosen_tc_without_sdk(tmp_path):
    from test_excel_workbook import _source
    from flir_player.excel_export import ExportData, ExportOptions
    from flir_player.workbook import TC_FIRST, write_workbook

    source = _source("Camera", 12, np.random.default_rng(8))
    prefill = tcmatch.clock_prefill(table(), lag_s=12.3, fps=10, frame=143,
                                   pairs={"T3": "TC1", "T1": "TC2"})
    data = ExportData(times=np.arange(12.0), sources=[source], options=ExportOptions(tc_prefill=prefill),
                      created=START, versions={"tool": "test", "sdk": ""})
    path = tmp_path / "book.xlsx"
    write_workbook(path, data)
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        tc, settings = book["TC Compare"], book["Settings"]
        assert tc["B6"].value == pytest.approx(-2)
        assert tc.cell(TC_FIRST + 1, 2).value == "T3"
        assert tc.cell(TC_FIRST + 1, 3).value == "T1"
        for name, chosen in (("TC1", "T3"), ("TC2", "T1")):
            row = next(r for r in range(1, settings.max_row + 1) if settings.cell(r, 3).value == name)
            assert settings.cell(row, 8).value == chosen
            assert settings.cell(row, 6).value is None  # no fitted emissivity override
        # Logger 2 is workbook 0: its raw T3 and T1 readings stay in the chosen columns.
        # Comparison formulas recalculate in Excel; their stored caches are deliberately #N/A.
        assert [tc.cell(TC_FIRST + 4, c).value for c in (1, 2, 3)] == [2, 64, 62]
    finally:
        book.close()


def test_default_pairs_prefer_the_cell_zones_over_other_rois_with_the_same_number():
    names = ("T1", "T2", "T3")
    rois = ("Box 3", "Cell 3", "Cell 6", "Cell 9")
    assert excel_dialog._default_tc_pairs(names, rois, ("Cell 3", "Cell 6", "Cell 9")) == (
        "T1 = Cell 3, T2 = Cell 6, T3 = Cell 9")
    # without zones, an ambiguous number pairs nothing
    assert excel_dialog._default_tc_pairs(names, rois) == "T2 = Cell 6, T3 = Cell 9"
