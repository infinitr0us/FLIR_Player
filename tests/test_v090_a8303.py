# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""0.9.0: A8303sc recordings (superframing ATS) with temperatures.

A fake handle follows what the 1006a A8303 file showed (local/notes/2026-10-18-a8303-answers): counts →
radiance by a quadratic; the SDK's radiance unit has the window and atmosphere taken out; temperatures follow
R / (exp(B/T) − 1) of the radiance, with the object parameters applied in radiance space and the camera's own
transmission constants; frames cycle through presets.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace as NS

import numpy as np
import openpyxl
import pytest

from flir_player.calibration import RADIANCE_SCALE, derive_calibration, sdk_values
from flir_player.excel_export import ExportOptions, SourceSpec, TcPrefill, run_export
from flir_player.models import RoiShape
from flir_player.radiometry import (
    UNIT_TRANSMISSION,
    AtmosphereConstants,
    Calibration,
    MeasurementParameters,
    PlanckConstants,
    RangeLimits,
    compensated_signal,
    fit_atmosphere,
    fit_radiance_curve,
    object_temperature,
)

A8303_ATMOSPHERE = AtmosphereConstants(X=0.0932162, alpha1=0.2014372, alpha2=0.00506231, beta1=0.00801709,
                                       beta2=0.00281940)
FILE_PARAMS = MeasurementParameters(emissivity=0.95, transmission=None, distance_m=3.3, humidity=0.3)
# per preset (1, 2): calibrated range (K), radiance curve (SDK units) and counts → radiance quadratic
PRESETS = {1: ((523.15, 873.15), PlanckConstants(19.104868, 3632.4014, 1.0, 0.0), (2.252e-10, 2.1435e-05, -0.04248)),
           2: ((773.15, 1473.15), PlanckConstants(22.873647, 3788.5186, 1.0, 0.0), (1.1e-10, 1.9e-05, -0.031))}


class _A8303Like:
    """A superframing ATS handle: frame i has preset 1 + i % 2 (or ``preset_of(i)``), 2 frames per second."""

    num_frames = 120

    def __init__(self, scene=None, clock_gap: float = 0.0, size: int = 8, preset_of=None,
                 transmission=None, clamped: bool = False):
        import fnv
        self.fnv = fnv
        self.width = self.height = size
        self.unit, self.temp_type = fnv.Unit.COUNTS, fnv.TempType.KELVIN
        self.supported_units = [fnv.Unit.COUNTS, fnv.Unit.RADIANCE_FACTORY, fnv.Unit.TEMPERATURE_FACTORY]
        self.reset_object_parameters()
        self.preset_of = preset_of or (lambda i: 1 + i % 2)
        info = [NS(available=False, calibrated=False, num_frames=0, min_temp=0.0, max_temp=0.0)]
        for preset, ((low, high), _curve, _quad) in PRESETS.items():
            frames = sum(1 for i in range(self.num_frames) if self.preset_of(i) == preset)
            info.append(NS(available=True, calibrated=True, min_temp=low, max_temp=high, frame_rate_valid=True,
                           frame_rate=54.2, num_frames=frames))
        self.source_info = NS(camera_model="A8303sc", preset_info=info)
        # each pixel's scene temperature (K) per frame: a module heating up, a little different per pixel
        self.scene = scene if scene is not None else (
            lambda i: np.linspace(560.0, 840.0, size * size).reshape(size, size) + 30.0 * np.sin(i / 7.0))
        self.clock_gap = clock_gap  # seconds the clock jumps after frame 60 (dropped frames)
        self.transmission = transmission  # (distance, humidity) → τ: a camera unlike FLIR's model
        self.clamped = clamped  # the SDK's status: every pixel clamped (a flame preset over a cold scene)
        self.get_frame(0)

    def get_preset(self, index):
        return self.preset_of(index)

    def reset_object_parameters(self):
        self.object_parameters = NS(**sdk_values(FILE_PARAMS), atmospheric_transmission=1.0)

    def _curve(self, preset):
        _range, curve, _quad = PRESETS[preset]
        return Calibration(planck=curve, atmosphere=A8303_ATMOSPHERE)

    def counts_for(self, index):
        _range, curve, (a2, a1, a0) = PRESETS[self.get_preset(index)]
        radiance = curve.signal(self.scene(index))  # what reaches the camera (window and air left out)
        return (-a1 + np.sqrt(a1 * a1 - 4 * a2 * (a0 - radiance))) / (2 * a2)

    def get_frame(self, index):
        fnv = self.fnv
        self.index = index
        seconds = index / 2.0 + (self.clock_gap if index >= 60 else 0.0)
        self.frame_info = NS(time=datetime(2026, 10, 6, 10, 47, 59) + timedelta(seconds=seconds))
        preset = self.get_preset(index)
        counts = self.counts_for(index)
        a2, a1, a0 = PRESETS[preset][2]
        radiance = a2 * counts * counts + a1 * counts + a0
        params = MeasurementParameters.from_sdk(vars(self.object_parameters))
        if self.transmission is not None and params.transmission is None and params.distance_m > 0:
            params = params.with_(transmission=self.transmission(params.distance_m, params.humidity))
        cal = self._curve(preset)
        if self.unit == fnv.Unit.COUNTS:
            self.final = counts
        elif self.unit == fnv.Unit.RADIANCE_FACTORY:
            self.final = compensated_signal(cal, params, radiance)
        else:
            self.final = object_temperature(cal, params, radiance)
        self.status = np.full((self.height, self.width), 2 if self.clamped else 0, np.int8)

    def update_frame(self):
        self.get_frame(self.index)

    def has_unit(self, unit):
        return unit in self.supported_units

    def close(self):
        pass


@pytest.fixture
def ats(tmp_path) -> Path:
    path = tmp_path / "1006_A8303sc.ats"
    path.write_bytes(b"synthetic, no FFF record")
    return path


# --- radiometry -------------------------------------------------------------------------------------


def test_the_atmosphere_fit_recovers_a_cameras_constants() -> None:
    grid = [(d, h, t) for d in (1.0, 3.0, 10.0, 30.0, 100.0) for h in (0.0, 0.5, 0.9) for t in (283.15, 303.15)]
    for constants in (A8303_ATMOSPHERE, AtmosphereConstants()):
        tau = [constants.transmission(*point) for point in grid]
        fitted, error = fit_atmosphere(*zip(*grid), tau)
        assert error < 1e-7
        for d, h, t in ((3.3, 0.3, 293.15), (50.0, 0.7, 298.15), (0.5, 0.0, 280.0)):
            assert fitted.transmission(d, h, t) == pytest.approx(constants.transmission(d, h, t), abs=1e-6)
    unit, error = fit_atmosphere(*zip(*grid), [1.0] * len(grid))
    assert unit == UNIT_TRANSMISSION and error == 0.0


def test_the_radiance_curve_has_two_constants() -> None:
    curve = PlanckConstants(1.9e6, 3632.4, 1.0, 0.0)
    kelvin = np.linspace(520.0, 900.0, 200)
    fitted = fit_radiance_curve(kelvin, curve.signal(kelvin))
    assert (fitted.F, fitted.O) == (1.0, 0.0)
    assert fitted.R == pytest.approx(curve.R, rel=1e-9) and fitted.B == pytest.approx(curve.B, rel=1e-9)
    # exact far below the fitted temperatures (the reflected and air terms are evaluated there)
    assert float(fitted.signal(293.15)) == pytest.approx(float(curve.signal(293.15)), rel=1e-7)


# --- calibration -------------------------------------------------------------------------------------


def test_the_radiance_route_verifies_an_a8303_preset(ats) -> None:
    pytest.importorskip("fnv")
    im = _A8303Like()
    frames = list(range(0, 120, 2))[::6]  # preset 1
    report = derive_calibration(im, ats, indices=frames, limits=RangeLimits(calibrated=PRESETS[1][0]))
    assert report.verified, report.message
    assert report.signal == "radiance" and "radiance" in report.message
    assert max(error for _label, error, _n in report.checks) < 1e-3
    assert any("own parameters" in label for label, _e, _n in report.checks)
    counts = im.counts_for(0)
    a2, a1, a0 = PRESETS[1][2]
    assert np.allclose(report.to_signal(counts), (a2 * counts ** 2 + a1 * counts + a0) * RADIANCE_SCALE, rtol=1e-6)
    tau = report.calibration.atmosphere.transmission(3.3, 0.3, 293.15)
    assert tau == pytest.approx(A8303_ATMOSPHERE.transmission(3.3, 0.3, 293.15), abs=1e-6)
    # temperatures at the recording's own parameters are the SDK's
    im.reset_object_parameters()
    im.unit = im.fnv.Unit.TEMPERATURE_FACTORY
    im.get_frame(4)
    ours = object_temperature(report.calibration, FILE_PARAMS, report.to_signal(im.counts_for(4)))
    assert np.max(np.abs(ours - im.final)) < 1e-3
    # without a range the counts route is tried first and the report says why it failed too
    assert derive_calibration(_A8303Like(), ats, indices=frames).signal == "radiance"
    assert derive_calibration(_A8303Like(), ats, indices=frames, signal="counts").signal == "counts"


def test_a_transmission_unlike_flirs_model_is_not_verified(ats) -> None:
    """Review 0.9.0: the transmission fit is checked on its own, not only through two of its points."""
    pytest.importorskip("fnv")
    frames = list(range(0, 120, 2))[::6]
    steps = (lambda _distance, humidity: 0.97 if humidity < 0.3 else 0.85)  # humidity steps, no distance
    report = derive_calibration(_A8303Like(transmission=steps), ats, indices=frames,
                                limits=RangeLimits(calibrated=PRESETS[1][0]))
    assert not report.verified and "transmission" in report.message


# --- export --------------------------------------------------------------------------------------------


def _export(tmp_path, ats, im, *, preset=None, prefill=None, ignition=10, options=None) -> openpyxl.Workbook:
    out = tmp_path / f"a8303_{preset}.xlsx"
    spec = SourceSpec(ats, (RoiShape(1, "cursor", ((3.5, 3.5),), "Cell 1"),
                            RoiShape(2, "rect", ((1.0, 1.0), (4.0, 4.0)), "Cell 2")),
                      ignition_frame=ignition, preset=preset)
    options = options or ExportOptions(step_s=1.0, tc_prefill=prefill)
    message = run_export(out, [spec], options, handles={ats.resolve(): im}, replace=[out])
    assert "1/1 recording(s) with live temperatures" in message
    return openpyxl.load_workbook(out, data_only=True)


def test_the_export_gives_an_a8303_preset_temperatures(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    book = _export(tmp_path, ats, _A8303Like(), preset=2)
    frames = [row[1] for row in book["Counts"].iter_rows(min_row=6, values_only=True) if row[1] is not None]
    assert frames and all(1 + (frame - 1) % 2 == 2 for frame in frames)  # preset 2 only
    largest = next(row for row in book["Validation"].iter_rows(values_only=True)
                   if row[0] == "Largest |difference| over checked samples (K)")
    assert largest[3] < 1e-3 and largest[4].startswith("PASS")
    headers = list(book["Counts"].iter_rows(min_row=5, max_row=5, values_only=True))[0]
    assert "Value (signal)" in headers and "Mean (signal)" in headers
    source = {row[0]: row[1] for row in book["Source"].iter_rows(values_only=True) if row[0]}
    assert source["Preset"].startswith("2 (500 to 1200 °C)") and "quadratic of the counts" in source["Signal"]
    assert list(book["Settings"].iter_rows(min_row=5, max_row=5, values_only=True))[0][1] == "Leave them blank (#N/A)"


def test_a_validation_that_checked_nothing_says_so(tmp_path, ats) -> None:
    """Review 0.9.0: with every sample clamped by the SDK the sheet said "0.0 PASS"."""
    pytest.importorskip("fnv")
    book = _export(tmp_path, ats, _A8303Like(clamped=True), preset=1)
    largest = next(row for row in book["Validation"].iter_rows(values_only=True)
                   if row[0] == "Largest |difference| over checked samples (K)")
    assert largest[3] == "–" and largest[4].startswith("NOT CHECKED")


def test_out_of_range_rows_start_blank(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    cold = _A8303Like(scene=lambda i: np.full((8, 8), 450.0 if i < 60 else 700.0) + np.arange(64).reshape(8, 8))
    book = _export(tmp_path, ats, cold, preset=1, ignition=0)
    rows = [row for row in book["Data"].iter_rows(min_row=6, values_only=True) if row[0] is not None]
    early = [row for row in rows if row[0] < 25]
    late = [row for row in rows if row[0] > 35]
    assert early and all(row[2] == "#N/A" for row in early)  # 177 °C, below preset 1's 250 °C
    assert late and all(isinstance(row[2], float) for row in late)
    # review 0.9.0: the Summary's values (as written) leave the blank rows out too
    summary = next(row for row in book["Summary"].iter_rows(values_only=True) if row[1] == "Cell 1")
    assert summary[3] >= 250.0 and summary[4] == pytest.approx(max(row[2] for row in late))


def test_a_frame_timed_tc_sheet_moves_onto_the_clock(tmp_path, ats) -> None:
    """The dialog places the TC readings by frame number / rate; the recording's clock gains 4 s after
    frame 60 (dropped frames), so at ignition frame 80 the axis is 4 s later than 80 / 2 fps."""
    pytest.importorskip("fnv")
    prefill = TcPrefill(times=(0.0, 1.0, 2.0), columns=((20.0, 21.0, 22.0),), names=("T1",),
                        roi_tc=(("Cell 1", "T1"),), offset_s=-30.0, window=(-30.0, -28.0),
                        roi_use=(("Cell 1", ((-30.0, -29.0),)),), frame=80, fps=2.0)
    options = ExportOptions(step_s=1.0, start_s=10.0, end_s=20.0, tc_prefill=prefill)
    book = _export(tmp_path, ats, _A8303Like(clock_gap=4.0), preset=1, ignition=80, options=options)
    offset = next(row[1] for row in book["TC Compare"].iter_rows(values_only=True)
                  if row[0] == "Logger time offset (s)")
    assert offset == pytest.approx(-30.0 + 80 / 2.0 - (80 / 2.0 + 4.0))
    # review 0.9.0: From/To placed the same way (e.g. the logger's span) move with the TC readings
    times = [row[0] for row in book["Data"].iter_rows(min_row=6, values_only=True) if row[0] is not None]
    assert (times[0], times[-1]) == (6, 16)


def test_rounded_tc_bounds_keep_their_end_seconds_on_the_clock(tmp_path, ats) -> None:
    """Review 0.9.0 (second round): the dialog rounds the TC span to 0.1 s on the frame axis; on the clock the
    span must keep its first and last logger second."""
    pytest.importorskip("fnv")
    fps = 1.874015748  # 80 / fps = 42.689076 s, while frame 80 is at clock second 44: delta = -1.310924 s
    offset = 44.0 - 80 / fps
    times = tuple(float(t) for t in range(10))
    prefill = TcPrefill(times=times, columns=(tuple(20.0 + t for t in times),), names=("T1",),
                        roi_tc=(("Cell 1", "T1"),), offset_s=offset, frame=80, fps=fps)
    options = ExportOptions(step_s=1.0, start_s=round(times[0] + offset, 1), end_s=round(times[-1] + offset, 1),
                            tc_prefill=prefill)
    book = _export(tmp_path, ats, _A8303Like(clock_gap=4.0), preset=1, ignition=80, options=options)
    rows = [row[0] for row in book["Data"].iter_rows(min_row=6, values_only=True) if row[0] is not None]
    assert (rows[0], rows[-1], len(rows)) == (0, 9, 10)


def test_superframing_areas_count_for_the_pixel_sheet_budget(tmp_path, ats) -> None:
    """Review 0.9.0: 42 areas of 400 pixels need more columns than a sheet has; the limit is lowered."""
    pytest.importorskip("fnv")
    rois = tuple(RoiShape(k + 1, "rect", ((x, y), (x + 20, y + 20)), f"Cell {k + 1}")
                 for k, (x, y) in enumerate((x, y) for y in range(0, 120, 20) for x in range(0, 140, 20)))
    out = tmp_path / "many.xlsx"
    spec = SourceSpec(ats, rois, ignition_frame=10, preset=1)
    message = run_export(out, [spec], ExportOptions(step_s=10.0), handles={ats.resolve(): _A8303Like(size=140)})
    assert "42 ROIs" in message
    book = openpyxl.load_workbook(out, read_only=True, data_only=True)
    source = {row[0]: row[1] for row in book["Source"].iter_rows(values_only=True) if row[0]}
    assert source["Pixel limit for exact means"] < 400


# --- player ----------------------------------------------------------------------------------------------


def test_the_player_reads_one_preset(ats, monkeypatch) -> None:
    pytest.importorskip("fnv")
    from flir_player import source as source_module
    from flir_player.processing import ProcessingState

    fake = _A8303Like()
    monkeypatch.setattr(source_module.fnv.file, "ImagerFile", lambda _path: fake)
    fake.rois = []
    fake.can_change_object_parameters = True
    fake.has_nuc = fake.has_bp = False
    src = source_module.FlirVideoSource()
    metadata = src.open(ats)
    assert [p.index for p in metadata.presets] == [1, 2] and src.preset == 1
    assert [src.read_frame(i).index for i in (0, 1, 2, 3, 119)] == [0, 2, 2, 4, 118]
    src.set_preset(2)
    assert [src.read_frame(i).index for i in (0, 1, 118)] == [1, 1, 119]
    assert src.preset_frames(0, 7) == [1, 3, 5, 7]
    src.set_processing(ProcessingState(temporal=("mean", 3)))
    assert src._earlier(9, 2) == [5, 7] and src._following(9) == 11
    src.set_preset(None)
    assert src.read_frame(1).index == 1
    with pytest.raises(ValueError, match="Preset 3"):
        src.set_preset(3)


def test_the_player_never_shows_another_preset(ats, monkeypatch) -> None:
    """Review 0.9.0: preset 1 is missing over frames 40-70 (dropped); frame 0 is preset 2."""
    pytest.importorskip("fnv")
    from flir_player import source as source_module

    preset_of = (lambda i: 2 if i == 0 or 40 <= i <= 70 else 1 + i % 2)
    fake = _A8303Like(preset_of=preset_of)
    fake.rois = []
    fake.can_change_object_parameters = True
    fake.has_nuc = fake.has_bp = False
    monkeypatch.setattr(source_module.fnv.file, "ImagerFile", lambda _path: fake)
    src = source_module.FlirVideoSource()
    metadata = src.open(ats)
    assert src.preset == 1 and src.take_first_packet().index == 2
    assert metadata.start_time == datetime(2026, 10, 6, 10, 47, 59)  # frame 0's stamp, not frame 2's
    assert metadata.last_frame(1) == 118 and metadata.last_frame(2) == 119 and metadata.last_frame() == 119
    assert src.read_frame(50).index == 72 and src.read_frame(50, snap=-1).index == 38
    assert src.read_frame(72, snap=-1).index == 72 and src.read_frame(71, snap=-1).index == 38
    assert src.read_frame(119).index == 118  # past the preset's last frame: its last one


def test_series_exports_write_the_frames_the_dialog_checked() -> None:
    """Review 0.9.0: the frames written are among those checked for files to replace."""
    from flir_player.decoder import series_frames

    assert series_frames([1, 3, 5, 7], 0, 7, 1) == [1, 3, 5, 7]
    assert series_frames([1, 3, 5, 7], 0, 7, 2) == []
    assert series_frames([0, 3, 6, 9], 0, 9, 3) == [0, 3, 6, 9]
    assert series_frames(list(range(10)), 2, 9, 4) == [2, 6]


def test_the_movie_dialog_starts_at_the_presets_rate(qapp) -> None:
    from flir_player.export_dialogs import ExportMovieDialog
    from flir_player.models import VideoMetadata

    metadata = VideoMetadata(path=Path("a.ats"), width=8, height=8, num_frames=120, start_time=None,
                             end_time=None, duration_seconds=40.0, nominal_fps=3.0)
    assert ExportMovieDialog(metadata).fps_spin.value() == 3.0
    assert ExportMovieDialog(metadata, fps=1.0).fps_spin.value() == 1.0


# --- the player's transport on a preset (review 0.9.0, second round) ---------------------------------------


def _window_on(qapp, ats, monkeypatch, fake):
    from conftest import wait_until
    from flir_player import source as source_module
    from flir_player.main_window import MainWindow

    fake.rois = []
    fake.can_change_object_parameters = True
    fake.has_nuc = fake.has_bp = False
    monkeypatch.setattr(source_module.fnv.file, "ImagerFile", lambda _path: fake)
    window = MainWindow()
    window.show()
    window.open_path(ats)
    assert wait_until(qapp, lambda: window.current_packet is not None and not window._busy)
    return window


def test_steps_walk_the_presets_own_frames(qapp, ats, monkeypatch) -> None:
    """A dropped preset-2 frame makes frames 2 and 3 both preset 1: a step must not skip either."""
    pytest.importorskip("fnv")
    from conftest import wait_until

    pattern = [1, 2, 1, 1, 2, 1]
    window = _window_on(qapp, ats, monkeypatch, _A8303Like(preset_of=lambda i: pattern[i] if i < 6 else 1 + i % 2))
    try:
        assert window._preset == 1
        visited = []
        for move in (None, 1, 1, -1, -1, -1):
            if move is None:
                window.seek_to(2)
            else:
                window.step_frames(move)
            previous = window.current_packet
            assert wait_until(qapp, lambda: window.current_packet is not previous and not window._busy)
            visited.append(window.current_packet.index)
        assert visited == [2, 3, 5, 3, 2, 0]
    finally:
        window.close()
        qapp.processEvents()


def test_playback_stays_inside_a_range_whose_end_is_another_presets(qapp, ats, monkeypatch) -> None:
    pytest.importorskip("fnv")
    from conftest import wait_until
    from flir_player.settings import app_settings

    app_settings().setValue("playback/loop", False)
    window = _window_on(qapp, ats, monkeypatch, _A8303Like())  # preset 1 = even frames
    shown = []
    try:
        window._change_loop(False)
        window.transport.slider.set_play_range(0, 5)
        window.seek_to(2)
        assert wait_until(qapp, lambda: window.current_packet.index == 2)
        original = window._present_packet
        window._present_packet = lambda packet: (shown.append(packet.index), original(packet))
        window.toggle_playback()
        assert wait_until(qapp, lambda: not window.playing, timeout=10.0)
        assert window.current_packet.index == 4 and max(shown) == 4
    finally:
        window.close()
        qapp.processEvents()


def test_play_again_after_a_preset_range_end_starts_the_range_over(qapp, ats, monkeypatch) -> None:
    """Review 0.9.0 (third round): stopped at the preset's last frame in the range (4 of 0-5), Play restarts."""
    pytest.importorskip("fnv")
    from conftest import wait_until

    window = _window_on(qapp, ats, monkeypatch, _A8303Like())  # preset 1 = even frames
    shown = []
    try:
        window._change_loop(False)
        window.transport.slider.set_play_range(0, 5)
        window.seek_to(0)
        assert wait_until(qapp, lambda: window.current_packet.index == 0)
        original = window._present_packet
        window._present_packet = lambda packet: (shown.append(packet.index), original(packet))
        window.toggle_playback()
        assert wait_until(qapp, lambda: not window.playing and window.current_packet.index == 4, timeout=10.0)
        shown.clear()
        window.toggle_playback()
        assert wait_until(qapp, lambda: 0 in shown, timeout=10.0)
        assert wait_until(qapp, lambda: 4 in shown, timeout=10.0) and max(shown) == 4
    finally:
        window.pause_playback(invalidate=True)
        window.close()
        qapp.processEvents()


def test_a_range_without_the_presets_frames_does_not_loop(qapp, ats, monkeypatch) -> None:
    """Review 0.9.0 (third round): frames 1-3 are all preset 2; Play with preset 1 shown must not keep seeking."""
    pytest.importorskip("fnv")
    from conftest import wait_until

    window = _window_on(qapp, ats, monkeypatch, _A8303Like(preset_of=lambda i: 2 if i in (1, 2) else 1 + i % 2))
    requests = []
    try:
        window.transport.slider.set_play_range(1, 3)
        original = window.decoder.request_frame
        window.decoder.request_frame = lambda index, *a, **k: (requests.append(index), original(index, *a, **k))
        window.toggle_playback()
        assert wait_until(qapp, lambda: window.current_packet.index == 4)
        for _ in range(20):
            qapp.processEvents()
        assert requests == [1] and not window.playing
    finally:
        window.close()
        qapp.processEvents()


def test_validation_errors_stay_check(tmp_path) -> None:
    """Review 0.9.0 (third round): samples whose formula fails are CHECK, not "nothing checked"."""
    from flir_player.excel_export import ExportData
    from flir_player.workbook import write_workbook
    from test_excel_workbook import _source

    for case in ("all", "one"):
        source = _source("Check", 12, np.random.default_rng(7))
        source.rois = source.rois[:1]  # one spot, two validation samples
        rows = source.validation_rows if case == "all" else source.validation_rows[:1]
        source.rois[0].counts["value"][rows] = 0.0  # LN of a negative argument
        out = tmp_path / f"errors_{case}.xlsx"
        write_workbook(out, ExportData(times=np.arange(12, dtype=float), sources=[source], options=ExportOptions(),
                                       created=datetime(2026, 10, 18), versions={}))
        values = openpyxl.load_workbook(out, data_only=True)["Validation"]
        formulas = openpyxl.load_workbook(out)["Validation"]
        assert (values["D4"].value, values["E4"].value) == ("#NUM!", "CHECK")
        assert "ISERROR" in formulas["D4"].value and formulas["E4"].value.startswith('=IF(ISERROR(D4),"CHECK"')
