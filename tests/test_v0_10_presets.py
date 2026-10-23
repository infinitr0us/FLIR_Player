# SPDX-FileCopyrightText: 2026 Yuchuan Li <Yuchuan.Li@outlook.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""0.10.0: presets joined into one range, a preset that disagrees left out, and camera-shutter frames.

The fake camera follows the A8303 fire tests (local/notes/2026-10-23-a8303-flames): the presets of a cycle are
17 ms apart and a cycle starts every second; counts → radiance by a quadratic per preset; preset 3 of the 1006 test a
carried a count offset (its cold frames read 1,205 °C); the shutter covered the sensor for a while (FlagInFOV = T).
"""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace as NS

import numpy as np
import openpyxl
import pytest

from flir_player.calibration import sdk_values
from flir_player.excel_export import ExportOptions, SourceSpec, TcPrefill, pick_preset, run_export
from flir_player.models import RoiShape
from flir_player.presets import (as_preset, base_preset, by_range, check_presets, describe, is_joined,
                                 shutter_in_view)
from flir_player.radiometry import (AtmosphereConstants, Calibration, MeasurementParameters, PlanckConstants,
                                    compensated_signal, object_temperature)

ATMOSPHERE = AtmosphereConstants(X=0.0932162, alpha1=0.2014372, alpha2=0.00506231, beta1=0.00801709,
                                 beta2=0.00281940)
FILE_PARAMS = MeasurementParameters(emissivity=0.95, transmission=None, distance_m=3.3, humidity=0.3)
# per preset: calibrated range (K), radiance curve and counts → radiance quadratic (as the 1006a A8303)
PRESETS = {1: ((523.15, 873.15), PlanckConstants(19.104868, 3632.4014, 1.0, 0.0), (2.252e-10, 2.1435e-05, -0.04248)),
           2: ((773.15, 1473.15), PlanckConstants(22.873647, 3788.5186, 1.0, 0.0), (6.42e-11, 1.5056e-04, -0.24085)),
           3: ((1123.15, 2273.15), PlanckConstants(24.411681, 3873.2168, 1.0, 0.0), (5.54e-09, 3.568e-04, -0.48656))}


class _Info(list):
    """A frame header: iterable name/value entries, and the clock."""

    def __init__(self, time, preset, flag):
        super().__init__([{"name": "Preset", "value": str(preset)}, {"name": "FlagInFOV", "value": "T" if flag else "F"}])
        self.time = time


class _Superframing:
    """A superframing ATS handle: cycle k (one frame per preset, 17 ms apart) starts at k seconds."""

    def __init__(self, scene=None, presets=(1, 2), cycles=60, pedestal=None, shutter=(), size=8, covered=(),
                 clamp_above=None):
        import fnv
        self.fnv = fnv
        self.presets = tuple(presets)
        self.num_frames = cycles * len(self.presets)
        self.width = self.height = size
        self.unit, self.temp_type = fnv.Unit.COUNTS, fnv.TempType.KELVIN
        self.supported_units = [fnv.Unit.COUNTS, fnv.Unit.RADIANCE_FACTORY, fnv.Unit.TEMPERATURE_FACTORY]
        self.reset_object_parameters()
        info = [NS(available=False, calibrated=False, num_frames=0, min_temp=0.0, max_temp=0.0)]
        for preset, ((low, high), _curve, _quad) in PRESETS.items():
            used = preset in self.presets
            info.append(NS(available=used, calibrated=True, min_temp=low, max_temp=high, frame_rate_valid=True,
                           frame_rate=58.0, num_frames=cycles if used else 0))
        self.source_info = NS(camera_model="A8303sc", preset_info=info)
        # each pixel's scene temperature (K) per cycle: 500 K to 1,100 K across the image, slowly varying
        self.scene = scene or (lambda k: np.linspace(500.0, 1100.0, size * size).reshape(size, size)
                               + 20.0 * np.sin(k / 5.0))
        self.pedestal = pedestal or {}  # preset → counts added (the 1006a preset 3)
        self.shutter = set(shutter)  # cycles whose frames the shutter covered
        self.covered = set(covered)  # single frames the shutter covered (its edge falling inside a cycle)
        self.clamp_above = clamp_above  # K above a preset's range where the SDK clamps high (status 4)
        self.get_frame(0)

    def where(self, index):
        """(cycle, position in the cycle, preset) of frame ``index``."""
        cycle, member = divmod(index, len(self.presets))
        return cycle, member, self.presets[member]

    def get_preset(self, index):
        return self.where(index)[2]

    def reset_object_parameters(self):
        self.object_parameters = NS(**sdk_values(FILE_PARAMS), atmospheric_transmission=1.0)

    def counts_for(self, index):
        cycle, _member, preset = self.where(index)
        _range, curve, (a2, a1, a0) = PRESETS[preset]
        flag = cycle in self.shutter or index in self.covered
        scene = np.full((self.height, self.width), 290.0) if flag else self.scene(cycle)
        radiance = curve.signal(scene)
        counts = (-a1 + np.sqrt(a1 * a1 - 4 * a2 * (a0 - radiance))) / (2 * a2)
        return counts + self.pedestal.get(preset, 0.0)

    def get_frame(self, index):
        fnv = self.fnv
        self.index = index
        cycle, member, preset = self.where(index)
        time = datetime(2026, 10, 6, 10, 47, 59) + timedelta(seconds=cycle + 0.017 * member)
        self.frame_info = _Info(time, preset, cycle in self.shutter or index in self.covered)
        counts = self.counts_for(index)
        (low, high), curve, (a2, a1, a0) = PRESETS[preset]
        radiance = a2 * counts * counts + a1 * counts + a0
        params = MeasurementParameters.from_sdk(vars(self.object_parameters))
        cal = Calibration(planck=curve, atmosphere=ATMOSPHERE)
        kelvin = object_temperature(cal, params, radiance)
        if self.unit == fnv.Unit.COUNTS:
            self.final = counts
        elif self.unit == fnv.Unit.RADIANCE_FACTORY:
            self.final = compensated_signal(cal, params, radiance)
        else:
            self.final = kelvin
        self.status = np.where((kelvin >= low) & (kelvin <= high), 0, 1).astype(np.int8)
        if self.clamp_above is not None:
            self.status[kelvin > high + self.clamp_above] = 4

    def update_frame(self):
        self.get_frame(self.index)

    def has_unit(self, unit):
        return unit in self.supported_units

    def close(self):
        pass


@pytest.fixture
def ats(tmp_path) -> Path:
    path = tmp_path / "1006_A8303sc.ats"
    path.write_bytes(b"synthetic, no FFF record")
    return path


# --- the choice ------------------------------------------------------------------------------------------


def test_presets_joined_are_a_tuple_lowest_first() -> None:
    assert base_preset((1, 2)) == 1 and base_preset(2) == 2 and base_preset(None) is None
    assert is_joined((1, 2)) and not is_joined(1) and not is_joined((1,))
    assert as_preset([1, 2]) == (1, 2) and as_preset([2]) == 2 and as_preset(None) is None
    ranges = {p: r for p, (r, _c, _q) in PRESETS.items()}
    assert by_range((3, 1, 2), ranges) == (1, 2, 3)
    assert describe((1, 2), ranges) == "Presets 1 and 2 joined (250 to 1200 °C)"
    assert describe(2) == "Preset 2"


def test_an_area_takes_the_lowest_preset_that_holds_it() -> None:
    assert pick_preset([0, 0]) == 0  # inside both: the lowest
    assert pick_preset([2, 0]) == 1  # above preset 1: preset 2
    assert pick_preset([3, 0]) == 1  # saturated in preset 1
    assert pick_preset([2, 1]) == 0  # spans the gap: preset 1 extrapolates upward, preset 2 has a floor below
    assert pick_preset([3, 1]) == 1
    assert pick_preset([1, 1]) == 0  # cold: the lowest, below its range
    assert pick_preset([3, 3]) == 1  # beyond every range: the highest


def test_the_shutter_header_is_read() -> None:
    assert shutter_in_view(NS(frame_info=_Info(None, 1, True)))
    assert not shutter_in_view(NS(frame_info=_Info(None, 1, False)))
    assert not shutter_in_view(NS(frame_info=NS(time=None)))  # no header entries (the T650's)


# --- the check -------------------------------------------------------------------------------------------


def test_presets_that_agree_are_joined() -> None:
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2, 3), size=40)
    ranges = {p: r for p, (r, _c, _q) in PRESETS.items()}
    check = check_presets(im, (1, 2, 3), ranges, cycles=8)
    assert check.kept == (1, 2, 3) and not check.left_out
    assert all(pair.agrees for pair in check.pairs)
    assert check.pairs[0].difference_k == pytest.approx(0.0, abs=5.0)  # 5 K bins; each preset's curve at ε 0.95


def test_an_offset_preset_is_left_out() -> None:
    """The 1006 test a: preset 3's counts 4,151 too high read every cold pixel as about 1,200 °C."""
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0})
    ranges = {p: r for p, (r, _c, _q) in PRESETS.items()}
    check = check_presets(im, (1, 2, 3), ranges, cycles=8)
    assert check.kept == (1, 2) and check.left_out == (3,)
    pair = check.pairs[-1]
    assert (pair.lower, pair.upper) == (2, 3) and pair.share > 0.5
    assert pair.difference_k > 400  # where both measure, preset 3 reads hundreds of kelvin hotter
    assert any(line.startswith("Preset 3 left out") for line in check.notes())


# --- the workbook ------------------------------------------------------------------------------------------


def _export(tmp_path, ats, im, *, preset=(1, 2), rois=None, options=None, name="joined") -> openpyxl.Workbook:
    out = tmp_path / f"{name}.xlsx"
    rois = rois or (RoiShape(1, "cursor", ((0.5, 0.5),), "Cold"),  # 500 K: preset 1, below its range
                    RoiShape(2, "cursor", ((1.5, 2.5),), "Battery"),  # ~660 K: preset 1
                    RoiShape(3, "cursor", ((7.5, 7.5),), "Flame"),  # 1,100 K: preset 2
                    RoiShape(4, "rect", ((5.0, 5.0), (8.0, 8.0)), "Box"))
    spec = SourceSpec(ats, rois, ignition_frame=0, preset=preset)
    options = options or ExportOptions(step_s=1.0, extremes=True)
    message = run_export(out, [spec], options, handles={ats.resolve(): im}, replace=[out])
    assert "1/1 recording(s) with live temperatures" in message
    return openpyxl.load_workbook(out, data_only=True)


def _header(book, sheet="Counts", row=5):
    return list(next(book[sheet].iter_rows(min_row=row, max_row=row, values_only=True)))


def test_a_joined_workbook_reads_each_area_in_its_own_preset(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    im = _Superframing()
    book = _export(tmp_path, ats, im)
    head = _header(book)
    assert head[1:4] == ["Frame (preset 1)", "Camera clock", "Frame (preset 2)"]
    presets = [i for i, h in enumerate(head) if h == "Preset"]
    assert len(presets) == 4
    rows = [r for r in book["Counts"].iter_rows(min_row=6, values_only=True) if r[0] is not None]
    used = {name: {r[col] for r in rows} for name, col in zip(("Cold", "Battery", "Flame", "Box"), presets)}
    assert used["Battery"] == {1} and used["Flame"] == {2} and used["Cold"] == {1}
    # each preset-2 frame is the next frame of the cycle (17 ms after its preset-1 frame)
    assert all(r[3] == r[1] + 1 for r in rows if r[1] is not None)
    # Data: the flame's temperature is the SDK's in preset 2, the cold spot below range (blank)
    data = [r for r in book["Data"].iter_rows(min_row=6, values_only=True) if r[0] is not None]
    dhead = _header(book, "Data", 4)
    flame = dhead.index(next(h for h in dhead if h and "Flame" in str(h)))
    cold = dhead.index(next(h for h in dhead if h and "Cold" in str(h)))
    im.reset_object_parameters()
    im.unit = im.fnv.Unit.TEMPERATURE_FACTORY
    im.get_frame(1)  # cycle 0, preset 2
    assert data[0][flame] == pytest.approx(float(im.final[7, 7]) - 273.15, abs=1e-3)
    assert all(r[cold] == "#N/A" for r in data)
    largest = next(row for row in book["Validation"].iter_rows(values_only=True)
                   if row[0] == "Largest |difference| over checked samples (K)")
    assert largest[3] < 1e-3 and largest[4].startswith("PASS")
    notes = [row[9] for row in book["Validation"].iter_rows(min_row=8, values_only=True) if row[9]]
    assert any(note.startswith("preset 2") for note in notes) and any(note.startswith("preset 1") for note in notes)
    source = {row[0]: row[1] for row in book["Source"].iter_rows(values_only=True) if row[0]}
    assert "Preset 2 joined" in source and "lowest preset whose range holds it" in source["Preset 2 joined"]
    assert source["Presets compared"].startswith("Preset 2")  # an 8 x 8 image: too few pixels to judge
    assert "left out" not in source["Presets compared"]


def test_a_joined_workbook_leaves_out_an_offset_preset(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0})
    rois = (RoiShape(1, "cursor", ((39.5, 39.5),), "Flame"),)
    book = _export(tmp_path, ats, im, preset=(1, 2, 3), rois=rois)
    head = _header(book)
    assert "Frame (preset 3)" not in head and "Frame (preset 2)" in head
    lines = [row[1] for row in book["Source"].iter_rows(values_only=True) if row[0] == "Presets compared"]
    assert any(line.startswith("Preset 3 left out") for line in lines)


def test_shutter_rows_have_no_counts(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    im = _Superframing(shutter=range(20, 30))
    book = _export(tmp_path, ats, im)
    dhead = _header(book, "Data", 5)
    status = dhead.index("Status")
    rows = [r for r in book["Data"].iter_rows(min_row=6, values_only=True) if r[0] is not None]
    covered = [r for r in rows if 20 <= r[0] < 30]
    assert len(covered) == 10 and all(r[status] == "Shutter in view" and r[status - 1] == "#N/A" for r in covered)
    assert all(r[status] != "Shutter in view" for r in rows if not 20 <= r[0] < 30)
    source = {row[0]: row[1] for row in book["Source"].iter_rows(values_only=True) if row[0]}
    assert source["Camera shutter"].startswith("10 rows")
    # the Summary's percentages count the shutter rows like its COUNTA does
    summary = next(row for row in book["Summary"].iter_rows(values_only=True) if row[1] == "Cold")
    assert summary[10] == pytest.approx(100.0 * 50 / 60)


def test_one_preset_still_reads_only_its_frames(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    book = _export(tmp_path, ats, _Superframing(), preset=2)
    head = _header(book)
    assert head[1:3] == ["Frame", "Camera clock"] and "Preset" not in head
    frames = [r[1] for r in book["Counts"].iter_rows(min_row=6, values_only=True) if r[1] is not None]
    assert frames and all(frame % 2 == 0 for frame in frames)  # 1-based: the second frame of each cycle


def test_tc_compare_and_the_match_follow_each_rows_preset(tmp_path, ats) -> None:
    pytest.importorskip("fnv")
    im = _Superframing()
    times = tuple(float(t) for t in range(60))
    scene = [im.scene(k) for k in range(60)]
    prefill = TcPrefill(times=times, columns=(tuple(float(s[7, 7]) - 273.15 for s in scene),
                                              tuple(float(s[2, 1]) - 273.15 for s in scene)),
                        names=("Tree 1 m", "T2"), roi_tc=(("Flame", "Tree 1 m"), ("Battery", "T2")))
    book = _export(tmp_path, ats, im, options=ExportOptions(step_s=1.0, tc_prefill=prefill))
    match = book["Emissivity Match"]
    own = {row[0]: row for row in match.iter_rows(values_only=True) if row[0] in ("Flame", "Battery")}
    # every row counts, preset 1 (battery) and preset 2 (flame) alike; the TCs are the scene, a blackbody seen at
    # the recorded ε 0.95, so each ROI's own best value is 1.00
    assert own["Battery"][4] == 60 and own["Flame"][4] == 60
    assert own["Battery"][8] == pytest.approx(1.0) and own["Flame"][8] == pytest.approx(1.0)


# --- the player ---------------------------------------------------------------------------------------------


def _player(ats, monkeypatch, fake):
    from flir_player import source as source_module

    fake.rois = []
    fake.can_change_object_parameters = True
    fake.has_nuc = fake.has_bp = False
    monkeypatch.setattr(source_module.fnv.file, "ImagerFile", lambda _path: fake)
    src = source_module.FlirVideoSource()
    src.open(ats)
    return src


def test_pixels_take_the_lowest_preset_that_holds_them() -> None:
    from flir_player.source import merge_presets

    ranges = {1: (523.15, 873.15), 2: (773.15, 1473.15)}
    kelvin_1 = np.array([[500.0, 700.0], [900.0, 950.0]])  # below, inside, above, clamped high
    status_1 = np.array([[1, 0, 1, 4]])
    kelvin_2 = np.array([[700.0, 702.0], [905.0, 1000.0]])
    status_2 = np.array([[1, 1, 0, 0]])
    values, status, which = merge_presets([(1, kelvin_1, status_1.reshape(2, 2)), (2, kelvin_2, status_2.reshape(2, 2))],
                                          ranges, "KELVIN")
    assert which.tolist() == [[1, 1], [2, 2]]
    assert values.tolist() == [[500.0, 700.0], [905.0, 1000.0]]
    assert status.tolist() == [[1, 0], [0, 0]]


def test_the_player_joins_presets_in_temperature_units(ats, monkeypatch) -> None:
    pytest.importorskip("fnv")
    fake = _Superframing(size=40)
    src = _player(ats, monkeypatch, fake)
    assert src.preset == (1, 2)
    counts = src.read_frame(0)
    assert counts.data[39, 39] == pytest.approx(fake.counts_for(0)[39, 39])  # counts: preset 1's frame as it is
    assert "temperature units join them" in counts.presets_note
    src.set_unit("temperature_factory_k")
    packet = src.read_frame(0)
    assert packet.index == 0 and "Presets 1 and 2 joined (250 to 1200 °C)" in packet.presets_note
    fake.unit = fake.fnv.Unit.TEMPERATURE_FACTORY
    fake.get_frame(1)
    hot = float(fake.final[39, 39])  # 1,100 K: preset 2's frame
    fake.get_frame(0)
    warm = float(fake.final[20, 0])  # ~800 K: inside preset 1
    assert packet.data[39, 39] == pytest.approx(hot) and packet.data[20, 0] == pytest.approx(warm)
    assert packet.maximum == pytest.approx(float(packet.data.max()))  # measured on the merged frame
    src.set_preset(1)
    assert src.read_frame(0).data[39, 39] != pytest.approx(hot)  # one preset: its own (extrapolated) value


def test_the_player_leaves_out_an_offset_preset(ats, monkeypatch) -> None:
    pytest.importorskip("fnv")
    fake = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0})
    src = _player(ats, monkeypatch, fake)
    assert src.preset == (1, 2, 3)
    src.set_unit("temperature_factory_k")
    packet = src.read_frame(0)
    assert "Presets 1 and 2 joined" in packet.presets_note and "Preset 3 left out" in packet.presets_note
    assert float(packet.data.max()) < 1200.0  # preset 3's ~1,400 K cold-frame reading never shows


def test_shutter_frames_are_flagged_in_the_player(ats, monkeypatch) -> None:
    pytest.importorskip("fnv")
    src = _player(ats, monkeypatch, _Superframing(shutter=range(5, 8)))
    assert src.read_frame(10).shutter and not src.read_frame(4).shutter  # frame 10 = cycle 5


# --- the choice in the windows -------------------------------------------------------------------------------


def _metadata_with_presets(tmp_path):
    from flir_player.models import PresetRange, VideoMetadata

    presets = (PresetRange(1, 100, 523.15, 873.15), PresetRange(2, 100, 773.15, 1473.15),
               PresetRange(3, 100, 1123.15, 2273.15))
    return VideoMetadata(path=tmp_path / "a8303.ats", width=64, height=48, num_frames=300, start_time=None,
                         end_time=None, duration_seconds=100.0, nominal_fps=3.0, camera_model="A8303sc",
                         camera_serial="", source_details=(), presets=presets)


def test_the_inspector_offers_the_presets_joined(qapp) -> None:
    from flir_player.widgets import InspectorPanel

    inspector = InspectorPanel()
    presets = _metadata_with_presets(Path(".")).presets
    emitted = []
    inspector.preset_changed.connect(emitted.append)
    inspector.set_presets(presets, (1, 2, 3))
    assert inspector.preset_combo.currentText() == "Presets joined: 250 to 2000 °C"
    inspector.preset_combo.setCurrentIndex(inspector.preset_combo.findData(2))
    inspector.preset_combo.setCurrentIndex(0)
    assert emitted == [2, (1, 2, 3)]
    inspector.select_preset(2)
    assert inspector.preset_combo.currentData() == 2


def test_the_excel_dialog_starts_on_the_players_choice(qapp, tmp_path) -> None:
    from flir_player.excel_dialog import ExcelExportDialog

    metadata = _metadata_with_presets(tmp_path)
    rois = (RoiShape(1, "rect", ((1.0, 1.0), (5.0, 5.0)), "Box 1"),)
    joined = ExcelExportDialog(metadata, rois, current_frame=0, preset=(1, 2, 3))
    assert joined._chosen_preset() == (1, 2, 3) and joined.preset_combo.currentText().startswith("Presets joined")
    one = ExcelExportDialog(metadata, rois, current_frame=0, preset=2)
    assert one._chosen_preset() == 2
    joined.close()
    one.close()


# --- review of 0.10.0 ------------------------------------------------------------------------------------------


RANGES = {p: r for p, (r, _c, _q) in PRESETS.items()}


def test_an_offset_preset_is_left_out_in_a_hot_scene() -> None:
    """Review: with the scene above preset 2's floor nothing contradicts; the median difference tells."""
    pytest.importorskip("fnv")
    hot = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0},
                        scene=lambda k: np.linspace(1200.0, 1600.0, 1600).reshape(40, 40))
    check = check_presets(hot, (1, 2, 3), RANGES, cycles=8)
    assert check.left_out == (3,) and check.pairs[-1].offset and check.pairs[-1].share < 0.5
    assert any(line.startswith("Preset 3 left out") and "against preset 2" in line for line in check.notes())


def test_the_check_reads_the_recordings_own_parameters() -> None:
    """Review: an emissivity set in the player must not change the verdict."""
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0})
    im.object_parameters.emissivity = 0.2
    assert check_presets(im, (1, 2, 3), RANGES, cycles=8).left_out == (3,)


def test_a_higher_frame_under_the_shutter_is_skipped() -> None:
    """Review: preset 1 exposed but preset 2's frame of the cycle covered must not fail preset 2."""
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2), size=40, covered={2 * k + 1 for k in range(60) if k != 30})
    assert check_presets(im, (1, 2), RANGES, cycles=8).kept == (1, 2)


def test_a_choice_that_skips_a_preset_is_still_checked() -> None:
    """Review: presets (1, 3) of a 1-2-3 cycle: preset 3's frame is two frames on."""
    pytest.importorskip("fnv")
    im = _Superframing(presets=(1, 2, 3), size=40, pedestal={3: 4151.0})
    check = check_presets(im, (1, 3), RANGES, cycles=8)
    assert check.left_out == (3,) and check.pairs[0].informative > 1000


def test_a_reference_from_the_recording_is_joined_too(ats, monkeypatch) -> None:
    """Review: a frame minus itself as the reference is zero with presets joined."""
    pytest.importorskip("fnv")
    src = _player(ats, monkeypatch, _Superframing(size=40))
    src.load_reference(ats, 0, "subtract")
    src.set_unit("temperature_factory_k")
    np.testing.assert_array_equal(src.read_frame(0).data, 0.0)
    src.set_preset(1)  # review round 2: the reference follows the presets shown
    np.testing.assert_array_equal(src.read_frame(0).data, 0.0)
    src.set_preset((1, 2))
    np.testing.assert_array_equal(src.read_frame(0).data, 0.0)
    for _ in range(2):  # review round 3: through preset 2 and back, the reference stays on frame 0's cycle
        src.set_preset(2)
        np.testing.assert_array_equal(src.read_frame(1).data, 0.0)  # frame 1: preset 2's frame of that cycle
        src.set_preset((1, 2))
        np.testing.assert_array_equal(src.read_frame(0).data, 0.0)
    assert src._reference_spec[1] == 0


def test_the_median_difference_takes_both_middle_values() -> None:
    """Review round 2: half the pixels agree, half read 700 K hotter: the median is 350 K, an offset."""
    from flir_player.presets import PairCheck, _Tally

    tally = _Tally()
    lower = (np.full(1000, 1200.0), np.zeros(1000, dtype=np.int8))
    upper = (np.r_[np.full(500, 1200.0), np.full(500, 1900.0)], np.zeros(1000, dtype=np.int8))
    tally.add(lower, upper, (773.15, 1473.15), (1123.15, 2273.15))
    assert tally.median() == pytest.approx(350.0, abs=5.0)
    pair = PairCheck(2, 3, tally.informative, tally.contradicting, tally.both, tally.median())
    assert pair.contradicting == 0 and pair.offset and not pair.agrees


def test_a_temporal_refetch_keeps_the_merged_clip_mask(ats, monkeypatch) -> None:
    """Review: reading the shown frame again (ROI edit) must not take the lowest preset's clipping."""
    pytest.importorskip("fnv")
    from flir_player.processing import ProcessingState

    src = _player(ats, monkeypatch, _Superframing(size=40, clamp_above=150.0))
    src.set_unit("temperature_factory_k")
    src.set_processing(ProcessingState(temporal=("average", 2)))
    first = src.read_frame(2)
    again = src.read_frame(2)
    np.testing.assert_array_equal(first.data, again.data)
    assert (first.clip_mask is None) == (again.clip_mask is None)
    if first.clip_mask is not None:
        np.testing.assert_array_equal(first.clip_mask, again.clip_mask)
    src.set_preset(1)  # preset 1 alone clamps the hottest pixels: the mask is not empty by construction
    assert src.read_frame(2).clip_mask is not None


class _Dropped(_Superframing):
    """Cycle 1 lost its preset-2 frame (dropped): frame 2 is cycle 1's preset 1, frame 3 cycle 2's."""

    def __init__(self, **kwargs):
        self.frames = None
        super().__init__(**kwargs)

    def where(self, index):
        if self.frames is None:
            k = len(self.presets)
            self.frames = [(c, m, p) for c in range(self.num_frames // k) for m, p in enumerate(self.presets)
                           if (c, p) != (1, 2)]
            self.num_frames = len(self.frames)
        return self.frames[index]


def test_a_reference_cycle_without_the_preset_is_refused(ats, monkeypatch) -> None:
    """Review round 4: no preset-2 frame in the reference's cycle: refuse, never another cycle's frame."""
    pytest.importorskip("fnv")
    fake = _Dropped(size=8)
    src = _player(ats, monkeypatch, fake)
    src.set_unit("temperature_factory_k")
    src.load_reference(ats, 2, "subtract")  # cycle 1's preset-1 frame
    with pytest.raises(ValueError, match="dropped"):
        src.set_preset(2)
    assert src.preset == (1, 2) and src._reference_spec[1] == 2  # the choice and the reference are kept
    np.testing.assert_array_equal(src.read_frame(2).data, 0.0)


def test_the_reference_label_follows_a_preset_change(qapp, ats, monkeypatch) -> None:
    """Review round 4: the inspector names the frame the reference now reads."""
    pytest.importorskip("fnv")
    from conftest import wait_until
    from flir_player import source as source_module
    from flir_player.main_window import MainWindow

    fake = _Superframing(size=8)
    fake.rois = []
    fake.can_change_object_parameters = True
    fake.has_nuc = fake.has_bp = False
    monkeypatch.setattr(source_module.fnv.file, "ImagerFile", lambda _path: fake)
    window = MainWindow()
    window.show()
    window.open_path(ats)
    try:
        assert wait_until(qapp, lambda: window.current_packet is not None and not window._busy)
        labels = []
        window.decoder.reference_ready.connect(labels.append)
        window.decoder.request_reference({"path": str(ats), "frame_index": 0, "op": "subtract"}, 0, 101)
        assert wait_until(qapp, lambda: labels and not window._busy)
        assert labels[-1].endswith("frame 1")
        window.decoder.request_preset(2, 0, 102)
        assert wait_until(qapp, lambda: len(labels) >= 2 and not window._busy)
        assert labels[-1].endswith("frame 2")  # preset 2's frame of the reference's cycle
    finally:
        window.close()
        qapp.processEvents()
